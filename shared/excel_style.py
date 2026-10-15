"""One house style for every Excel workbook the app writes.

Change the font or colors here and every forecast workbook, extract and
validation report follows. The font matches the Actuals validation reports.
"""

from __future__ import annotations

from copy import copy

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

FONT_NAME = "Aptos"
BODY_SIZE = 10
TITLE_SIZE = 16
# Cells at least this large in a source workbook are treated as titles: they
# get the house font but keep their size.
TITLE_THRESHOLD = 12

NAVY = "243F5C"
SLATE = "44546A"
MUTED = "6B7785"
BORDER = "D0D7DE"
BAND = "F7F9FB"
WHITE = "FFFFFF"

# status: (fill, text color)
STATUS_COLORS = {
    "PASS": ("E2F0D9", "1E6B1E"),
    "RECONCILED": ("E2F0D9", "1E6B1E"),
    "FAIL": ("F4CCCC", "9C0006"),
    "REVIEW REQUIRED": ("F4CCCC", "9C0006"),
    "ACTION REQUIRED": ("FCE4D6", "9C4A00"),
    "NOT VALIDATED": ("FFF2CC", "8A5A00"),
    "INCOMPLETE": ("FFF2CC", "8A5A00"),
    "NOT SET UP": ("EDEDED", "595959"),
    "NOT RUN": ("EDEDED", "595959"),
    "NO RESULTS": ("EDEDED", "595959"),
}

MONEY = '#,##0.00;[Red](#,##0.00);"–"'
COUNT = '0;-0;"–"'

THIN = Side(style="thin", color=BORDER)
GRID = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
UNDERLINE = Border(bottom=THIN)


def font(size: float = BODY_SIZE, **kwargs) -> Font:
    return Font(name=FONT_NAME, size=size, **kwargs)


def fill(color: str) -> PatternFill:
    return PatternFill("solid", fgColor=color)


def status_style(cell, status: str) -> None:
    background, text = STATUS_COLORS.get(str(status or "").upper(), ("FFFFFF", "000000"))
    cell.fill = fill(background)
    cell.font = font(bold=True, color=text)
    cell.alignment = Alignment(horizontal="center", vertical="center")


def header_row(ws, row: int, first_col: int, last_col: int) -> None:
    for col in range(first_col, last_col + 1):
        cell = ws.cell(row, col)
        cell.font = font(bold=True, color=WHITE)
        cell.fill = fill(NAVY)
        cell.border = GRID
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def normalize_fonts(ws) -> int:
    """Give every cell of ``ws`` the house font and body size.

    Bold, italic, underline, strike-through and color are kept, so headers
    and totals still stand out. Title-sized cells keep their size. Returns
    the number of cells changed.
    """
    changed = 0
    for row in ws.iter_rows():
        for cell in row:
            current = cell.font
            size = current.sz if current.sz and current.sz >= TITLE_THRESHOLD else BODY_SIZE
            if current.name == FONT_NAME and current.sz == size:
                continue
            new = copy(current)
            new.name = FONT_NAME
            new.sz = size
            cell.font = new
            changed += 1
    return changed


def finish_sheet(ws, *, landscape: bool = True, gridlines: bool = False) -> None:
    """Common page setup: gridlines, print orientation, fit to one page wide."""
    ws.sheet_view.showGridLines = gridlines
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
