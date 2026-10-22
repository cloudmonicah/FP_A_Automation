from __future__ import annotations

import json
import os
from datetime import date, datetime
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def _counts(results):
    out = {"PASS": 0, "FAIL": 0, "ACTION REQUIRED": 0, "NOT VALIDATED": 0}
    for x in results:
        s = str(x.get("status") or "").upper()
        out[s] = out.get(s, 0) + 1
    return out


def merge_saved_results(
    output_folder,
    forecast_month,
    forecast_year,
    forecast_type,
    categories_run,
    results,
) -> list[dict]:
    """Keep the latest results of every category for the month's report.

    The results are saved beside the report as a small .results.json file.
    Running one category replaces that category's rows; results saved for
    another forecast month are ignored.
    """
    period = date(forecast_year, forecast_month, 1)
    store_path = Path(output_folder) / f"Forecast_Validation_{period:%b}_{forecast_year}.results.json"

    saved = {}
    try:
        data = json.loads(store_path.read_text(encoding="utf-8"))
        if data.get("forecast_type") == forecast_type and isinstance(data.get("categories"), dict):
            for category, rows in data["categories"].items():
                # An earlier version saved {"validated_at": ..., "results": [...]}.
                if isinstance(rows, dict):
                    rows = rows.get("results")
                if isinstance(rows, list):
                    saved[category] = [row for row in rows if isinstance(row, dict)]
    except (OSError, ValueError, AttributeError):
        pass

    def plain(value):
        return value.isoformat() if isinstance(value, (date, datetime)) else value

    for category in categories_run:
        saved[category] = [
            {key: plain(value) for key, value in item.items()}
            for item in results
            if item.get("category", "") == category
        ]

    Path(output_folder).mkdir(parents=True, exist_ok=True)
    staged = store_path.with_name(f"~staging_{store_path.name}")
    staged.write_text(
        json.dumps({"forecast_type": forecast_type, "categories": saved}, indent=1, default=str),
        encoding="utf-8",
    )
    os.replace(staged, store_path)

    order = ["OPEX", "Direct Cost", "Revenue", "Other"]
    ordered = [c for c in order if c in saved] + [c for c in saved if c not in order]
    return [item for category in ordered for item in saved[category]]


def write_forecast_validation_report(
    validation_result,
    output_folder,
    fiscal_year,
    forecast_month,
    forecast_year,
    forecast_type,
    tolerance=1000.00,
    forecast_start=None,
):
    output_folder = Path(output_folder)
    output_folder.mkdir(parents=True, exist_ok=True)
    period = datetime(forecast_year, forecast_month, 1)
    path = output_folder / f"Forecast_Validation_{period:%b}_{forecast_year}.xlsx"
    results = list(validation_result.get("results") or [])
    counts = _counts(results)
    overall = (
        "RECONCILED"
        if not (
            counts.get("FAIL")
            or counts.get("ACTION REQUIRED")
            or counts.get("NOT VALIDATED")
        )
        else "REVIEW REQUIRED"
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "Validation Summary"
    detail = wb.create_sheet("Validation Detail")

    navy = PatternFill("solid", fgColor="1F4E78")
    blue = PatternFill("solid", fgColor="D9EAF7")
    green = PatternFill("solid", fgColor="E2F0D9")
    red = PatternFill("solid", fgColor="FCE4D6")
    yellow = PatternFill("solid", fgColor="FFF2CC")
    edge = Border(bottom=Side(style="thin", color="B7B7B7"))

    ws.merge_cells("A1:F1")
    c = ws["A1"]
    c.value = "FORECAST VALIDATION REPORT"
    c.font = Font(size=18, bold=True, color="FFFFFF")
    c.fill = navy
    ws.row_dimensions[1].height = 28

    rows = [
        ("Forecast Period", period.strftime("%B %Y")),
        ("Fiscal Year", fiscal_year),
        ("Forecast Start", forecast_start or period),
        ("Forecast Type", forecast_type),
        ("Tolerance", tolerance),
        ("Routes Reported", len(results)),
        ("Passed", counts.get("PASS", 0)),
        ("Failed", counts.get("FAIL", 0)),
        ("Action Required", counts.get("ACTION REQUIRED", 0)),
        ("Not Validated", counts.get("NOT VALIDATED", 0)),
        ("Overall Status", overall),
        ("Last Updated", datetime.now().strftime("%m/%d/%Y %I:%M %p")),
    ]
    for r, (label, value) in enumerate(rows, start=3):
        ws.cell(r, 1, label).font = Font(bold=True)
        ws.cell(r, 2, value)
        ws.cell(r, 1).border = edge
        ws.cell(r, 2).border = edge
    ws["B5"].number_format = "mm/dd/yyyy"
    ws["B7"].number_format = "$#,##0.00;[Red]-$#,##0.00"
    ws["B13"].fill = green if overall == "RECONCILED" else yellow
    ws["B13"].font = Font(bold=True)
    ws["A15"] = "Status Guide"
    ws["A15"].font = Font(bold=True)
    guides = [
        (
            "PASS",
            "Excel and individual Tableau Grand Totals reconcile within tolerance.",
        ),
        ("FAIL", "Validation completed but the variance exceeds tolerance."),
        ("ACTION REQUIRED", "Validation could not complete and requires review."),
        (
            "NOT VALIDATED",
            "Not checked automatically: the individual Tableau tab needs an extra "
            "department filter first (see Note).",
        ),
    ]
    for r, (s, m) in enumerate(guides, start=16):
        ws.cell(r, 1, s).font = Font(bold=True)
        ws.cell(r, 2, m)
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 75

    headers = [
        "Forecast Workbook",
        "Worksheet",
        "Individual Tableau Workbook",
        "Individual Tableau Worksheet",
        "Excel Grand Total",
        "Tableau Grand Total",
        "Variance",
        "Status",
        "Period Mode",
        "Forecast Start",
        "Forecast Type",
        "Note",
        "Category",
        "Validated At",
    ]
    for col, h in enumerate(headers, 1):
        c = detail.cell(1, col, h)
        c.font = Font(bold=True)
        c.fill = blue
        c.border = edge
        c.alignment = Alignment(horizontal="center", vertical="center")

    for r, item in enumerate(results, 2):
        vals = [
            item.get("forecast_workbook"),
            item.get("worksheet"),
            item.get("individual_tableau_workbook"),
            item.get("individual_tableau_worksheet"),
            item.get("excel_grand_total"),
            item.get("tableau_grand_total"),
            item.get("variance"),
            item.get("status"),
            item.get("period_mode"),
            item.get("forecast_start"),
            item.get("forecast_type"),
            item.get("note"),
            item.get("category"),
            item.get("validated_at"),
        ]
        for col, v in enumerate(vals, 1):
            detail.cell(r, col, v)
            detail.cell(r, col).border = edge
        for col in (5, 6, 7):
            detail.cell(r, col).number_format = "$#,##0.00;[Red]-$#,##0.00"
        s = str(item.get("status") or "").upper()
        sc = detail.cell(r, 8)
        sc.font = Font(bold=True)
        sc.fill = green if s == "PASS" else red if s == "FAIL" else yellow

    detail.freeze_panes = "A2"
    detail.auto_filter.ref = detail.dimensions
    widths = {
        1: 46,
        2: 28,
        3: 42,
        4: 38,
        5: 18,
        6: 20,
        7: 16,
        8: 18,
        9: 18,
        10: 18,
        11: 26,
        12: 70,
        13: 16,
        14: 20,
    }
    for col, w in widths.items():
        detail.column_dimensions[get_column_letter(col)].width = w
    for row in detail.iter_rows():
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)

    # Save beside the report and swap it in, so a report open in Excel is
    # reported clearly and a half-written file is never left behind.
    staged = path.with_name(f"~saving_{path.name}")
    wb.save(staged)
    try:
        os.replace(staged, path)
    except PermissionError as exc:
        staged.unlink(missing_ok=True)
        from shared.friendly_errors import PlainLanguageError

        raise PlainLanguageError(
            f'The validation report "{path.name}" is open (in Excel, Teams or the browser), '
            "so it could not be updated. Close it everywhere and click Validate again."
        ) from exc
    wb.close()
    return path
