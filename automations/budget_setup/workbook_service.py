from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook


def normalize_text(value: Any) -> str:
    return str(value or "").strip().lower()


def is_formula(value: Any) -> bool:
    return isinstance(value, str) and value.startswith("=")


def find_header(
    worksheet,
    names: set[str],
) -> tuple[int | None, int | None]:
    for row_index in range(1, worksheet.max_row + 1):
        for column_index in range(1, worksheet.max_column + 1):
            value = normalize_text(
                worksheet.cell(
                    row=row_index,
                    column=column_index,
                ).value
            )

            if value in names:
                return row_index, column_index

    return None, None


def find_month_columns(
    worksheet,
    header_row: int,
) -> list[int]:
    grand_total_column = None

    for column_index in range(1, worksheet.max_column + 1):
        value = normalize_text(
            worksheet.cell(
                row=header_row,
                column=column_index,
            ).value
        )

        if value in {
            "grand total",
            "grandtotal",
            "total",
        }:
            grand_total_column = column_index
            break

    if grand_total_column is None:
        return []

    if grand_total_column <= 12:
        return []

    return list(
        range(
            grand_total_column - 12,
            grand_total_column,
        )
    )


def build_fiscal_month_headers(
    target_fiscal_year: str,
) -> list[str]:
    fiscal_number = int(target_fiscal_year.replace("FY", ""))

    start_year = 2000 + fiscal_number - 1
    end_year = start_year + 1

    return [
        f"Jul-{start_year % 100:02d}",
        f"Aug-{start_year % 100:02d}",
        f"Sep-{start_year % 100:02d}",
        f"Oct-{start_year % 100:02d}",
        f"Nov-{start_year % 100:02d}",
        f"Dec-{start_year % 100:02d}",
        f"Jan-{end_year % 100:02d}",
        f"Feb-{end_year % 100:02d}",
        f"Mar-{end_year % 100:02d}",
        f"Apr-{end_year % 100:02d}",
        f"May-{end_year % 100:02d}",
        f"Jun-{end_year % 100:02d}",
    ]


def update_forecast_type(
    worksheet,
    header_row: int,
    forecast_column: int,
    target_fiscal_year: str,
) -> int:
    worksheet.cell(
        row=header_row,
        column=forecast_column,
    ).value = "Forecast Type"

    replacement_value = f"{target_fiscal_year} Forecast"

    rows_updated = 0

    for row_index in range(
        header_row + 1,
        worksheet.max_row + 1,
    ):
        cell = worksheet.cell(
            row=row_index,
            column=forecast_column,
        )

        current_value = normalize_text(cell.value)

        if not current_value:
            continue

        if is_formula(cell.value):
            continue

        if "budget" in current_value or "forecast" in current_value:
            cell.value = replacement_value
            rows_updated += 1

    return rows_updated


def update_month_headers(
    worksheet,
    header_row: int,
    month_columns: list[int],
    target_fiscal_year: str,
) -> None:
    headers = build_fiscal_month_headers(target_fiscal_year)

    for column_index, header in zip(
        month_columns,
        headers,
    ):
        worksheet.cell(
            row=header_row,
            column=column_index,
        ).value = header


def convert_budget_workbook(
    workbook_path: str | Path,
    target_fiscal_year: str,
) -> dict:
    """
    Convert a copied budget workbook into the initial
    fiscal-year forecast workbook.

    All monthly budget values are retained.
    """
    workbook_path = Path(workbook_path)

    workbook = load_workbook(
        workbook_path,
        data_only=False,
    )

    worksheets_updated = 0
    forecast_values_updated = 0
    skipped_sheets: list[str] = []

    for worksheet in workbook.worksheets:
        header_row, forecast_column = find_header(
            worksheet,
            {
                "budget type",
                "forecast type",
            },
        )

        if header_row is None or forecast_column is None:
            skipped_sheets.append(f"{worksheet.title}: Budget Type not found")
            continue

        month_columns = find_month_columns(
            worksheet,
            header_row,
        )

        if len(month_columns) != 12:
            skipped_sheets.append(f"{worksheet.title}: 12 month columns not found")
            continue

        updated_values = update_forecast_type(
            worksheet=worksheet,
            header_row=header_row,
            forecast_column=forecast_column,
            target_fiscal_year=target_fiscal_year,
        )

        update_month_headers(
            worksheet=worksheet,
            header_row=header_row,
            month_columns=month_columns,
            target_fiscal_year=target_fiscal_year,
        )

        forecast_values_updated += updated_values
        worksheets_updated += 1

        print(
            f"{worksheet.title}: Updated successfully. "
            "All monthly budget values retained."
        )

    workbook.save(workbook_path)

    return {
        "worksheets_updated": worksheets_updated,
        "forecast_values_updated": forecast_values_updated,
        "skipped_sheets": skipped_sheets,
    }


def update_forecast_workbook(
    workbook_path: str | Path,
    target_fiscal_year: str,
    forecast_month: int,
    forecast_year: int,
) -> dict:
    """
    Prepare an existing fiscal-year forecast workbook
    for the selected forecast month.

    Existing month values, formulas, headers, and formatting
    are preserved. Only Forecast Type values are updated.
    """
    workbook_path = Path(workbook_path)

    workbook = load_workbook(
        workbook_path,
        data_only=False,
    )

    forecast_type_value = (
        f"{target_fiscal_year} Forecast_" f"{forecast_month:02d}" f"{forecast_year}"
    )

    worksheets_updated = 0
    forecast_values_updated = 0
    skipped_sheets: list[str] = []

    for worksheet in workbook.worksheets:
        header_row, forecast_column = find_header(
            worksheet,
            {
                "forecast type",
            },
        )

        if header_row is None or forecast_column is None:
            skipped_sheets.append(f"{worksheet.title}: Forecast Type not found")
            continue

        rows_updated = 0

        for row_index in range(
            header_row + 1,
            worksheet.max_row + 1,
        ):
            cell = worksheet.cell(
                row=row_index,
                column=forecast_column,
            )

            if is_formula(cell.value):
                continue

            current_value = normalize_text(cell.value)

            if not current_value:
                continue

            if "forecast" in current_value:
                cell.value = forecast_type_value
                rows_updated += 1

        if rows_updated > 0:
            worksheets_updated += 1
            forecast_values_updated += rows_updated

            print(
                f"{worksheet.title}: "
                f"Forecast Type updated to "
                f"{forecast_type_value}. "
                "Existing monthly values retained."
            )

    workbook.save(workbook_path)

    return {
        "worksheets_updated": worksheets_updated,
        "forecast_values_updated": forecast_values_updated,
        "forecast_type": forecast_type_value,
        "skipped_sheets": skipped_sheets,
    }


# ============================================================
# FORMULA AUDIT
# ============================================================


def get_formula_cells(
    workbook_path: str | Path,
) -> dict[str, dict[str, str]]:
    """
    Return every Excel formula contained in a workbook.

    Example:

        {
            "ALIV": {
                "T2": "=SUM(H2:S2)",
                "H12": "=SUM(H2:H11)",
            }
        }

    This function DOES NOT modify the workbook.
    """

    workbook_path = Path(workbook_path)

    workbook = load_workbook(
        workbook_path,
        data_only=False,
    )

    formulas = {}

    try:

        for worksheet in workbook.worksheets:

            sheet_formulas = {}

            for row in worksheet.iter_rows():

                for cell in row:

                    if isinstance(cell.value, str) and cell.value.startswith("="):

                        sheet_formulas[cell.coordinate] = cell.value

            formulas[worksheet.title] = sheet_formulas

    finally:

        workbook.close()

    return formulas


def count_formulas(
    formula_map: dict[str, dict[str, str]],
) -> int:
    """
    Count all formulas across the workbook.
    """

    return sum(len(sheet_formulas) for sheet_formulas in formula_map.values())


def compare_formula_preservation(
    source_formulas: dict[str, dict[str, str]],
    destination_formulas: dict[str, dict[str, str]],
) -> dict:
    """
    Compare formulas before and after workbook processing.

    A formula is considered preserved only when the same
    worksheet/cell still contains the same formula.
    """

    missing_formulas = []
    changed_formulas = []

    for sheet_name, formulas in source_formulas.items():

        destination_sheet = destination_formulas.get(
            sheet_name,
            {},
        )

        for coordinate, source_formula in formulas.items():

            destination_formula = destination_sheet.get(coordinate)

            if destination_formula is None:

                missing_formulas.append(
                    {
                        "sheet": sheet_name,
                        "cell": coordinate,
                        "formula": source_formula,
                    }
                )

                continue

            if destination_formula != source_formula:

                changed_formulas.append(
                    {
                        "sheet": sheet_name,
                        "cell": coordinate,
                        "before": source_formula,
                        "after": destination_formula,
                    }
                )

    return {
        "source_formula_count": (count_formulas(source_formulas)),
        "destination_formula_count": (count_formulas(destination_formulas)),
        "missing_formula_count": (len(missing_formulas)),
        "changed_formula_count": (len(changed_formulas)),
        "missing_formulas": missing_formulas,
        "changed_formulas": changed_formulas,
        "formulas_preserved": (not missing_formulas and not changed_formulas),
    }
