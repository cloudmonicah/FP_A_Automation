from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any
import csv

from openpyxl import load_workbook

# ============================================================
# TABLEAU -> EXCEL FORECAST UPDATE SERVICE
# ============================================================
#
# PURPOSE
# -------
# Shared helper functions used by the Tableau -> Excel forecast
# automation.
#
# IMPORTANT BUSINESS RULES
# ------------------------
# 1. Excel workbook headers are NOT renamed to Tableau headers.
# 2. Different Excel workbooks may use different header names.
# 3. Approved aliases are used to understand those differences.
# 4. Fiscal Year is supplied by the user from the GUI.
# 5. Forecast Year remains a separate GUI value.
# 6. Jul-Jun month columns are identified dynamically.
# 7. Existing formulas must never be overwritten.
# 8. Existing account/descriptive fields remain unchanged.
# 9. Existing workbook structure remains intact.
#
# NOTE
# ----
# Workbook-wide formatting, fiscal-header correction and freeze
# pane handling are performed by the standalone forecast service.
# This module provides the shared parsing/mapping functions used
# by that service.
# ============================================================


# ============================================================
# TABLEAU -> EXCEL FIELD DEFINITIONS
# ============================================================
#
# Tableau has relatively consistent field names.
#
# Excel does NOT.
#
# For example, Tableau may return:
#
#     GL Business Code
#
# while one Excel workbook may contain:
#
#     Bus Code Name
#
# and another may contain:
#
#     Business Code
#
# The alias configuration below lets Python understand that
# these columns represent the same business field WITHOUT
# changing the existing Excel header.
# ============================================================


KEY_COLUMN_MAPPING = {
    "GL Business Code": "Bus Code Name",
    "Company": "Company",
    "GL FS Category": "GL FS Category",
    "GL Account Mapping": "GL Account Mapping",
    "GL Sub Segment code": "GL Sub Segment Code",
    "GL Sub Account Description": "GL Sub Account Description",
}


# ============================================================
# EXCEL HEADER ALIASES
# ============================================================
#
# Add new aliases here if another departmental workbook uses
# different wording.
#
# IMPORTANT:
# These aliases are used only for MATCHING.
# The Excel header itself is NOT renamed.
# ============================================================


EXCEL_HEADER_ALIASES = {
    "Bus Code Name": [
        "Bus Code Name",
        "GL Business Code",
        "Business Code",
        "Bus Code",
        "Business Code Name",
    ],
    "Company": [
        "Company",
        "Company Name",
        "Entity",
        "Entity Name",
    ],
    "GL FS Category": [
        "GL FS Category",
        "FS Category",
        "Financial Statement Category",
    ],
    "GL Account Mapping": [
        "GL Account Mapping",
        "Account Mapping",
        "GL Mapping",
    ],
    "GL Sub Segment Code": [
        "GL Sub Segment Code",
        "GL Sub Segment code",
        "GL Sub Segment",
        "Sub Segment Code",
        "Sub Segment",
        "Account",
        "Account Code",
        "GL Account",
    ],
    "GL Sub Account Description": [
        "GL Sub Account Description",
        "GL Sub Account Desc",
        "Sub Account Description",
        "Sub Account Desc",
        "Account Description",
        "Description",
    ],
}


# ============================================================
# MONTH CONFIGURATION
# ============================================================


MONTH_NAMES = {
    "jan": 1,
    "january": 1,
    "feb": 2,
    "february": 2,
    "mar": 3,
    "march": 3,
    "apr": 4,
    "april": 4,
    "may": 5,
    "jun": 6,
    "june": 6,
    "jul": 7,
    "july": 7,
    "aug": 8,
    "august": 8,
    "sep": 9,
    "sept": 9,
    "september": 9,
    "oct": 10,
    "october": 10,
    "nov": 11,
    "november": 11,
    "dec": 12,
    "december": 12,
}


MONTH_ABBREVIATIONS = {
    1: "Jan",
    2: "Feb",
    3: "Mar",
    4: "Apr",
    5: "May",
    6: "Jun",
    7: "Jul",
    8: "Aug",
    9: "Sep",
    10: "Oct",
    11: "Nov",
    12: "Dec",
}


# Fiscal-year display order.
#
# Cable Bahamas fiscal year:
#
#     July -> June
#
FISCAL_MONTH_ORDER = (
    7,
    8,
    9,
    10,
    11,
    12,
    1,
    2,
    3,
    4,
    5,
    6,
)


# ============================================================
# NORMALIZE TEXT
# ============================================================


def normalize_text(
    value: Any,
) -> str:
    """
    Normalize text so Tableau and Excel values can be compared
    reliably.

    Examples:

        " OTHER STAFF COSTS "
            ->
        "other staff costs"

        "GL   Business   Code"
            ->
        "gl business code"
    """

    return " ".join(str(value or "").strip().lower().split())


# ============================================================
# FIND EXCEL COLUMN BY ALIAS
# ============================================================


def find_excel_column_by_alias(
    header_lookup: dict[str, int],
    canonical_header: str,
) -> int | None:
    """
    Find an Excel column using any approved alias.

    Example:

        canonical_header:
            "Bus Code Name"

        Workbook may actually contain:
            "Business Code"

    The function returns the existing Excel column number.

    IMPORTANT:
    The workbook header is NOT changed.
    """

    aliases = EXCEL_HEADER_ALIASES.get(
        canonical_header,
        [canonical_header],
    )

    for alias in aliases:

        normalized_alias = normalize_text(alias)

        column_index = header_lookup.get(normalized_alias)

        if column_index is not None:
            return column_index

    return None


# ============================================================
# PARSE FISCAL YEAR
# ============================================================


def parse_fiscal_year(
    fiscal_year: str | int,
) -> int:
    """
    Convert the Fiscal Year selected by the user into the
    four-digit fiscal ending year.

    Accepted examples:

        FY27 -> 2027
        27   -> 2027
        2027 -> 2027

        FY28 -> 2028

    IMPORTANT:
    Fiscal Year is NOT the same thing as Forecast Year.

    Example:

        Fiscal Year:
            FY27

        Forecast Month:
            August

        Forecast Year:
            2026

    The fiscal calendar is Jul-26 through Jun-27 while Tableau
    Forecast Start is 08/01/2026.
    """

    text = str(fiscal_year).strip().upper().replace("FY", "")

    if not text.isdigit():

        raise ValueError(
            f'Invalid fiscal year "{fiscal_year}". '
            'Expected a value such as "FY27" or "2027".'
        )

    year = int(text)

    # --------------------------------------------------------
    # Convert a two-digit fiscal year to four digits.
    #
    # FY27 -> 2027
    # FY28 -> 2028
    # --------------------------------------------------------

    if year < 100:
        year += 2000

    return year


# ============================================================
# BUILD FISCAL MONTH KEYS
# ============================================================


def build_fiscal_month_keys(
    fiscal_year: str | int,
) -> dict[int, tuple[int, int]]:
    """
    Build the correct (year, month) value for every month in
    the selected Fiscal Year.

    Example:

        fiscal_year = FY27

    returns:

        7  -> (2026, 7)
        8  -> (2026, 8)
        9  -> (2026, 9)
        10 -> (2026, 10)
        11 -> (2026, 11)
        12 -> (2026, 12)

        1  -> (2027, 1)
        2  -> (2027, 2)
        3  -> (2027, 3)
        4  -> (2027, 4)
        5  -> (2027, 5)
        6  -> (2027, 6)
    """

    fiscal_end_year = parse_fiscal_year(fiscal_year)

    fiscal_start_year = fiscal_end_year - 1

    month_keys = {}

    for month_number in FISCAL_MONTH_ORDER:

        # ----------------------------------------------------
        # July-December belong to the calendar year BEFORE
        # the fiscal ending year.
        #
        # January-June belong to the fiscal ending year.
        # ----------------------------------------------------

        if month_number >= 7:
            calendar_year = fiscal_start_year

        else:
            calendar_year = fiscal_end_year

        month_keys[month_number] = (
            calendar_year,
            month_number,
        )

    return month_keys


# ============================================================
# BUILD FISCAL MONTH HEADERS
# ============================================================


def build_fiscal_month_headers(
    fiscal_year: str | int,
) -> dict[int, str]:
    """
    Build the approved Excel month headers for the selected
    Fiscal Year.

    FY27:

        Jul_26
        Aug_26
        Sep_26
        Oct_26
        Nov_26
        Dec_26
        Jan_27
        Feb_27
        Mar_27
        Apr_27
        May_27
        Jun_27

    FY28 automatically becomes:

        Jul_27
        ...
        Jun_28

    No code change is required for the next fiscal year.
    """

    fiscal_end_year = parse_fiscal_year(fiscal_year)

    fiscal_start_year = fiscal_end_year - 1

    headers = {}

    for month_number in FISCAL_MONTH_ORDER:

        # ----------------------------------------------------
        # Determine which calendar year belongs to the month.
        # ----------------------------------------------------

        if month_number >= 7:
            calendar_year = fiscal_start_year

        else:
            calendar_year = fiscal_end_year

        # ----------------------------------------------------
        # Excel display format:
        #
        #     Jul_26
        #     Jan_27
        # ----------------------------------------------------

        year_suffix = str(calendar_year)[-2:]

        headers[month_number] = f"{MONTH_ABBREVIATIONS[month_number]}" f"_{year_suffix}"

    return headers


# ============================================================
# NORMALIZE MONTH HEADER
# ============================================================


def normalize_month_header(
    value: Any,
) -> tuple[int, int] | None:
    """
    Convert a Tableau or Excel month header into:

        (year, month)

    Supported examples:

        "August 2026"
            ->
        (2026, 8)

        "Aug-26"
            ->
        (2026, 8)

        "Aug_26"
            ->
        (2026, 8)

        "Aug/26"
            ->
        (2026, 8)

        Excel datetime
            ->
        (year, month)

    IMPORTANT:
    This function reads what a header CURRENTLY says.

    The standalone forecast service will later use the
    user-selected Fiscal Year to correct incorrect Excel years.
    """

    if value is None:
        return None

    # --------------------------------------------------------
    # Real Excel datetime.
    # --------------------------------------------------------

    if isinstance(value, datetime):

        return (
            value.year,
            value.month,
        )

    text = str(value).strip()

    if not text:
        return None

    # --------------------------------------------------------
    # Normalize common separators.
    #
    # Examples:
    #
    #     Aug-26
    #     Aug_26
    #     Aug/26
    #
    # become:
    #
    #     Aug 26
    # --------------------------------------------------------

    parts = text.replace("-", " ").replace("_", " ").replace("/", " ").split()

    if len(parts) != 2:
        return None

    month_text = parts[0].strip().lower()

    year_text = parts[1].strip()

    if month_text not in MONTH_NAMES:
        return None

    try:
        year = int(year_text)

    except ValueError:
        return None

    # --------------------------------------------------------
    # Convert two-digit year:
    #
    #     26 -> 2026
    #     27 -> 2027
    # --------------------------------------------------------

    if year < 100:
        year += 2000

    return (
        year,
        MONTH_NAMES[month_text],
    )


# ============================================================
# GET MONTH NUMBER FROM HEADER
# ============================================================


def get_month_number_from_header(
    value: Any,
) -> int | None:
    """
    Identify the MONTH from an Excel/Tableau header while
    intentionally ignoring the year.

    This is important because an existing Excel workbook may
    contain an incorrect year.

    Example:

        Jan_26

    inside FY27 is still recognized as JANUARY.

    The standalone service can therefore correct it to:

        Jan_27

    rather than failing because the old year was wrong.
    """

    if value is None:
        return None

    # --------------------------------------------------------
    # Excel date/datetime.
    # --------------------------------------------------------

    if isinstance(value, datetime):
        return value.month

    text = str(value).strip()

    if not text:
        return None

    # --------------------------------------------------------
    # Normalize common separators.
    # --------------------------------------------------------

    parts = text.replace("-", " ").replace("_", " ").replace("/", " ").split()

    if not parts:
        return None

    month_text = parts[0].strip().lower()

    return MONTH_NAMES.get(month_text)


# ============================================================
# CONVERT TABLEAU NUMBER
# ============================================================


def convert_number(
    value: Any,
) -> Any:
    """
    Convert Tableau values into numeric Excel values.

    Examples:

        "24,033"
            ->
        24033

        "$1,354"
            ->
        1354

        "(500)"
            ->
        -500

    Non-numeric values are returned unchanged.
    """

    if value is None:
        return None

    if isinstance(
        value,
        (
            int,
            float,
        ),
    ):
        return value

    text = str(value).strip()

    if not text:
        return None

    # --------------------------------------------------------
    # Accounting negative:
    #
    #     (500)
    #
    # means:
    #
    #     -500
    # --------------------------------------------------------

    negative = text.startswith("(") and text.endswith(")")

    cleaned = (
        text.replace("$", "").replace(",", "").replace("(", "").replace(")", "").strip()
    )

    try:

        number = float(cleaned)

        if negative:
            number *= -1

        return number

    except ValueError:

        # ----------------------------------------------------
        # If Tableau returned text rather than a number,
        # preserve the value.
        # ----------------------------------------------------

        return value


# ============================================================
# READ TABLEAU EXPORT
# ============================================================


def read_tableau_export(
    export_path: str | Path,
) -> list[dict[str, Any]]:
    """
    Read a Tableau Crosstab exported as CSV or XLSX.

    This function remains available for older/test workflows.
    The current production workflow primarily reads Tableau
    through the Tableau Desktop crosstab service.
    """

    export_path = Path(export_path)

    if not export_path.exists():

        raise FileNotFoundError("Tableau export was not found:\n" f"{export_path}")

    # ========================================================
    # CSV EXPORT
    # ========================================================

    if export_path.suffix.lower() == ".csv":

        with export_path.open(
            "r",
            encoding="utf-8-sig",
            newline="",
        ) as file:

            reader = csv.DictReader(file)

            return [dict(row) for row in reader]

    # ========================================================
    # EXCEL EXPORT
    # ========================================================

    if export_path.suffix.lower() == ".xlsx":

        workbook = load_workbook(
            export_path,
            data_only=True,
            read_only=True,
        )

        try:

            worksheet = workbook.active

            rows = worksheet.iter_rows(values_only=True)

            try:

                headers = [str(value or "").strip() for value in next(rows)]

            except StopIteration:

                return []

            records = []

            for row in rows:

                record = {}

                for index, value in enumerate(row):

                    if index >= len(headers):
                        continue

                    header = headers[index]

                    if not header:
                        continue

                    record[header] = value

                records.append(record)

            return records

        finally:

            workbook.close()

    raise ValueError("Tableau export must be .csv or .xlsx.")


# ============================================================
# FIND EXCEL HEADER ROW
# ============================================================


def find_excel_header_row(
    worksheet,
) -> int:
    """
    Find the Excel forecast table header row.

    Forecast Type is used as the anchor because its COLUMN
    position may vary between workbooks and the header row
    itself may also vary.

    We therefore do NOT assume:

        row 1
        column A
        H2 freeze pane

    The actual workbook structure is discovered dynamically.
    """

    for row_index in range(
        1,
        worksheet.max_row + 1,
    ):

        for column_index in range(
            1,
            worksheet.max_column + 1,
        ):

            value = normalize_text(
                worksheet.cell(
                    row=row_index,
                    column=column_index,
                ).value
            )

            if value == "forecast type":
                return row_index

    raise RuntimeError(
        "Forecast Type header was not found " f'in worksheet "{worksheet.title}".'
    )


# ============================================================
# BUILD EXCEL HEADER LOOKUP
# ============================================================


def build_excel_header_lookup(
    worksheet,
    header_row: int,
) -> dict[str, int]:
    """
    Build a lookup of the EXISTING Excel headers.

    Example:

        {
            "forecast type": 1,
            "bus code name": 2,
            "company": 3,
            "comments": 8,
            "jul_26": 9,
            ...
        }

    IMPORTANT:
    Nothing is renamed here.
    """

    lookup = {}

    for column_index in range(
        1,
        worksheet.max_column + 1,
    ):

        value = worksheet.cell(
            row=header_row,
            column=column_index,
        ).value

        normalized = normalize_text(value)

        if normalized:

            lookup[normalized] = column_index

    return lookup


# ============================================================
# BUILD EXCEL MONTH LOOKUP
# ============================================================


def build_excel_month_lookup(
    worksheet,
    header_row: int,
) -> dict[tuple[int, int], int]:
    """
    Find month columns in Excel using their CURRENT header.

    Example:

        Jul_26 -> (2026, 7)
        Aug_26 -> (2026, 8)

    IMPORTANT:
    This function reports what currently exists.

    If an FY27 workbook incorrectly contains:

        Jan_26

    this function reports:

        (2026, 1)

    The standalone service will correct the header using the
    user-selected Fiscal Year before final month mapping.
    """

    lookup = {}

    for column_index in range(
        1,
        worksheet.max_column + 1,
    ):

        header_value = worksheet.cell(
            row=header_row,
            column=column_index,
        ).value

        month_key = normalize_month_header(header_value)

        if month_key:

            lookup[month_key] = column_index

    return lookup


# ============================================================
# BUILD EXCEL MONTH-NUMBER LOOKUP
# ============================================================


def build_excel_month_number_lookup(
    worksheet,
    header_row: int,
) -> dict[int, int]:
    """
    Locate financial columns by MONTH NUMBER while ignoring
    the year currently shown in Excel.

    This is the safer lookup for correcting fiscal-year headers.

    Example workbook:

        H = Comments
        I = Jul_26
        J = Aug_26
        ...
        O = Jan_26    <- wrong year

    returns approximately:

        {
            7: 9,
            8: 10,
            ...
            1: 15,
        }

    This allows FY27 logic to correct O from Jan_26 to Jan_27.
    """

    lookup = {}

    duplicate_months = []

    for column_index in range(
        1,
        worksheet.max_column + 1,
    ):

        header_value = worksheet.cell(
            row=header_row,
            column=column_index,
        ).value

        month_number = get_month_number_from_header(header_value)

        if month_number is None:
            continue

        # ----------------------------------------------------
        # There should only be one financial column for each
        # month in the forecast table.
        # ----------------------------------------------------

        if month_number in lookup:

            duplicate_months.append(month_number)

            continue

        lookup[month_number] = column_index

    if duplicate_months:

        duplicate_names = [
            MONTH_ABBREVIATIONS[month] for month in sorted(set(duplicate_months))
        ]

        raise RuntimeError(
            "Duplicate fiscal month columns were found in "
            f'worksheet "{worksheet.title}": '
            f"{', '.join(duplicate_names)}"
        )

    return lookup


# ============================================================
# BUILD TABLEAU MONTH LOOKUP
# ============================================================


def build_tableau_month_lookup(
    tableau_rows: list[dict[str, Any]],
) -> dict[tuple[int, int], str]:
    """
    Find Tableau month columns using their full year/month.

    Example:

        {
            (2026, 7): "July 2026",
            (2026, 8): "August 2026",
        }
    """

    if not tableau_rows:
        return {}

    lookup = {}

    for header in tableau_rows[0].keys():

        month_key = normalize_month_header(header)

        if month_key:

            lookup[month_key] = header

    return lookup


# ============================================================
# BUILD TABLEAU MONTH-NUMBER LOOKUP
# ============================================================


def build_tableau_month_number_lookup(
    tableau_rows: list[dict[str, Any]],
) -> dict[int, str]:
    """
    Locate Tableau financial columns by month number.

    The fiscal year is validated separately by the production
    standalone service.

    This helper is useful when Excel contains an incorrect
    year but the month itself is still correct.
    """

    if not tableau_rows:
        return {}

    lookup = {}

    for header in tableau_rows[0].keys():

        month_number = get_month_number_from_header(header)

        if month_number is None:
            continue

        if month_number in lookup:

            raise RuntimeError(
                "Tableau returned more than one column for "
                f"{MONTH_ABBREVIATIONS[month_number]}."
            )

        lookup[month_number] = header

    return lookup


# ============================================================
# BUILD TABLEAU ROW KEY
# ============================================================


def build_tableau_row_key(
    tableau_row: dict[str, Any],
) -> tuple[str, ...]:
    """
    Create a normalized Tableau matching key.

    This helper remains available for services/tests that use
    the complete business-key matching strategy.
    """

    return tuple(
        normalize_text(tableau_row.get(tableau_header))
        for tableau_header in KEY_COLUMN_MAPPING
    )


# ============================================================
# BUILD EXCEL ROW KEY
# ============================================================


def build_excel_row_key(
    worksheet,
    row_index: int,
    header_lookup: dict[str, int],
) -> tuple[str, ...]:
    """
    Create the equivalent normalized Excel matching key.

    Approved header aliases are used so departmental workbooks
    do not have to use identical column names.

    IMPORTANT:
    Existing Excel headers and values are NOT changed.
    """

    key_values = []

    for excel_header in KEY_COLUMN_MAPPING.values():

        column_index = find_excel_column_by_alias(
            header_lookup,
            excel_header,
        )

        if column_index is None:

            raise RuntimeError(
                f'Excel field "{excel_header}" '
                "could not be located using any approved alias."
            )

        value = worksheet.cell(
            row=row_index,
            column=column_index,
        ).value

        key_values.append(normalize_text(value))

    return tuple(key_values)


# ============================================================
# UPDATE EXCEL WORKBOOK
# ============================================================


def update_forecast_from_rows(
    tableau_rows: list[dict],
    forecast_workbook_path: str | Path,
    preferred_sheet_name: str = "Forecast",
) -> dict:
    """
    Legacy/direct forecast-update function.

    This function is retained for compatibility with existing
    tests and any older parts of the project.

    The newer production workflow is handled by:

        tableau_standalone_forecast_service.py

    IMPORTANT BUSINESS RULES
    ------------------------
    This function:

    - does NOT rename Excel descriptive headers
    - does NOT change Forecast Type
    - does NOT change business/account fields
    - does NOT change comments
    - does NOT change Grand Total
    - preserves formulas
    - updates matching Jul-Jun values only

    The production standalone service performs the additional:

    - selected Fiscal Year correction
    - workbook-specific routing
    - formatting standardization
    - auto-sizing
    - dynamic July freeze panes
    """

    forecast_workbook_path = Path(forecast_workbook_path)

    # ========================================================
    # VALIDATE INPUT
    # ========================================================

    if not forecast_workbook_path.exists():

        raise FileNotFoundError(
            "Forecast workbook was not found:\n" f"{forecast_workbook_path}"
        )

    if not tableau_rows:

        raise RuntimeError("No Tableau rows were supplied.")

    # ========================================================
    # OPEN EXCEL FORECAST WORKBOOK
    # ========================================================

    workbook = load_workbook(
        forecast_workbook_path,
        data_only=False,
    )

    try:

        # ----------------------------------------------------
        # Prefer the Forecast worksheet when available.
        # Otherwise use the active worksheet.
        # ----------------------------------------------------

        if preferred_sheet_name in workbook.sheetnames:

            worksheet = workbook[preferred_sheet_name]

        else:

            worksheet = workbook.active

        # ====================================================
        # FIND TABLE STRUCTURE
        # ====================================================

        header_row = find_excel_header_row(worksheet)

        excel_headers = build_excel_header_lookup(
            worksheet,
            header_row,
        )

        excel_months = build_excel_month_lookup(
            worksheet,
            header_row,
        )

        tableau_months = build_tableau_month_lookup(tableau_rows)

        # ====================================================
        # MATCH MONTH COLUMNS
        # ====================================================

        common_months = sorted(set(excel_months) & set(tableau_months))

        if not common_months:

            raise RuntimeError(
                "No matching fiscal month columns were found "
                "between Tableau and Excel."
            )

        # ----------------------------------------------------
        # A complete forecast update requires all 12 months.
        # ----------------------------------------------------

        if len(common_months) != 12:

            raise RuntimeError(
                "Forecast update stopped because exactly "
                "12 matching fiscal month columns were expected. "
                f"Found: {len(common_months)}."
            )

        # ====================================================
        # LOCATE EXCEL DETAIL ROWS
        # ====================================================

        sub_segment_column = find_excel_column_by_alias(
            excel_headers,
            "GL Sub Segment Code",
        )

        description_column = find_excel_column_by_alias(
            excel_headers,
            "GL Sub Account Description",
        )

        if sub_segment_column is None:

            raise RuntimeError(
                "Excel Sub Segment/Account column could not "
                "be located using the approved header aliases."
            )

        if description_column is None:

            raise RuntimeError(
                "Excel Description column could not be "
                "located using the approved header aliases."
            )

        excel_detail_rows = []

        for row_index in range(
            header_row + 1,
            worksheet.max_row + 1,
        ):

            sub_segment_value = worksheet.cell(
                row=row_index,
                column=sub_segment_column,
            ).value

            description_value = worksheet.cell(
                row=row_index,
                column=description_column,
            ).value

            # ------------------------------------------------
            # A financial detail row must contain both an
            # account/sub-segment and a description.
            #
            # Neither field is modified.
            # ------------------------------------------------

            if sub_segment_value in (
                None,
                "",
            ):
                continue

            if description_value in (
                None,
                "",
            ):
                continue

            excel_detail_rows.append(row_index)

        # ====================================================
        # SAFETY CHECK
        # ====================================================

        tableau_row_count = len(tableau_rows)

        excel_row_count = len(excel_detail_rows)

        print()
        print("=" * 70)
        print("FORECAST ROW VALIDATION")
        print("=" * 70)

        print(f"Tableau detail rows: " f"{tableau_row_count}")

        print(f"Excel detail rows:   " f"{excel_row_count}")

        if tableau_row_count != excel_row_count:

            print("Row counts match:    NO")

            print("=" * 70)

            raise RuntimeError(
                "Forecast update stopped because Tableau and "
                "Excel detail-row counts do not match. "
                f"Tableau rows: {tableau_row_count}; "
                f"Excel rows: {excel_row_count}."
            )

        print("Row counts match:    YES")

        print("=" * 70)

        # ====================================================
        # UPDATE JUL-JUN ONLY
        # ====================================================

        cells_updated = 0

        formulas_preserved = 0

        for position, tableau_row in enumerate(tableau_rows):

            excel_row = excel_detail_rows[position]

            for month_key in common_months:

                tableau_column = tableau_months[month_key]

                excel_column = excel_months[month_key]

                tableau_value = tableau_row.get(tableau_column)

                excel_cell = worksheet.cell(
                    row=excel_row,
                    column=excel_column,
                )

                # --------------------------------------------
                # Never overwrite an existing Excel formula.
                # --------------------------------------------

                if isinstance(
                    excel_cell.value,
                    str,
                ) and excel_cell.value.startswith("="):

                    formulas_preserved += 1

                    continue

                # --------------------------------------------
                # Only the financial month value is changed.
                # --------------------------------------------

                excel_cell.value = convert_number(tableau_value)

                cells_updated += 1

        # ====================================================
        # RECALCULATE WHEN EXCEL OPENS
        # ====================================================

        try:

            workbook.calculation.fullCalcOnLoad = True

            workbook.calculation.forceFullCalc = True

            workbook.calculation.calcMode = "auto"

        except Exception:

            # ------------------------------------------------
            # Calculation settings differ slightly between
            # openpyxl versions.
            #
            # Formula cells are still preserved if these
            # properties are unavailable.
            # ------------------------------------------------

            pass

        # ====================================================
        # SAVE
        # ====================================================

        workbook.save(forecast_workbook_path)

        # ====================================================
        # RESULT
        # ====================================================

        print()
        print("=" * 70)
        print("MONTHLY FORECAST VALUES UPDATED")
        print("=" * 70)

        print(f"Rows updated:       " f"{tableau_row_count}")

        print(f"Months updated:     " f"{len(common_months)}")

        print(f"Cells updated:      " f"{cells_updated}")

        print(f"Formulas preserved: " f"{formulas_preserved}")

        print("Modified columns:   Jul-Jun ONLY")

        print("=" * 70)

        return {
            "tableau_rows": tableau_row_count,
            "rows_matched": tableau_row_count,
            "rows_not_found": 0,
            "cells_updated": cells_updated,
            "formulas_preserved": formulas_preserved,
            "duplicate_excel_keys": 0,
            "unmatched_rows": [],
            "months_updated": [f"{year}-{month:02d}" for year, month in common_months],
        }

    finally:

        # ----------------------------------------------------
        # Always close the workbook, even when validation or
        # processing fails.
        # ----------------------------------------------------

        workbook.close()
