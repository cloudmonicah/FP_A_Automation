"""Capture individual Direct Cost Tableau views to OPEX-style Excel crosstabs.

Raw captures are NOT reconciled amounts. Existing financial validation stays unchanged.
"""
from __future__ import annotations
import csv
import re
import time
import xml.etree.ElementTree as ET
from datetime import date
from io import StringIO
from pathlib import Path
from uuid import uuid4
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from automations.actuals.mapping import INDIVIDUAL_DIRECT_COST_MAPPINGS


def _next_month(d):
    return date(d.year + (d.month == 12), d.month % 12 + 1, 1)


def _prepare(source, target, worksheet, closed_period):
    tree = ET.parse(source)
    root = tree.getroot()
    sheet = next((s for s in root.findall('./worksheets/worksheet') if s.get('name') == worksheet), None)
    if sheet is None:
        raise ValueError(f'Worksheet not found: {worksheet}')
    names = {c.get('name') for c in sheet.iter('column')
             if re.sub('[^a-z0-9]', '', c.get('caption', '').lower()) == 'forecaststart'}
    edits = 0
    if len(names) == 1:
        start = _next_month(closed_period)
        for col in root.iter('column'):
            if col.get('name') in names and 'param-domain-type' in col.attrib:
                col.set('value', f'#{start:%Y-%m-%d}#')
                calc = col.find('calculation')
                if calc is not None:
                    calc.set('formula', f'#{start:%Y-%m-%d}#')
                edits += 1
    tree.write(target, encoding='utf-8', xml_declaration=True)
    return edits


def _matrix(text):
    """Preserve Tableau's multi-row period headings (do not DictReader them)."""
    rows = list(csv.reader(StringIO(str(text).lstrip('\ufeff')), delimiter='\t'))
    while rows and not any(str(v).strip() for v in rows[-1]):
        rows.pop()
    if not rows:
        raise ValueError('Empty Tableau crosstab')
    width = max(map(len, rows))
    return [row + [''] * (width - len(row)) for row in rows]


def _sheet_name(index, company, line):
    clean = re.sub(r'[\\/*?:\[\]]', '_', f'{index:03d}_{company}_{line}')
    return clean[:31]


def _write_matrix(sheet, matrix):
    """Match OPEX: original full Tableau crosstab, only row 1 highlighted."""
    from decimal import Decimal, InvalidOperation
    for row in matrix:
        values = []
        for value in row:
            text = str(value).strip()
            if text and not text.endswith('%') and not text.startswith('='):
                try:
                    values.append(float(Decimal(text.replace(',', '').replace('$', ''))))
                    continue
                except (InvalidOperation, ValueError):
                    pass
            values.append(text)
        sheet.append(values)
    for cell in sheet[1]:
        cell.fill = PatternFill('solid', fgColor='703070')
        cell.font = Font(bold=True, color='FFFFFF')
    sheet.freeze_panes = 'E3'
    sheet.sheet_view.showGridLines = False
    for col in range(1, sheet.max_column + 1):
        sheet.column_dimensions[get_column_letter(col)].width = 27 if col <= 4 else 19
    if sheet.max_row >= 2:
        sheet.auto_filter.ref = f'A2:{get_column_letter(sheet.max_column)}{sheet.max_row}'


def run_individual_direct_cost_live(*, output_paths, closed_period, tableau_settings=None, wait_seconds=8):
    """Capture all routes as separate XLSX files and one indexed multi-tab workbook."""
    from automations.actuals.tableau_forecast_service import open_tableau_workbook
    from automations.actuals.tableau_desktop_export_service import get_tableau_crosstab_text
    from automations.actuals.tableau_desktop_close_service import close_tableau_forecast_workbook

    raw_base = Path(output_paths['tableau_raw_copies'])
    folder = raw_base / f'Individual_Direct_Cost_{closed_period:%Y_%m}_{uuid4().hex[:8]}'
    folder.mkdir(parents=True, exist_ok=True)
    combined = Workbook()
    index_sheet = combined.active
    index_sheet.title = 'Capture Index'
    index_sheet.append(['Sheet', 'Company', 'Source TWB', 'Worksheet', 'Rows', 'Raw XLSX'])
    completed = 0
    routes = list(INDIVIDUAL_DIRECT_COST_MAPPINGS)
    for index, route in enumerate(routes, 1):
        company, line = route['company'], route['fs_line']
        source = Path(route['individual_tableau_workbook'])
        worksheet = route['individual_tableau_worksheet']
        name = _sheet_name(index, company, line)
        temp = source.with_name(f'{source.stem}_DC_Live_{uuid4().hex[:8]}.twb')
        opened = False
        status, note, count, raw_path = 'ERROR', '', 0, ''
        print(f'[DIRECT COST {index}/{len(routes)}] {company} / {line} | {worksheet}', flush=True)
        try:
            if not source.is_file():
                raise FileNotFoundError(source)
            edits = _prepare(source, temp, worksheet, closed_period)
            open_tableau_workbook(tableau_workbook_path=str(temp))
            opened = True
            time.sleep(wait_seconds)
            text = get_tableau_crosstab_text(tableau_workbook_path=str(temp), worksheet_name=worksheet)
            matrix = _matrix(text)
            count = max(0, len(matrix) - 2)
            if count == 0:
                raise ValueError('No data rows')
            single = Workbook()
            single_sheet = single.active
            single_sheet.title = name
            _write_matrix(single_sheet, matrix)
            raw = folder / f'{name}.xlsx'
            single.save(raw)
            single.close()
            raw_path = str(raw)
            _write_matrix(combined.create_sheet(name), matrix)
            status = 'CAPTURED'
            note = f'Forecast Start edits: {edits}; raw crosstab only, financial reconciliation NOT VERIFIED'
            completed += 1
        except Exception as exc:
            note = f'{type(exc).__name__}: {exc}'
            print(f'[DIRECT COST] ERROR: {note}', flush=True)
        finally:
            if opened:
                try:
                    close_tableau_forecast_workbook(tableau_workbook_path=str(temp))
                except Exception as exc:
                    note += f'; Tableau close warning: {exc}'
            if temp.exists():
                try:
                    temp.unlink()
                except OSError:
                    note += '; temporary workbook cleanup required'
            index_sheet.append([name, company, str(source), worksheet, count, raw_path])
            combined.save(folder / 'Individual_Direct_Cost_Crosstabs.xlsx')
    combined.close()
    return {'status': 'REVIEW REQUIRED', 'captured': completed, 'total': len(routes),
            'audit_workbook': str(folder / 'Individual_Direct_Cost_Crosstabs.xlsx'),
            'crosstab_workbook': str(folder / 'Individual_Direct_Cost_Crosstabs.xlsx'),
            'raw_folder': str(folder)}
