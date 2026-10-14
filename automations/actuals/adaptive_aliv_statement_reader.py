"""Discover ALIV monthly/YTD actuals by headings and labels (no fixed cells)."""
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
import re
from openpyxl import load_workbook


def norm(v):
    return ' '.join(re.sub(r'^\d+\s+', '', str(v or '')).upper().replace('&',' AND ').split())


def _date(v):
    return v if isinstance(v,(date,datetime)) else None


def read_aliv_period(path, closed_period):
    wb=load_workbook(Path(path),read_only=True,data_only=True)
    try:
        if 'P & L_' not in wb: raise ValueError('ALIV P & L_ worksheet not found')
        ws=wb['P & L_']
        rows=list(ws.iter_rows(values_only=True))
        from automations.actuals.statement_columns import locate_aliv_columns
        header,month_col,ytd_col=locate_aliv_columns(rows,closed_period,Path(path).name)
        out={}
        for r in range(header+1,len(rows)):
            row=rows[r]
            label=norm(row[1] if len(row)>1 else None)
            if not label: continue
            month=row[month_col] if len(row)>month_col else None
            ytd=row[ytd_col] if len(row)>ytd_col else None
            if label in out:
                # duplicate labels are ambiguous; do not silently pick one
                out[label]=None
            elif isinstance(month,(int,float,Decimal)) and (
                isinstance(ytd,(int,float,Decimal)) or closed_period.month == 7
            ):
                # July is month 1 of the fiscal year: YTD is by definition the monthly
                # actual, so a missing/different YTD cell never blocks the run.
                if closed_period.month == 7:
                    ytd = month
                out[label]=(float(month),float(ytd))
        return out
    finally: wb.close()


def get_aliv_line(path,closed_period,label):
    values=read_aliv_period(path,closed_period)
    key=norm(label)
    if key not in values or values[key] is None:
        raise ValueError(f'ALIV statement line missing/ambiguous: {label}')
    return values[key]
