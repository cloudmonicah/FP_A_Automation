from __future__ import annotations

import csv
import getpass
import re
import time

from datetime import (
    date,
    datetime,
)

from decimal import (
    Decimal,
    ROUND_HALF_UP,
)

from pathlib import Path

from typing import (
    Any,
)

from openpyxl import Workbook

from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)

import csv

from io import StringIO

# ============================================================
# PROJECT IMPORTS
# ============================================================

from automations.actuals.validation_config import (
    DIRECT_COST_MAPPINGS,
    DIRECT_COST_TABLEAU_WORKBOOK,
    DIRECT_COST_TABLEAU_WORKSHEET,
    AGG_SUM,
    AGG_TOTAL,
)

from automations.actuals.direct_cost_financial_statement_reader import (
    read_aliv_direct_cost,
)

# ============================================================
# VALIDATION STATUS CONSTANTS
# ============================================================

STATUS_PASS = "PASS"
STATUS_FAIL = "FAIL"
STATUS_ACTION_REQUIRED = "ACTION REQUIRED"
STATUS_REVIEW_REQUIRED = "REVIEW REQUIRED"


# ============================================================
# EXCEL STYLING
# ============================================================

PURPLE_FILL = PatternFill(
    fill_type="solid",
    fgColor="5B1760",
)

LIGHT_PURPLE_FILL = PatternFill(
    fill_type="solid",
    fgColor="EADCF0",
)

PASS_FILL = PatternFill(
    fill_type="solid",
    fgColor="C6EFCE",
)

FAIL_FILL = PatternFill(
    fill_type="solid",
    fgColor="FFC7CE",
)

ACTION_FILL = PatternFill(
    fill_type="solid",
    fgColor="FFEB9C",
)

WHITE_FONT = Font(
    color="FFFFFF",
    bold=True,
)

THIN_SIDE = Side(
    style="thin",
    color="D9D9D9",
)

THIN_BORDER = Border(
    left=THIN_SIDE,
    right=THIN_SIDE,
    top=THIN_SIDE,
    bottom=THIN_SIDE,
)


# ============================================================
# TEXT NORMALIZATION
# ============================================================


def _normalize(
    value: Any,
) -> str:
    """
    Normalize Tableau and mapping labels for comparison.
    """

    if value is None:
        return ""

    text = str(value).strip().upper()

    text = text.replace(
        "&",
        "AND",
    )

    text = re.sub(
        r"\s+",
        " ",
        text,
    )

    return text


# ============================================================
# NUMBER PARSING
# ============================================================


def _to_number(
    value: Any,
) -> float | None:
    """
    Convert Tableau-formatted values into Python numbers.

    Examples:
        "127,943"  -> 127943
        "(500)"    -> -500
        "$48,572"  -> 48572
        ""         -> None
    """

    if value is None:
        return None

    if isinstance(
        value,
        (int, float),
    ):
        return float(value)

    text = str(value).strip()

    if not text:
        return None

    if text.lower() in {
        "nan",
        "none",
        "null",
    }:
        return None

    negative = text.startswith("(") and text.endswith(")")

    text = (
        text.replace("$", "")
        .replace(",", "")
        .replace("%", "")
        .replace("(", "")
        .replace(")", "")
        .strip()
    )

    try:

        number = float(text)

    except ValueError:

        return None

    if negative:
        number *= -1

    return number


# ============================================================
# EXCEL-STYLE ROUNDING
# ============================================================


def _tableau_to_000(
    value: float,
) -> int:
    """
    Tableau stores values in whole dollars.

    Convert to $000 and use Excel-style ROUND HALF UP.
    """

    decimal_value = Decimal(str(value)) / Decimal("1000")

    return int(
        decimal_value.quantize(
            Decimal("1"),
            rounding=ROUND_HALF_UP,
        )
    )


def _fs_to_000(
    value: float,
) -> int:
    """
    Financial Statement values are already stored in $000.

    Only round to the nearest whole $000.
    """

    decimal_value = Decimal(str(value))

    return int(
        decimal_value.quantize(
            Decimal("1"),
            rounding=ROUND_HALF_UP,
        )
    )


# ============================================================
# FORECAST START
# ============================================================


def _get_forecast_start_from_closed_period(
    closed_period: date | datetime,
) -> date:
    """
    Forecast Start is always the first day of the month
    immediately following the closed actuals period.

    Example:
        August 2026 actuals
        -> Forecast Start = 09/01/2026
    """

    if closed_period.month == 12:

        return date(
            closed_period.year + 1,
            1,
            1,
        )

    return date(
        closed_period.year,
        closed_period.month + 1,
        1,
    )


# ============================================================
# RESOLVE CLOSED-MONTH ACTUAL COLUMN
# ============================================================


def _find_month_actual_column(
    row: dict[str, Any],
    closed_period: date | datetime,
) -> str:
    """
    Find the Tableau Actual column for the selected month.

    Expected examples:
        July 2026 Actual
        August 2026 Actual
        August 2026\nActual
        August 2026 - Actual

    The function intentionally rejects:
        Budget
        Var $
        Var %
        Grand Total
    """

    month_name = closed_period.strftime("%B %Y").upper()

    candidates = []

    for key in row.keys():

        normalized_key = _normalize(key)

        if month_name not in normalized_key:
            continue

        if "ACTUAL" not in normalized_key:
            continue

        if "BUDGET" in normalized_key:
            continue

        if "VAR" in normalized_key:
            continue

        if "GRAND TOTAL" in normalized_key:
            continue

        candidates.append(key)

    if len(candidates) == 1:

        return candidates[0]

    if not candidates:

        available_columns = "\n".join(f" - {key}" for key in row.keys())

        raise KeyError(
            "Could not find the Tableau Actual column for "
            f"{closed_period:%B %Y}.\n\n"
            "Available columns:\n"
            f"{available_columns}"
        )

    raise RuntimeError(
        "More than one possible Tableau Actual column was "
        f"found for {closed_period:%B %Y}:\n"
        + "\n".join(f" - {column}" for column in candidates)
    )


# ============================================================
# RESOLVE DIMENSION COLUMN
# ============================================================


def _find_dimension_column(
    row: dict[str, Any],
    expected_name: str,
) -> str:
    """
    Find a Tableau dimension column using normalized names.
    """

    expected = _normalize(expected_name)

    for key in row.keys():

        if _normalize(key) == expected:

            return key

    raise KeyError(f"Required Tableau column " f"'{expected_name}' was not found.")


# ============================================================
# BUILD TABLEAU DIRECT COST VALUES
# ============================================================


def build_direct_cost_tableau_values(
    rows: list[dict[str, Any]],
    *,
    closed_period: date | datetime,
) -> dict[str, dict[str, Any]]:
    """
    Aggregate Direct Cost detail rows by:

        Company
        GL Account Mapping

    The worksheet contains GL-level detail. Repeated values in the
    Company and GL Account Mapping columns may be blank in copied
    Tableau crosstabs, so the parser forward-fills those dimensions.

    Example:

        GL Account Mapping = ROAMING COSTS

        127,943
        221,559
          1,000
         21,059
        -------
        371,561

    This function returns:

        {
            "ALIV": {
                "categories": {...},
                "total": ...
            },

            "CBL": {
                "categories": {...},
                "total": ...
            }
        }

    Company total is calculated from ALL detail rows rather than
    relying on Tableau's displayed subtotal row. This prevents
    subtotal formatting from affecting the comparison.
    """

    if not rows:

        raise RuntimeError("No Tableau Direct Cost rows were returned.")

    first_row = rows[0]

    company_column = _find_dimension_column(
        first_row,
        "Company",
    )

    mapping_column = _find_dimension_column(
        first_row,
        "GL Account Mapping",
    )

    actual_column = _find_month_actual_column(
        first_row,
        closed_period,
    )

    print()
    print("=" * 70)
    print("DIRECT COST TABLEAU PARSER")
    print("=" * 70)
    print(
        "Closed Period :",
        closed_period.strftime("%B %Y"),
    )
    print(
        "Actual Column :",
        actual_column,
    )
    print(
        "Company Col   :",
        company_column,
    )
    print(
        "Mapping Col   :",
        mapping_column,
    )
    print("=" * 70)

    results = {
        "ALIV": {
            "categories": {},
            "total": 0.0,
        },
        "CBL": {
            "categories": {},
            "total": 0.0,
        },
    }

    current_company = ""
    current_mapping = ""

    for row in rows:

        raw_company = row.get(company_column)

        raw_mapping = row.get(mapping_column)

        # ----------------------------------------------------
        # FORWARD-FILL COMPANY
        # ----------------------------------------------------
        #
        # Tableau often displays Company only once at the top
        # of each grouped block.
        # ----------------------------------------------------

        normalized_company = _normalize(raw_company)

        if normalized_company in {
            "ALIV",
            "CBL",
        }:

            current_company = normalized_company

        # ----------------------------------------------------
        # STOP / IGNORE NON-COMPANY ROWS
        # ----------------------------------------------------

        if current_company not in results:
            continue

        # ----------------------------------------------------
        # FORWARD-FILL GL ACCOUNT MAPPING
        # ----------------------------------------------------

        normalized_mapping = _normalize(raw_mapping)

        if normalized_mapping and normalized_mapping not in {
            "TOTAL",
            "GRAND TOTAL",
        }:

            current_mapping = normalized_mapping

        # ----------------------------------------------------
        # IGNORE TABLEAU SUBTOTAL / GRAND TOTAL ROWS
        # ----------------------------------------------------
        #
        # We calculate the company total ourselves from the
        # underlying detail rows.
        # ----------------------------------------------------

        row_dimension_text = " | ".join(
            _normalize(value)
            for key, value in row.items()
            if key != actual_column and value is not None
        )

        if "GRAND TOTAL" in row_dimension_text:

            continue

        # A row containing only the Tableau subtotal should not
        # be counted again.
        if normalized_mapping in {
            "TOTAL",
            "GRAND TOTAL",
        }:

            continue

        # ----------------------------------------------------
        # ACTUAL VALUE
        # ----------------------------------------------------

        actual_value = _to_number(row.get(actual_column))

        if actual_value is None:
            continue

        # ----------------------------------------------------
        # REQUIRE A CATEGORY
        # ----------------------------------------------------

        if not current_mapping:
            continue

        company_result = results[current_company]

        categories = company_result["categories"]

        # ----------------------------------------------------
        # AGGREGATE EVERY DETAIL ROW UNDER THE SAME
        # GL ACCOUNT MAPPING.
        # ----------------------------------------------------

        categories[current_mapping] = (
            categories.get(
                current_mapping,
                0.0,
            )
            + actual_value
        )

        # ----------------------------------------------------
        # COMPANY TOTAL
        # ----------------------------------------------------

        company_result["total"] += actual_value

    # ========================================================
    # PRINT AUDIT TOTALS
    # ========================================================

    for company in (
        "ALIV",
        "CBL",
    ):

        print()
        print(f"{company} DIRECT COST TABLEAU VALUES")
        print("-" * 70)

        for (
            category,
            amount,
        ) in sorted(results[company]["categories"].items()):

            print(f"{category:<35}" f"{amount:>18,.2f}")

        print("-" * 70)
        print(
            f"{company + ' CALCULATED TOTAL':<35}"
            f"{results[company]['total']:>18,.2f}"
        )

    print("=" * 70)

    return results


# ============================================================
# FINANCIAL STATEMENT VALUE LOOKUP
# ============================================================


def _resolve_fs_value(
    *,
    fs_line: str,
    fs_values: dict[str, float],
) -> float | None:
    """
    Resolve a Financial Statement value using normalized labels.
    """

    target = _normalize(fs_line)

    for (
        extracted_line,
        value,
    ) in fs_values.items():

        if _normalize(extracted_line) == target:

            return value

    return None


# ============================================================
# COMPARE ONE COMPANY
# ============================================================


def compare_direct_cost_company(
    *,
    company: str,
    fs_values: dict[str, float],
    tableau_values: dict[str, Any],
) -> dict[str, Any]:
    """
    Compare Direct Cost Financial Statement lines against
    aggregated Tableau detail.

    TOTAL mappings use the total of all Tableau detail rows.

    SUM mappings add all mapped GL Account Mapping categories.
    """

    company = _normalize(company)

    details: list[dict[str, Any]] = []

    for mapping in DIRECT_COST_MAPPINGS:

        if _normalize(mapping.get("company")) != company:

            continue

        fs_line = mapping["fs_line"]

        aggregation = mapping["aggregation"]

        mapped_categories = [
            _normalize(category)
            for category in mapping.get(
                "tableau_categories",
                [],
            )
        ]

        # ====================================================
        # FINANCIAL STATEMENT
        # ====================================================

        fs_raw = _resolve_fs_value(
            fs_line=(fs_line),
            fs_values=(fs_values),
        )

        fs_000 = _fs_to_000(fs_raw) if fs_raw is not None else None

        result = {
            "company": company,
            "statement_line": fs_line,
            "tableau_categories": (mapped_categories),
            "aggregation": aggregation,
            "fs_raw_000": fs_raw,
            "fs_rounded_000": fs_000,
            "tableau_raw_dollars": None,
            "tableau_rounded_000": None,
            "variance_000": None,
            "status": None,
            "note": "",
        }

        # ====================================================
        # FINANCIAL STATEMENT CONTROL
        # ====================================================

        if fs_raw is None:

            result["status"] = STATUS_ACTION_REQUIRED

            result["note"] = "Financial Statement line was not found."

            details.append(result)

            continue

        # ====================================================
        # TABLEAU TOTAL
        # ====================================================

        if aggregation == AGG_TOTAL:

            tableau_raw = tableau_values.get("total")

            if tableau_raw is None:

                result["status"] = STATUS_ACTION_REQUIRED

                result["note"] = "Tableau company total could not be calculated."

                details.append(result)

                continue

        # ====================================================
        # TABLEAU SUM MAPPING
        # ====================================================

        elif aggregation == AGG_SUM:

            category_values = tableau_values.get("categories", {})

            missing_categories = []

            tableau_raw = 0.0

            for category in mapped_categories:

                if category not in category_values:

                    missing_categories.append(category)

                    continue

                tableau_raw += category_values[category]

            if missing_categories:

                result["status"] = STATUS_ACTION_REQUIRED

                result["tableau_raw_dollars"] = (
                    tableau_raw if tableau_raw != 0 else None
                )

                result["note"] = "Missing Tableau category: " + ", ".join(
                    missing_categories
                )

                details.append(result)

                continue

        else:

            result["status"] = STATUS_ACTION_REQUIRED

            result["note"] = (
                "Unsupported Direct Cost aggregation type: " f"{aggregation}"
            )

            details.append(result)

            continue

        # ====================================================
        # COMPARE
        # ====================================================

        tableau_000 = _tableau_to_000(tableau_raw)

        variance_000 = tableau_000 - fs_000

        status = STATUS_PASS if variance_000 == 0 else STATUS_FAIL

        if variance_000 == 0:

            note = "Reconciled."

        elif variance_000 > 0:

            note = (
                f"Tableau is {variance_000:,} ($000) "
                "higher than the Financial Statement."
            )

        else:

            note = (
                f"Tableau is {abs(variance_000):,} ($000) "
                "lower than the Financial Statement."
            )

        result.update(
            {
                "tableau_raw_dollars": (tableau_raw),
                "tableau_rounded_000": (tableau_000),
                "variance_000": (variance_000),
                "status": (status),
                "note": (note),
            }
        )

        details.append(result)

    # ========================================================
    # COMPANY SUMMARY
    # ========================================================

    passed = sum(1 for item in details if item["status"] == STATUS_PASS)

    failed = sum(1 for item in details if item["status"] == STATUS_FAIL)

    action_required = sum(
        1 for item in details if item["status"] == STATUS_ACTION_REQUIRED
    )

    company_status = (
        STATUS_PASS
        if (failed == 0 and action_required == 0)
        else STATUS_REVIEW_REQUIRED
    )

    return {
        "company": company,
        "status": company_status,
        "summary": {
            "passed": passed,
            "failed": failed,
            "action_required": (action_required),
        },
        "details": details,
    }


# ============================================================
# RAW TABLEAU CSV
# ============================================================


def _save_tableau_rows_csv(
    *,
    rows: list[dict],
    output_file: str | Path,
) -> Path:
    """
    Save the exact parsed Tableau rows for audit evidence.
    """

    output_file = Path(output_file)

    output_file.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    if not rows:

        raise RuntimeError("No Direct Cost Tableau rows were available.")

    headers = list(rows[0].keys())

    with output_file.open(
        "w",
        encoding="utf-8-sig",
        newline="",
    ) as handle:

        writer = csv.DictWriter(
            handle,
            fieldnames=(headers),
        )

        writer.writeheader()

        writer.writerows(rows)

    return output_file


# ============================================================
# DIRECT COST TABLEAU CROSSTAB PARSER
# ============================================================


def _parse_direct_cost_crosstab_text(
    clipboard_text: str,
) -> list[dict[str, Any]]:
    """
    Parse the Direct Cost Tableau crosstab.

    The Tableau export contains two header rows:

        Row 1:
            Company
            GL Business Code
            GL Account Mapping
            GL Sub Segment code
            GL Sub Account Description
            July 2026
            ...
            August 2026
            ...
            Grand Total
            ...

        Row 2:
            blank
            blank
            blank
            blank
            blank
            Actual
            Budget
            Var $
            Var %
            Actual
            Budget
            Var $
            Var %
            ...

    We combine both header rows into flat names such as:

        Company
        GL Business Code
        GL Account Mapping
        July 2026 Actual
        July 2026 Budget
        August 2026 Actual
        August 2026 Budget
        Grand Total Actual
    """

    if not clipboard_text.strip():

        raise RuntimeError("Direct Cost Tableau copied an empty crosstab.")

    # --------------------------------------------------------
    # Convert clipboard text into a raw matrix.
    # --------------------------------------------------------

    reader = csv.reader(
        StringIO(clipboard_text),
        delimiter="\t",
    )

    raw_rows = [row for row in reader if any(str(value).strip() for value in row)]

    if len(raw_rows) < 3:

        raise RuntimeError(
            "Direct Cost Tableau crosstab does not contain "
            "enough rows to identify its headers."
        )

    # --------------------------------------------------------
    # Print the first few raw rows so we can verify exactly
    # what Tableau placed on the clipboard.
    # --------------------------------------------------------

    print()
    print("=" * 70)
    print("DIRECT COST RAW TABLEAU HEADER DIAGNOSTIC")
    print("=" * 70)

    for index, row in enumerate(
        raw_rows[:4],
        start=1,
    ):

        print(f"RAW ROW {index}:")

        for column_number, value in enumerate(
            row,
            start=1,
        ):

            print(f"  {column_number:02}: " f"{value!r}")

    print("=" * 70)

    # --------------------------------------------------------
    # Direct Cost uses:
    #
    # Row 1 = month / dimension headers
    # Row 2 = measure headers
    #
    # We will combine them.
    # --------------------------------------------------------

    header_row_1 = raw_rows[0]

    header_row_2 = raw_rows[1]

    column_count = max(
        len(header_row_1),
        len(header_row_2),
    )

    # --------------------------------------------------------
    # Extend both header rows so indexes always line up.
    # --------------------------------------------------------

    header_row_1 = header_row_1 + [""] * (column_count - len(header_row_1))

    header_row_2 = header_row_2 + [""] * (column_count - len(header_row_2))

    # --------------------------------------------------------
    # Tableau may show a month name only on the first column
    # within a month group.
    #
    # Example:
    #
    # July 2026 | blank | blank | blank
    # Actual    | Budget| Var $ | Var %
    #
    # Forward-fill the month heading across those measure
    # columns.
    # --------------------------------------------------------

    combined_headers = []

    current_parent_header = ""

    dimension_headers = {
        "COMPANY",
        "GL BUSINESS CODE",
        "GL FS CATEGORY",
        "GL ACCOUNT MAPPING",
        "GL SUB SEGMENT CODE",
        "GL SUB ACCOUNT DESCRIPTION",
    }

    for column_index in range(column_count):

        top_header = str(header_row_1[column_index]).strip()

        lower_header = str(header_row_2[column_index]).strip()

        normalized_top = _normalize(top_header)

        # ----------------------------------------------------
        # Dimension columns stand by themselves.
        # ----------------------------------------------------

        if normalized_top in dimension_headers:

            combined_headers.append(top_header)

            current_parent_header = ""

            continue

        # ----------------------------------------------------
        # If Tableau supplied a parent/month heading, remember
        # it for the following Actual/Budget/Variance columns.
        # ----------------------------------------------------

        if top_header:

            current_parent_header = top_header

        # ----------------------------------------------------
        # Build final flattened header.
        # ----------------------------------------------------

        if current_parent_header and lower_header:

            combined_header = f"{current_parent_header} " f"{lower_header}"

        elif top_header:

            combined_header = top_header

        elif lower_header:

            combined_header = lower_header

        else:

            combined_header = f"Unnamed Column " f"{column_index + 1}"

        combined_headers.append(combined_header)

    # --------------------------------------------------------
    # Verify the resulting headers.
    # --------------------------------------------------------

    print()
    print("=" * 70)
    print("DIRECT COST FLATTENED TABLEAU COLUMNS")
    print("=" * 70)

    for header in combined_headers:

        print(
            " -",
            header,
        )

    print("=" * 70)

    # --------------------------------------------------------
    # Build dictionaries from the remaining data rows.
    # --------------------------------------------------------

    results = []

    for raw_row in raw_rows[2:]:

        raw_row = raw_row + [""] * (column_count - len(raw_row))

        row_dict = {}

        for (
            column_index,
            header,
        ) in enumerate(combined_headers):

            row_dict[header] = raw_row[column_index]

        results.append(row_dict)

    if not results:

        raise RuntimeError(
            "No Direct Cost Tableau data rows were found "
            "after parsing the multi-row header."
        )

    print(
        "Direct Cost parsed data rows:",
        len(results),
    )

    return results


# ============================================================
# GENERAL SHEET STYLE
# ============================================================


def _style_sheet(
    worksheet,
) -> None:
    """
    Apply basic readable formatting.
    """

    worksheet.freeze_panes = "A2"

    for cell in worksheet[1]:

        cell.fill = PURPLE_FILL

        cell.font = WHITE_FONT

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    for column_cells in worksheet.columns:

        max_length = 0

        column_letter = column_cells[0].column_letter

        for cell in column_cells:

            text = "" if cell.value is None else str(cell.value)

            max_length = max(
                max_length,
                len(text),
            )

        worksheet.column_dimensions[column_letter].width = min(
            max(
                max_length + 2,
                12,
            ),
            45,
        )


# ============================================================
# RAW TABLEAU EXCEL
# ============================================================


def _save_tableau_rows_workbook(
    *,
    rows: list[dict],
    output_file: str | Path,
) -> Path:
    """
    Save the same Direct Cost Tableau extract to Excel.
    """

    output_file = Path(output_file)

    output_file.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    if not rows:

        raise RuntimeError("No Direct Cost Tableau rows were available.")

    workbook = Workbook()

    worksheet = workbook.active

    worksheet.title = "Direct Cost Tableau Extract"

    headers = list(rows[0].keys())

    worksheet.append(headers)

    for row in rows:

        worksheet.append([row.get(header) for header in headers])

    _style_sheet(worksheet)

    workbook.save(output_file)

    workbook.close()

    return output_file


# ============================================================
# DIRECT COST DETAIL VALIDATION SHEET
# ============================================================


def _write_company_sheet(
    workbook: Workbook,
    *,
    company_result: dict[str, Any],
    closed_period: date | datetime,
    forecast_start: date,
) -> None:

    company = company_result["company"]

    worksheet = workbook.create_sheet(title=f"{company} Direct Cost")

    # ========================================================
    # TITLE
    # ========================================================

    worksheet.merge_cells("A1:F1")

    worksheet["A1"] = f"{company} DIRECT COST ACTUALS VALIDATION"

    worksheet["A1"].fill = PURPLE_FILL

    worksheet["A1"].font = Font(
        color="FFFFFF",
        bold=True,
        size=14,
    )

    worksheet["A1"].alignment = Alignment(
        horizontal="center",
        vertical="center",
    )

    # ========================================================
    # PERIOD INFORMATION
    # ========================================================

    worksheet["A2"] = "Closed Period"

    worksheet["B2"] = closed_period.strftime("%B %Y")

    worksheet["A3"] = "Tableau Forecast Start"

    worksheet["B3"] = forecast_start.strftime("%m/%d/%Y")

    # ========================================================
    # COLUMN HEADERS
    # ========================================================

    headers = [
        "Financial Statement Line",
        "Financial Statement ($000)",
        "Tableau ($000)",
        "Variance ($000)",
        "Status",
        "Exception / Note",
    ]

    header_row = 5

    for (
        column_number,
        header,
    ) in enumerate(
        headers,
        start=1,
    ):

        cell = worksheet.cell(
            row=header_row,
            column=column_number,
            value=header,
        )

        cell.fill = PURPLE_FILL
        cell.font = WHITE_FONT
        cell.border = THIN_BORDER

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    # ========================================================
    # VALIDATION DETAIL
    # ========================================================

    current_row = header_row + 1

    for item in company_result["details"]:

        values = [
            item["statement_line"],
            item["fs_rounded_000"],
            item["tableau_rounded_000"],
            item["variance_000"],
            item["status"],
            item["note"],
        ]

        for (
            column_number,
            value,
        ) in enumerate(
            values,
            start=1,
        ):

            cell = worksheet.cell(
                row=current_row,
                column=column_number,
                value=value,
            )

            cell.border = THIN_BORDER

            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )

        # ====================================================
        # STATUS FORMATTING
        # ====================================================

        status_cell = worksheet.cell(
            row=current_row,
            column=5,
        )

        if item["status"] == STATUS_PASS:

            status_cell.fill = PASS_FILL

        elif item["status"] == STATUS_FAIL:

            status_cell.fill = FAIL_FILL

        else:

            status_cell.fill = ACTION_FILL

        # ====================================================
        # NUMBER FORMATTING
        # ====================================================

        worksheet.cell(
            row=current_row,
            column=2,
        ).number_format = "#,##0"

        worksheet.cell(
            row=current_row,
            column=3,
        ).number_format = "#,##0"

        worksheet.cell(
            row=current_row,
            column=4,
        ).number_format = "+#,##0;-#,##0;0"

        current_row += 1

    # ========================================================
    # COLUMN WIDTHS
    # ========================================================
    #
    # Match the simplified OPEX presentation.
    # ========================================================

    widths = {
        "A": 34,
        "B": 22,
        "C": 20,
        "D": 20,
        "E": 20,
        "F": 50,
    }

    for (
        column,
        width,
    ) in widths.items():

        worksheet.column_dimensions[column].width = width

    # ========================================================
    # FREEZE HEADER
    # ========================================================

    worksheet.freeze_panes = "A6"


# ============================================================
# EXECUTIVE SUMMARY
# ============================================================


def _write_executive_summary(
    workbook: Workbook,
    *,
    closed_period: date | datetime,
    forecast_start: date,
    aliv_result: dict[str, Any],
    cbl_result: dict[str, Any],
) -> None:
    """
    Write management-facing summary.
    """

    worksheet = workbook.active

    worksheet.title = "Executive Summary"

    worksheet.merge_cells("A1:H1")

    worksheet["A1"] = "DIRECT COST ACTUALS VALIDATION"

    worksheet["A1"].fill = PURPLE_FILL

    worksheet["A1"].font = Font(
        color="FFFFFF",
        bold=True,
        size=16,
    )

    worksheet["A1"].alignment = Alignment(
        horizontal="center",
    )

    worksheet.merge_cells("A2:H2")

    worksheet["A2"] = (
        f"Closed Period: {closed_period:%B %Y}"
        "   |   "
        f"Tableau Forecast Start: "
        f"{forecast_start:%m/%d/%Y}"
    )

    worksheet["A2"].fill = LIGHT_PURPLE_FILL

    worksheet["A2"].alignment = Alignment(
        horizontal="center",
    )

    headers = [
        "Company",
        "Status",
        "Passed",
        "Failed",
        "Action Required",
    ]

    for (
        index,
        header,
    ) in enumerate(
        headers,
        start=1,
    ):

        cell = worksheet.cell(
            row=5,
            column=index,
            value=header,
        )

        cell.fill = PURPLE_FILL

        cell.font = WHITE_FONT

        cell.border = THIN_BORDER

    row_number = 6

    for result in (
        aliv_result,
        cbl_result,
    ):

        values = [
            result["company"],
            result["status"],
            result["summary"]["passed"],
            result["summary"]["failed"],
            result["summary"]["action_required"],
        ]

        for (
            column_number,
            value,
        ) in enumerate(
            values,
            start=1,
        ):

            worksheet.cell(
                row=row_number,
                column=(column_number),
                value=value,
            ).border = THIN_BORDER

        status_cell = worksheet.cell(
            row=row_number,
            column=2,
        )

        if result["status"] == STATUS_PASS:

            status_cell.fill = PASS_FILL

        else:

            status_cell.fill = ACTION_FILL

        row_number += 1

    worksheet["A10"] = "Exceptions Requiring Review"

    worksheet["A10"].font = Font(
        bold=True,
        size=12,
    )

    exception_headers = [
        "Company",
        "Financial Statement Line",
        "FS ($000)",
        "Tableau ($000)",
        "Variance ($000)",
        "Status",
        "Exception / Note",
        "Owner / Action",
    ]

    for (
        column_number,
        header,
    ) in enumerate(
        exception_headers,
        start=1,
    ):

        cell = worksheet.cell(
            row=11,
            column=(column_number),
            value=header,
        )

        cell.fill = PURPLE_FILL

        cell.font = WHITE_FONT

        cell.border = THIN_BORDER

    exception_row = 12

    for result in (
        aliv_result,
        cbl_result,
    ):

        for item in result["details"]:

            if item["status"] == STATUS_PASS:

                continue

            values = [
                item["company"],
                item["statement_line"],
                item["fs_rounded_000"],
                item["tableau_rounded_000"],
                item["variance_000"],
                item["status"],
                item["note"],
                "",
            ]

            for (
                column_number,
                value,
            ) in enumerate(
                values,
                start=1,
            ):

                cell = worksheet.cell(
                    row=exception_row,
                    column=(column_number),
                    value=value,
                )

                cell.border = THIN_BORDER

                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True,
                )

            exception_row += 1

    if exception_row == 12:

        worksheet.merge_cells("A12:H13")

        worksheet["A12"] = "No validation exceptions identified."

        worksheet["A12"].fill = PASS_FILL

        worksheet["A12"].alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    widths = {
        "A": 14,
        "B": 34,
        "C": 16,
        "D": 18,
        "E": 18,
        "F": 20,
        "G": 48,
        "H": 24,
    }

    for (
        column,
        width,
    ) in widths.items():

        worksheet.column_dimensions[column].width = width


# ============================================================
# WRITE VALIDATION WORKBOOK
# ============================================================


def write_direct_cost_validation_workbook(
    *,
    output_file: str | Path,
    closed_period: date | datetime,
    forecast_start: date,
    aliv_result: dict[str, Any],
    cbl_result: dict[str, Any],
) -> Path:
    """
    Create the completed Direct Cost validation workbook.
    """

    from automations.actuals.actuals_report_layout import write_area_workbook

    return write_area_workbook(
        area="Direct Cost",
        output_file=output_file,
        closed_period=closed_period,
        forecast_start=forecast_start,
        aliv_result=aliv_result,
        cbl_result=cbl_result,
    )


# ============================================================
# LIVE DIRECT COST VALIDATION
# ============================================================


def run_live_direct_cost_validation(
    *,
    aliv_fs_path: str | Path,
    cbl_fs_path: str | Path,
    output_paths: dict[str, Path],
    closed_period: date | datetime,
    tableau_workbook_path: str | Path,
    open_tableau_first: bool = False,
    tableau_open_wait_seconds: int = 8,
) -> dict[str, Any]:

    # ========================================================
    # RUN AUDIT INFORMATION
    # ========================================================

    run_timestamp = datetime.now()

    run_by = getpass.getuser()

    # ========================================================
    # 1. READ FINANCIAL STATEMENTS
    # ========================================================

    fs_results = {"ALIV": read_aliv_direct_cost(aliv_fs_path, closed_period)}

    # ========================================================
    # 2. FORECAST START
    # ========================================================

    from automations.actuals.cbl_statement_periods import read_cbl_area

    fs_results["CBL"] = read_cbl_area(cbl_fs_path, "Direct Cost", closed_period)[
        "MONTH"
    ]

    forecast_start = _get_forecast_start_from_closed_period(closed_period)

    print()
    print("=" * 70)
    print("DIRECT COST TABLEAU ACTUALS PERIOD")
    print("=" * 70)

    print(
        "Actuals Month  :",
        closed_period.strftime("%B %Y"),
    )

    print(
        "Forecast Start :",
        forecast_start.strftime("%m/%d/%Y"),
    )

    print("=" * 70)

    # ========================================================
    # 3. UPDATE FORECAST START IN TWB
    # ========================================================

    from uuid import uuid4
    from automations.actuals.actuals_report_layout import prepare_group_copy

    source_tableau_path = Path(tableau_workbook_path)
    from shared.tableau_temp_copies import temp_twb_path

    tableau_workbook_path = temp_twb_path(source_tableau_path, "GroupActuals")
    prepare_group_copy(
        source_tableau_path,
        tableau_workbook_path,
        DIRECT_COST_TABLEAU_WORKSHEET,
        closed_period,
    )

    # ========================================================
    # 4. OPEN TABLEAU
    # ========================================================

    if open_tableau_first:

        from automations.forecast.tableau_service import (
            open_tableau_workbook,
        )

        open_tableau_workbook(tableau_workbook_path=(tableau_workbook_path))

        print("Mapped Direct Cost Tableau workbook opened:")

        print(tableau_workbook_path)

        print("Waiting for Tableau Desktop...")

        time.sleep(tableau_open_wait_seconds)

    # ========================================================
    # 5. EXTRACT TABLEAU CROSSTAB
    # ========================================================

    from automations.actuals.tableau_desktop_export_service import (
        get_tableau_crosstab_text,
    )

    # ========================================================
    # COPY THE RAW DIRECT COST CROSSTAB
    # ========================================================
    #
    # Direct Cost has hierarchical column headers, so we do not
    # use the standard flat Tableau parser used by OPEX.
    # ========================================================

    from automations.actuals.tableau_desktop_close_service import close_tableau_forecast_workbook

    try:
        tableau_clipboard_text = get_tableau_crosstab_text(
            tableau_workbook_path=str(tableau_workbook_path),
            worksheet_name=DIRECT_COST_TABLEAU_WORKSHEET,
        )
    finally:
        close_tableau_forecast_workbook(
            tableau_workbook_path=str(tableau_workbook_path)
        )
        Path(tableau_workbook_path).unlink(missing_ok=True)

    # ========================================================
    # PARSE DIRECT COST MULTI-ROW HEADERS
    # ========================================================

    tableau_rows = _parse_direct_cost_crosstab_text(tableau_clipboard_text)

    if not tableau_rows:

        raise RuntimeError("Direct Cost Tableau extraction returned no rows.")

    # ========================================================
    # OUTPUT PATHS
    # ========================================================

    period_tag = closed_period.strftime("%Y-%m")

    raw_folder = Path(output_paths["tableau_raw_copies"])

    direct_cost_folder = Path(output_paths["direct_cost"])

    logs_folder = Path(output_paths["logs"])

    # ========================================================
    # 6. SAVE RAW TABLEAU
    # ========================================================

    raw_csv = raw_folder / ("Direct_Cost_Tableau_Raw_" f"{period_tag}.csv")

    _save_tableau_rows_csv(
        rows=(tableau_rows),
        output_file=(raw_csv),
    )

    raw_excel = raw_folder / ("Direct_Cost_Tableau_Extract_" f"{period_tag}.xlsx")

    _save_tableau_rows_workbook(
        rows=(tableau_rows),
        output_file=(raw_excel),
    )

    # ========================================================
    # 7. BUILD TABLEAU CATEGORY TOTALS
    # ========================================================

    tableau_values = build_direct_cost_tableau_values(
        tableau_rows,
        closed_period=(closed_period),
    )

    from automations.actuals.opex_validation_service import _validate_group_period_headers

    _validate_group_period_headers(tableau_rows, closed_period)
    ytd_rows = []
    month_header = closed_period.strftime("%B %Y") + " Actual"
    for row in tableau_rows:
        keys = [k for k in row if _normalize(k) == "GRAND TOTAL ACTUAL"]
        if len(keys) != 1:
            raise ValueError(
                "Expected one Grand Total Actual column for Direct Cost YTD."
            )
        ytd_rows.append({**row, month_header: row[keys[0]]})
    tableau_ytd_values = build_direct_cost_tableau_values(
        ytd_rows, closed_period=closed_period
    )

    # ========================================================
    # 8. COMPARE ALIV
    # ========================================================

    aliv_result = compare_direct_cost_company(
        company="ALIV",
        fs_values=(fs_results["ALIV"]),
        tableau_values=(tableau_values["ALIV"]),
    )

    # ========================================================
    # 9. COMPARE CBL
    # ========================================================

    cbl_result = compare_direct_cost_company(
        company="CBL",
        fs_values=(fs_results["CBL"]),
        tableau_values=(tableau_values["CBL"]),
    )

    # ========================================================
    # 10. OVERALL STATUS
    # ========================================================

    overall_status = (
        STATUS_PASS
        if (
            aliv_result["status"] == STATUS_PASS and cbl_result["status"] == STATUS_PASS
        )
        else STATUS_REVIEW_REQUIRED
    )

    # ========================================================
    # 11. WRITE VALIDATION WORKBOOK
    # ========================================================

    validation_file = direct_cost_folder / (
        "Direct_Cost_Actuals_Validation_" f"{period_tag}.xlsx"
    )

    write_direct_cost_validation_workbook(
        output_file=(validation_file),
        closed_period=(closed_period),
        forecast_start=(forecast_start),
        aliv_result=(aliv_result),
        cbl_result=(cbl_result),
    )

    # ========================================================
    # 12. BUILD RESULT OBJECT
    # ========================================================

    from automations.actuals.opex_validation_service import _write_group_actuals_sheet

    monthly = {
        c: {**v["categories"], f"{c} TOTAL": v["total"]}
        for c, v in tableau_values.items()
    }
    ytd = {
        c: {**v["categories"], f"{c} TOTAL": v["total"]}
        for c, v in tableau_ytd_values.items()
    }
    _write_group_actuals_sheet(
        validation_file,
        monthly,
        ytd,
        closed_period,
        area="Direct Cost",
        aliv_fs_path=aliv_fs_path,
        cbl_fs_path=cbl_fs_path,
    )

    from automations.actuals.actuals_report_layout import report_status

    overall_status = report_status(validation_file)

    result = {
        "area": ("Direct Cost"),
        "closed_period": (closed_period.strftime("%Y-%m-%d")),
        "forecast_start": (forecast_start.strftime("%Y-%m-%d")),
        "run_date": (run_timestamp.strftime("%Y-%m-%d")),
        "run_time": (run_timestamp.strftime("%H:%M:%S")),
        "run_by": (run_by),
        "overall_status": (overall_status),
        "tableau_workbook": str(source_tableau_path),
        "tableau_worksheet": (DIRECT_COST_TABLEAU_WORKSHEET),
        "financial_statements": {
            "ALIV": str(aliv_fs_path),
            "CBL": str(cbl_fs_path),
        },
        "validation_workbook": str(validation_file),
        "raw_tableau_csv": str(raw_csv),
        "raw_tableau_workbook": str(raw_excel),
        "companies": {
            "ALIV": (aliv_result),
            "CBL": (cbl_result),
        },
    }

    # ========================================================
    # 13. TIMESTAMPED AUDIT LOG
    # ========================================================

    log_file = logs_folder / (
        "Direct_Cost_Validation_"
        f"{period_tag}_"
        f"{run_timestamp:%Y%m%d_%H%M%S}"
        ".log"
    )

    log_file.write_text(
        "\n".join(
            [
                ("DIRECT COST ACTUALS VALIDATION"),
                ("=" * 70),
                "",
                ("RUN INFORMATION"),
                (f"Run Date: " f"{run_timestamp:%m/%d/%Y}"),
                (f"Run Time: " f"{run_timestamp:%I:%M:%S %p}"),
                (f"Run By: " f"{run_by}"),
                "",
                ("VALIDATION PERIOD"),
                (f"Closed Period: " f"{closed_period:%B %Y}"),
                (f"Forecast Start: " f"{forecast_start:%m/%d/%Y}"),
                "",
                ("SOURCE FILES"),
                (f"ALIV Financial Statement: " f"{Path(aliv_fs_path).resolve()}"),
                (f"CBL Financial Statement: " f"{Path(cbl_fs_path).resolve()}"),
                (f"Tableau Workbook: " f"{Path(tableau_workbook_path).resolve()}"),
                (f"Tableau Worksheet: " f"{DIRECT_COST_TABLEAU_WORKSHEET}"),
                "",
                ("VALIDATION RESULTS"),
                (f"Overall Status: " f"{overall_status}"),
                (f"ALIV Status: " f"{aliv_result['status']}"),
                (f"CBL Status: " f"{cbl_result['status']}"),
                "",
                ("OUTPUT FILES"),
                (f"Validation Workbook: " f"{validation_file}"),
                (f"Raw Tableau CSV: " f"{raw_csv}"),
                (f"Raw Tableau Workbook: " f"{raw_excel}"),
                "",
                ("=" * 70),
            ]
        ),
        encoding="utf-8",
    )

    result["log_file"] = str(log_file)

    # ========================================================
    # COMPLETE
    # ========================================================

    print()
    print("=" * 70)
    print("DIRECT COST ACTUALS VALIDATION COMPLETE")
    print("=" * 70)

    print(
        "Overall:",
        overall_status,
    )

    print(
        "ALIV:",
        aliv_result["status"],
    )

    print(
        "CBL:",
        cbl_result["status"],
    )

    print(
        "Workbook:",
        validation_file,
    )

    print(
        "Audit Log:",
        log_file,
    )

    print("=" * 70)

    return result
