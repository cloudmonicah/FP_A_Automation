from __future__ import annotations
from shared.friendly_errors import explain

from collections import defaultdict
from copy import copy
from pathlib import Path
from datetime import datetime
import os
import re

from openpyxl import load_workbook
from openpyxl.styles import Border, Side
from openpyxl.utils import get_column_letter, range_boundaries

from automations.actuals.tableau_excel_update_service import (
    FISCAL_MONTH_ORDER,
    MONTH_ABBREVIATIONS,
    build_excel_header_lookup,
    build_excel_month_number_lookup,
    build_fiscal_month_headers,
    build_fiscal_month_keys,
    find_excel_column_by_alias,
    find_excel_header_row,
    get_month_number_from_header,
    normalize_text as shared_normalize_text,
    parse_fiscal_year,
)
from automations.forecast.mapping import FORECAST_ROUTES

STANDARD_FONT_NAME = "Aptos"
STANDARD_FONT_SIZE = 12
MAX_AUTOFIT_WIDTH = 60
AUTOFIT_PADDING = 2

TABLEAU_FIELD_ALIASES = {
    "business_code": [
        "GL Business Code",
        "Bus Code Name",
        "Business Code",
        "Bus Code",
    ],
    "company": [
        "Company",
        "Company Name",
        "Entity",
    ],
    "fs_category": [
        "GL FS Category",
        "FS Category",
        "Financial Statement Category",
    ],
    "account_mapping": [
        "GL Account Mapping",
        "Account Mapping",
        "GL Mapping",
    ],
    "sub_segment": [
        "GL Sub Segment code",
        "GL Sub Segment Code",
        "GL Sub Segment",
        "Sub Segment Code",
        "Sub Segment",
        "Account",
        "Account Code",
        "GL Account",
    ],
    "description": [
        "GL Sub Account Description",
        "GL Sub Account Desc",
        "Sub Account Description",
        "Sub Account Desc",
        "Account Description",
        "Description",
    ],
    "grand_total": [
        "Grand Total",
        "Total",
    ],
}

CANONICAL_FIELD_MAPPING = {
    "business_code": "Bus Code Name",
    "company": "Company",
    "fs_category": "GL FS Category",
    "account_mapping": "GL Account Mapping",
    "sub_segment": "GL Sub Segment Code",
    "description": "GL Sub Account Description",
}


def resolve_target_filename(
    mapped_filename: str,
    forecast_month: str,
) -> str:

    month = forecast_month.strip().title()[:3]

    valid_months = {
        "Jan",
        "Feb",
        "Mar",
        "Apr",
        "May",
        "Jun",
        "Jul",
        "Aug",
        "Sep",
        "Oct",
        "Nov",
        "Dec",
    }

    if month not in valid_months:

        raise ValueError(f'Invalid forecast month "{forecast_month}".')

    pattern = r"_(Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)" r"\s*Forecast"

    resolved, count = re.subn(
        pattern,
        f"_{month} Forecast",
        mapped_filename,
        flags=re.IGNORECASE,
    )

    if count == 0:

        raise RuntimeError(
            "Could not identify a forecast month in mapped workbook: "
            f"{mapped_filename}"
        )

    return resolved


def normalize_text(value) -> str:

    return shared_normalize_text(value)


def normalize_sheet_name(value) -> str:

    return re.sub(
        r"[^a-z0-9]+",
        "",
        str(value or "").strip().lower(),
    )


def convert_excel_value(value):

    if value is None:
        return None

    if isinstance(value, (int, float)):
        return value

    text = str(value).strip()

    if not text:
        return None

    number_text = text.replace(",", "").replace("$", "")

    if number_text.startswith("(") and number_text.endswith(")"):

        number_text = "-" + number_text[1:-1]

    try:

        number = float(number_text)

        if number.is_integer():
            return int(number)

        return number

    except ValueError:

        return text


def find_tableau_field(
    row: dict,
    canonical_field: str,
):

    aliases = TABLEAU_FIELD_ALIASES.get(
        canonical_field,
        [],
    )

    # --------------------------------------------------------
    # Build normalized lookup once for this Tableau row.
    # --------------------------------------------------------

    normalized_headers = {normalize_text(header): header for header in row.keys()}

    for alias in aliases:

        actual_header = normalized_headers.get(normalize_text(alias))

        if actual_header is not None:

            return (
                actual_header,
                row.get(actual_header),
            )

    return (
        None,
        None,
    )


def is_tableau_grand_total_row(
    row: dict,
) -> bool:

    _, business_code = find_tableau_field(
        row,
        "business_code",
    )

    _, sub_segment = find_tableau_field(
        row,
        "sub_segment",
    )

    _, description = find_tableau_field(
        row,
        "description",
    )

    if normalize_text(business_code) == "grand total":
        return True

    if (
        normalize_text(sub_segment) == "total"
        and normalize_text(description) == "total"
    ):
        return True

    if normalize_text(description) == "grand total":
        return True

    return False


def get_tableau_detail_rows(
    tableau_rows: list[dict],
) -> list[dict]:

    detail_rows = []

    for row in tableau_rows:

        if is_tableau_grand_total_row(row):
            continue

        if not any(value not in (None, "") for value in row.values()):
            continue

        detail_rows.append(row)

    return detail_rows


def get_tableau_company_value(
    row: dict,
):

    _, company = find_tableau_field(
        row,
        "company",
    )

    return company


def filter_rows_for_company(
    tableau_rows: list[dict],
    company: str,
) -> list[dict]:

    wanted = normalize_text(company)

    filtered = []

    for row in get_tableau_detail_rows(tableau_rows):

        actual = normalize_text(get_tableau_company_value(row))

        if actual == wanted:
            filtered.append(row)

    return filtered


def build_workbook_groups(
    forecast_month: str,
) -> dict[str, list[dict]]:

    groups = defaultdict(list)

    for configured_route in FORECAST_ROUTES:

        route = dict(configured_route)

        route["target_workbook"] = resolve_target_filename(
            mapped_filename=route["workbook"],
            forecast_month=forecast_month,
        )

        groups[route["target_workbook"]].append(route)

    return dict(groups)


def get_configured_excel_sheet(
    workbook,
    configured_sheet_name: str,
):

    if configured_sheet_name in workbook.sheetnames:

        return workbook[configured_sheet_name]

    wanted = normalize_sheet_name(configured_sheet_name)

    matches = [
        worksheet
        for worksheet in workbook.worksheets
        if normalize_sheet_name(worksheet.title) == wanted
    ]

    if len(matches) == 1:
        return matches[0]

    if not matches:

        raise RuntimeError(
            "Configured Excel worksheet was not found. "
            f"Expected={configured_sheet_name!r}; "
            f"Available={workbook.sheetnames}"
        )

    raise RuntimeError(
        "Configured Excel worksheet is ambiguous after normalization. "
        f"Expected={configured_sheet_name!r}; "
        f"Matches={[ws.title for ws in matches]}"
    )


def get_forecast_structure(
    worksheet,
    fiscal_year,
) -> dict:
    """
    Identify the structure of a forecast worksheet.

    IMPORTANT:
    Some older forecast workbooks contain month-header formulas
    such as:

        Jul-26
        =EDATE(H1,1)
        =EDATE(I1,1)
        ...

    openpyxl does not calculate those formulas.

    If all 12 fiscal months cannot be identified normally,
    this function uses the recognized July column as the start
    of the fiscal-month block and maps the next 11 consecutive
    columns through June.

    The headers are later replaced with REAL Excel dates by
    finalize_forecast_sheet().
    """

    # ========================================================
    # FIND HEADER ROW
    # ========================================================

    header_row = find_excel_header_row(worksheet)

    # ========================================================
    # BUILD HEADER LOOKUP
    # ========================================================

    header_lookup = build_excel_header_lookup(
        worksheet,
        header_row,
    )

    # ========================================================
    # TRY NORMAL MONTH DETECTION FIRST
    # ========================================================

    month_columns = build_excel_month_number_lookup(
        worksheet,
        header_row,
    )

    # ========================================================
    # FALLBACK FOR OLD =EDATE(...) MONTH HEADERS
    # ========================================================
    #
    # If all 12 months were not recognized, but July was
    # recognized, use July as the start of the fiscal block.
    #
    # Fiscal order:
    #
    # Jul Aug Sep Oct Nov Dec Jan Feb Mar Apr May Jun
    #
    # This does NOT retain the EDATE formulas.
    # finalize_forecast_sheet() will replace these headers
    # with actual datetime values.
    # ========================================================

    missing_months = [
        month for month in FISCAL_MONTH_ORDER if month not in month_columns
    ]

    if missing_months:

        july_column = month_columns.get(7)

        if isinstance(july_column, int):

            # ------------------------------------------------
            # Make sure 12 consecutive columns physically
            # exist beginning with July.
            # ------------------------------------------------

            last_fiscal_column = july_column + len(FISCAL_MONTH_ORDER) - 1

            if last_fiscal_column <= worksheet.max_column:

                month_columns = {}

                for offset, month_number in enumerate(FISCAL_MONTH_ORDER):
                    month_columns[month_number] = july_column + offset

                print(
                    f"[MONTH HEADER REPAIR] "
                    f'"{worksheet.title}": '
                    "formula-based fiscal headers detected. "
                    "Jul-Jun columns were mapped from the "
                    f"July starting column {july_column}."
                )

    # ========================================================
    # FINAL MONTH VALIDATION
    # ========================================================

    missing_months = [
        MONTH_ABBREVIATIONS[month]
        for month in FISCAL_MONTH_ORDER
        if (
            month not in month_columns
            or not isinstance(
                month_columns.get(month),
                int,
            )
        )
    ]

    if missing_months:

        raise RuntimeError(
            f'Worksheet "{worksheet.title}" does not contain '
            "a complete Jul-Jun forecast month block. "
            f"Missing: {', '.join(missing_months)}"
        )

    # ========================================================
    # FORECAST TYPE
    # ========================================================

    forecast_type_column = header_lookup.get(normalize_text("Forecast Type"))

    if forecast_type_column is None:

        raise RuntimeError(
            f"Forecast Type column was not found in " f'"{worksheet.title}".'
        )

    # ========================================================
    # DESCRIPTIVE COLUMNS
    # ========================================================

    descriptive_columns = {}

    for (
        canonical_tableau_field,
        canonical_excel_field,
    ) in CANONICAL_FIELD_MAPPING.items():

        column_index = find_excel_column_by_alias(
            header_lookup,
            canonical_excel_field,
        )

        if column_index is not None:

            descriptive_columns[canonical_tableau_field] = column_index

    # ========================================================
    # GRAND TOTAL COLUMN
    # ========================================================

    grand_total_column = None

    grand_total_aliases = {
        "grand total",
        "total",
    }

    for (
        normalized_header,
        column_index,
    ) in header_lookup.items():

        if normalized_header in grand_total_aliases:

            grand_total_column = column_index
            break

    # ========================================================
    # RETURN STRUCTURE
    # ========================================================

    return {
        "header_row": header_row,
        "header_lookup": header_lookup,
        "forecast_type_column": forecast_type_column,
        "descriptive_columns": descriptive_columns,
        "month_columns": month_columns,
        "grand_total_column": grand_total_column,
        "fiscal_year": parse_fiscal_year(fiscal_year),
    }


def get_existing_forecast_type(
    worksheet,
    structure: dict,
):

    forecast_type_column = structure["forecast_type_column"]

    first_data_row = structure["header_row"] + 1

    for row_number in range(
        first_data_row,
        worksheet.max_row + 1,
    ):

        value = worksheet.cell(
            row=row_number,
            column=forecast_type_column,
        ).value

        if value is None:
            continue

        text = str(value).strip()

        if text:
            return value

    return None


def apply_style(
    cell,
    style,
):

    cell.font = copy(style["font"])

    cell.fill = copy(style["fill"])

    cell.border = copy(style["border"])

    cell.alignment = copy(style["alignment"])

    cell.number_format = style["number_format"]

    cell.protection = copy(style["protection"])


MATCH_STRATEGIES = (
    (
        "full",
        (
            "business_code",
            "company",
            "fs_category",
            "account_mapping",
            "sub_segment",
            "description",
        ),
    ),
    (
        "description_primary",
        (
            "business_code",
            "company",
            "account_mapping",
            "sub_segment",
            "description",
        ),
    ),
    (
        "primary",
        (
            "business_code",
            "company",
            "account_mapping",
            "sub_segment",
        ),
    ),
    (
        "secondary",
        (
            "business_code",
            "company",
            "sub_segment",
        ),
    ),
    (
        "tertiary",
        (
            "business_code",
            "company",
            "account_mapping",
        ),
    ),
    (
        "description",
        (
            "business_code",
            "company",
            "description",
        ),
    ),
)


def is_excel_formula(value) -> bool:
    """Return True when an Excel cell contains a formula."""

    return isinstance(value, str) and value.startswith("=")


def is_excel_total_row(worksheet, row_number: int, structure: dict) -> bool:
    """Return True for a recognized Total or Grand Total row."""

    descriptive_columns = structure["descriptive_columns"]

    for column_number in descriptive_columns.values():
        value = normalize_text(
            worksheet.cell(row=row_number, column=column_number).value
        )

        if value in {"total", "grand total"}:
            return True

    return False


def build_tableau_fiscal_month_lookup(
    tableau_rows: list[dict],
    fiscal_year,
) -> dict[int, str]:
    """
    Build a PARTIAL month-number -> Tableau-header mapping.

    Tableau is allowed to omit fiscal months. Missing months are
    handled as blank financial values during the update rather than
    failing the workbook.
    """

    if not tableau_rows:
        raise RuntimeError("No Tableau rows were supplied.")

    # Validate the GUI Fiscal Year even though month matching itself
    # is based on month number.
    build_fiscal_month_keys(fiscal_year)

    month_lookup = {}
    first_row = tableau_rows[0]

    for header in first_row.keys():
        month_number = get_month_number_from_header(header)

        if month_number is None:
            continue

        if month_number in month_lookup:
            raise RuntimeError(
                "Tableau returned more than one column for "
                f"{MONTH_ABBREVIATIONS[month_number]}."
            )

        month_lookup[month_number] = header

    if not month_lookup:
        raise RuntimeError(
            "Tableau did not return any recognizable fiscal month columns."
        )

    missing_months = [
        MONTH_ABBREVIATIONS[month]
        for month in FISCAL_MONTH_ORDER
        if month not in month_lookup
    ]

    if missing_months:
        print(
            "[PARTIAL MONTHS] Tableau is missing: "
            + ", ".join(missing_months)
            + ". Those Excel month values will be blanked."
        )

    return month_lookup


def standardize_fiscal_month_headers(
    worksheet,
    structure: dict,
    fiscal_year,
):
    """
    Write the Jul-Jun fiscal headers as REAL Excel date values.

    Example for FY27:

        Jul-26 -> underlying value 2026-07-01
        Aug-26 -> underlying value 2026-08-01
        ...
        Jun-27 -> underlying value 2027-06-01

    Excel displays the values as mmm-yy while retaining a true
    date value underneath.
    """

    # Convert values such as "FY27" to the fiscal ending year.
    fiscal_year_number = parse_fiscal_year(fiscal_year)

    # FY27 begins in July 2026 and ends in June 2027.
    start_year = fiscal_year_number - 1

    end_year = fiscal_year_number

    month_columns = structure["month_columns"]

    header_row = structure["header_row"]

    for month_number in FISCAL_MONTH_ORDER:

        column_number = month_columns[month_number]

        # July-December belong to the calendar year before the
        # fiscal ending year.
        if month_number >= 7:

            calendar_year = start_year

        # January-June belong to the fiscal ending year.
        else:

            calendar_year = end_year

        header_cell = worksheet.cell(
            row=header_row,
            column=column_number,
        )

        # Store a genuine Python datetime so openpyxl writes a
        # genuine Excel date serial.
        header_cell.value = datetime(
            calendar_year,
            month_number,
            1,
        )

        # Display exactly like Jul-26, Aug-26, Sep-26, etc.
        header_cell.number_format = "mmm-yy"


def get_forecast_excel_table(worksheet, structure: dict):
    """
    Return the Excel Table/ListObject that contains the forecast header.

    If a workbook uses a formatted range instead of a real Excel Table,
    None is returned and the service falls back to the detected total row.
    """

    header_row = structure["header_row"]
    july_column = structure["month_columns"][7]

    for table in worksheet.tables.values():
        min_col, min_row, max_col, max_row = range_boundaries(table.ref)

        if min_row == header_row and min_col <= july_column <= max_col:
            return table

    return None


def find_excel_grand_total_row(worksheet, structure: dict):
    """Find the first recognized total row below the forecast header."""

    first_data_row = structure["header_row"] + 1

    for row_number in range(first_data_row, worksheet.max_row + 1):
        if is_excel_total_row(worksheet, row_number, structure):
            return row_number

    return None


def get_data_region(worksheet, structure: dict) -> dict:
    """
    Describe the current forecast detail region and safe insertion point.

    For a real Excel Table, insertion occurs before its totals row when
    totals are shown, otherwise directly after the last table data row.
    For a normal range, insertion occurs before Grand Total.
    """

    table = get_forecast_excel_table(worksheet, structure)
    first_data_row = structure["header_row"] + 1

    if table is not None:
        min_col, min_row, max_col, max_row = range_boundaries(table.ref)
        totals_shown = bool(getattr(table, "totalsRowShown", False))

        if totals_shown:
            insertion_row = max_row
            last_detail_row = max_row - 1
        else:
            insertion_row = max_row + 1
            last_detail_row = max_row

        return {
            "table": table,
            "table_ref_before": table.ref,
            "table_min_col": min_col,
            "table_max_col": max_col,
            "table_min_row": min_row,
            "table_max_row": max_row,
            "totals_shown": totals_shown,
            "first_data_row": first_data_row,
            "last_detail_row": last_detail_row,
            "insertion_row": insertion_row,
            "grand_total_row": (
                max_row
                if totals_shown
                else find_excel_grand_total_row(
                    worksheet,
                    structure,
                )
            ),
        }

    grand_total_row = find_excel_grand_total_row(worksheet, structure)

    if grand_total_row is None:
        insertion_row = worksheet.max_row + 1
        last_detail_row = worksheet.max_row
    else:
        insertion_row = grand_total_row
        last_detail_row = grand_total_row - 1

    return {
        "table": None,
        "table_ref_before": None,
        "table_min_col": 1,
        "table_max_col": worksheet.max_column,
        "table_min_row": structure["header_row"],
        "table_max_row": grand_total_row or worksheet.max_row,
        "totals_shown": False,
        "first_data_row": first_data_row,
        "last_detail_row": last_detail_row,
        "insertion_row": insertion_row,
        "grand_total_row": grand_total_row,
    }


def rebuild_total_formulas(
    worksheet,
    structure: dict,
):
    """
    Rebuild all forecast formulas after Tableau detail rows
    have been written.

    Creates:

    1. A Grand Total formula for every detail row.
    2. A SUM formula at the bottom of every fiscal month.
    3. An overall Grand Total formula.

    IMPORTANT:
    - Month columns are handled individually.
    - No min()/max() calculation is performed across month
      columns.
    - Missing/invalid month-column positions are skipped safely.
    """

    # ========================================================
    # FIND CURRENT GRAND TOTAL ROW
    # ========================================================

    total_row = find_excel_grand_total_row(
        worksheet,
        structure,
    )

    if total_row is None:
        print(f'[FORMULAS] "{worksheet.title}": ' "No Grand Total row found.")
        return None

    # ========================================================
    # DETAIL REGION
    # ========================================================

    first_detail_row = structure["header_row"] + 1

    last_detail_row = total_row - 1

    if last_detail_row < first_detail_row:
        return total_row

    month_columns = structure["month_columns"]

    # ========================================================
    # GET VALID FISCAL MONTH COLUMNS
    # ========================================================
    #
    # Do NOT use:
    #
    #     min(month_columns.values())
    #     max(month_columns.values())
    #
    # because some workbooks can temporarily contain None for
    # a month-column position.
    # ========================================================

    valid_month_columns = []

    for month_number in FISCAL_MONTH_ORDER:

        column_number = month_columns.get(month_number)

        if isinstance(column_number, int):
            valid_month_columns.append(column_number)

    if not valid_month_columns:
        raise RuntimeError(
            f'Worksheet "{worksheet.title}" has no valid ' "fiscal month columns."
        )

    # ========================================================
    # DETAIL ROW GRAND TOTAL FORMULAS
    # ========================================================

    grand_total_column = structure.get("grand_total_column")

    if isinstance(grand_total_column, int):

        month_references = [
            get_column_letter(column_number) for column_number in valid_month_columns
        ]

        # ----------------------------------------------------
        # Use the individual recognized month cells rather
        # than assuming Jul-Jun are always one uninterrupted
        # range.
        # ----------------------------------------------------

        for row_number in range(
            first_detail_row,
            last_detail_row + 1,
        ):

            references = [
                f"{column_letter}{row_number}" for column_letter in month_references
            ]

            worksheet.cell(
                row=row_number,
                column=grand_total_column,
            ).value = (
                "=SUM(" + ",".join(references) + ")"
            )

    # ========================================================
    # MONTHLY TOTAL FORMULAS
    # ========================================================

    for month_number in FISCAL_MONTH_ORDER:

        column_number = month_columns.get(month_number)

        # ----------------------------------------------------
        # Skip anything that isn't an actual Excel column.
        # ----------------------------------------------------

        if not isinstance(column_number, int):
            continue

        column_letter = get_column_letter(column_number)

        worksheet.cell(
            row=total_row,
            column=column_number,
        ).value = (
            f"=SUM("
            f"{column_letter}{first_detail_row}:"
            f"{column_letter}{last_detail_row}"
            f")"
        )

    # ========================================================
    # OVERALL GRAND TOTAL
    # ========================================================

    if isinstance(grand_total_column, int):

        grand_total_letter = get_column_letter(grand_total_column)

        worksheet.cell(
            row=total_row,
            column=grand_total_column,
        ).value = (
            f"=SUM("
            f"{grand_total_letter}{first_detail_row}:"
            f"{grand_total_letter}{last_detail_row}"
            f")"
        )

    # ========================================================
    # CONFIRMATION
    # ========================================================

    print(
        f'[FORMULAS REBUILT] "{worksheet.title}": '
        f"detail rows {first_detail_row}-"
        f"{last_detail_row}; "
        f"Grand Total row {total_row}."
    )

    return total_row


def update_existing_forecast_rows(
    worksheet,
    tableau_datasets,
    structure: dict,
    fiscal_year,
    workbook_name: str,
) -> dict:

    # ========================================================
    # 1. COLLECT ALL TABLEAU DETAIL ROWS
    # ========================================================

    combined_rows = []

    dataset_diagnostics = []

    is_revenue_workbook = "revenue" in workbook_name.lower()

    if is_revenue_workbook:
        print(
            f'[REVENUE SIGN RULE] "{workbook_name}": '
            "Tableau financial values will be multiplied by -1."
        )

    for tableau_worksheet_name, tableau_rows in tableau_datasets:

        # Remove Tableau Grand Total and completely blank rows.
        detail_rows = get_tableau_detail_rows(tableau_rows)

        if not detail_rows:
            raise RuntimeError(
                "No usable Tableau detail rows were returned for "
                f'"{tableau_worksheet_name}".'
            )

        tableau_months = build_tableau_fiscal_month_lookup(
            tableau_rows=detail_rows,
            fiscal_year=fiscal_year,
        )

        print(
            f'Rebuilding from "{tableau_worksheet_name}": '
            f"{len(detail_rows)} Tableau detail row(s)"
        )

        dataset_diagnostics.append(
            {
                "tableau_worksheet": tableau_worksheet_name,
                "detail_rows": len(detail_rows),
            }
        )

        # Keep the Tableau row order exactly as exported.
        for tableau_row in detail_rows:

            combined_rows.append(
                {
                    "tableau_worksheet": tableau_worksheet_name,
                    "tableau_row": tableau_row,
                    "tableau_months": tableau_months,
                }
            )

    tableau_detail_count = len(combined_rows)

    if tableau_detail_count == 0:
        raise RuntimeError(
            f'No Tableau detail rows are available for "{worksheet.title}".'
        )

    print(
        f'[TABLEAU SOURCE] "{worksheet.title}": '
        f"{tableau_detail_count} total detail row(s)"
    )

    # ========================================================
    # 2. CAPTURE EXISTING WORKBOOK INFORMATION
    # ========================================================

    header_row = structure["header_row"]

    first_data_row = header_row + 1

    forecast_type_column = structure["forecast_type_column"]

    descriptive_columns = structure["descriptive_columns"]

    month_columns = structure["month_columns"]

    grand_total_column = structure.get("grand_total_column")

    existing_forecast_type = get_existing_forecast_type(
        worksheet=worksheet,
        structure=structure,
    )

    if existing_forecast_type is None:
        raise RuntimeError(
            f'Forecast Type could not be determined in "{worksheet.title}".'
        )

    print(f'[FORECAST TYPE] "{worksheet.title}": ' f"{existing_forecast_type}")

    region = get_data_region(
        worksheet=worksheet,
        structure=structure,
    )

    excel_table = region["table"]

    old_first_data_row = region["first_data_row"]

    old_last_detail_row = region["last_detail_row"]

    old_detail_count = max(
        0,
        old_last_detail_row - old_first_data_row + 1,
    )

    print(
        f'[EXCEL BEFORE] "{worksheet.title}": '
        f"{old_detail_count} existing detail row(s)"
    )

    style_source_row = first_data_row

    row_style_template = {}

    for column_number in range(
        1,
        worksheet.max_column + 1,
    ):

        source_cell = worksheet.cell(
            row=style_source_row,
            column=column_number,
        )

        row_style_template[column_number] = {
            "font": copy(source_cell.font),
            "fill": copy(source_cell.fill),
            "border": copy(source_cell.border),
            "alignment": copy(source_cell.alignment),
            "number_format": source_cell.number_format,
            "protection": copy(source_cell.protection),
        }

    source_row_height = worksheet.row_dimensions[style_source_row].height

    row_difference = tableau_detail_count - old_detail_count

    # --------------------------------------------------------
    # TABLEAU HAS MORE ROWS THAN EXCEL
    # --------------------------------------------------------

    if row_difference > 0:

        # Insert the additional rows immediately after the
        # current detail region and before Grand Total.
        insertion_row = old_last_detail_row + 1

        worksheet.insert_rows(
            insertion_row,
            amount=row_difference,
        )

        print(
            f"[RESIZE] Inserted {row_difference} detail row(s) "
            f'into "{worksheet.title}".'
        )

    # --------------------------------------------------------
    # TABLEAU HAS FEWER ROWS THAN EXCEL
    # --------------------------------------------------------

    elif row_difference < 0:

        rows_to_delete = abs(row_difference)

        # Delete from the bottom of the OLD detail region.
        #
        # This keeps the header and beginning of the financial
        # detail area in their original positions.
        delete_start_row = first_data_row + tableau_detail_count

        worksheet.delete_rows(
            delete_start_row,
            amount=rows_to_delete,
        )

        print(
            f"[RESIZE] Removed {rows_to_delete} old detail row(s) "
            f'from "{worksheet.title}".'
        )

    else:

        print(
            f'[RESIZE] "{worksheet.title}" already has the '
            "required number of detail rows."
        )

    # ========================================================
    # 6. APPLY DETAIL-ROW FORMATTING TO THE FINAL REGION
    # ========================================================

    final_last_detail_row = first_data_row + tableau_detail_count - 1

    for row_number in range(
        first_data_row,
        final_last_detail_row + 1,
    ):

        for column_number, style in row_style_template.items():

            target_cell = worksheet.cell(
                row=row_number,
                column=column_number,
            )

            apply_style(
                target_cell,
                style,
            )

        if source_row_height is not None:
            worksheet.row_dimensions[row_number].height = source_row_height

    # ========================================================
    # 7. CLEAR OLD DETAIL VALUES

    for row_number in range(
        first_data_row,
        final_last_detail_row + 1,
    ):

        # Clear descriptive fields.
        for excel_column in descriptive_columns.values():

            worksheet.cell(
                row=row_number,
                column=excel_column,
            ).value = None

        # Clear Forecast Type before writing the approved value.
        worksheet.cell(
            row=row_number,
            column=forecast_type_column,
        ).value = None

        # Clear all fiscal month values.
        for month_number in FISCAL_MONTH_ORDER:

            excel_column = month_columns[month_number]

            worksheet.cell(
                row=row_number,
                column=excel_column,
            ).value = None

        # Clear row Grand Total.
        if grand_total_column is not None:

            worksheet.cell(
                row=row_number,
                column=grand_total_column,
            ).value = None

    # ========================================================
    # 8. WRITE TABLEAU ROWS DIRECTLY INTO EXCEL
    # ========================================================

    values_written = 0

    for row_offset, item in enumerate(combined_rows):

        excel_row = first_data_row + row_offset

        tableau_row = item["tableau_row"]

        tableau_months = item["tableau_months"]

        # ----------------------------------------------------
        # FORECAST TYPE
        # ----------------------------------------------------

        worksheet.cell(
            row=excel_row,
            column=forecast_type_column,
        ).value = existing_forecast_type

        values_written += 1

        # ----------------------------------------------------
        # DESCRIPTIVE TABLEAU FIELDS
        # ----------------------------------------------------

        for (
            canonical_field,
            excel_column,
        ) in descriptive_columns.items():

            _, tableau_value = find_tableau_field(
                tableau_row,
                canonical_field,
            )

            worksheet.cell(
                row=excel_row,
                column=excel_column,
            ).value = convert_excel_value(tableau_value)

            values_written += 1

        # ----------------------------------------------------
        # JUL-JUN FINANCIAL VALUES
        # ----------------------------------------------------

        for month_number in FISCAL_MONTH_ORDER:

            excel_column = month_columns[month_number]

            tableau_header = tableau_months.get(month_number)

            # If Tableau did not export this fiscal month,
            # explicitly write a blank.
            if tableau_header is None:

                new_value = None

            else:

                new_value = convert_excel_value(tableau_row.get(tableau_header))

                if (
                    is_revenue_workbook
                    and isinstance(new_value, (int, float))
                    and not isinstance(new_value, bool)
                ):
                    new_value = new_value * -1

            worksheet.cell(
                row=excel_row,
                column=excel_column,
            ).value = new_value

            values_written += 1

        # ----------------------------------------------------
        # ROW GRAND TOTAL FORMULA
        # ----------------------------------------------------

        if grand_total_column is not None:

            month_cells = [
                (f"{get_column_letter(month_columns[month_number])}" f"{excel_row}")
                for month_number in FISCAL_MONTH_ORDER
            ]

            worksheet.cell(
                row=excel_row,
                column=grand_total_column,
            ).value = (
                "=SUM(" + ",".join(month_cells) + ")"
            )

    # ========================================================
    # 9. RESIZE THE REAL EXCEL TABLE
    # ========================================================

    if excel_table is not None:

        (
            table_min_col,
            table_min_row,
            table_max_col,
            _,
        ) = range_boundaries(excel_table.ref)

        totals_shown = bool(
            getattr(
                excel_table,
                "totalsRowShown",
                False,
            )
        )

        if totals_shown:

            # Header + Tableau detail rows + totals row.
            table_max_row = header_row + tableau_detail_count + 1

        else:

            # Header + Tableau detail rows.
            table_max_row = header_row + tableau_detail_count

        excel_table.ref = (
            f"{get_column_letter(table_min_col)}"
            f"{table_min_row}:"
            f"{get_column_letter(table_max_col)}"
            f"{table_max_row}"
        )

        print(f'[TABLE RESIZED] "{worksheet.title}": ' f"{excel_table.ref}")

    # ========================================================
    # 10. REBUILD BOTTOM GRAND TOTAL FORMULAS
    # ========================================================

    rebuild_total_formulas(
        worksheet=worksheet,
        structure=structure,
    )

    # ========================================================
    # 11. VALIDATE FINAL DETAIL ROW COUNT
    # ========================================================

    final_region = get_data_region(
        worksheet=worksheet,
        structure=structure,
    )

    final_first_data_row = final_region["first_data_row"]

    final_last_detail_row = final_region["last_detail_row"]

    excel_detail_count = max(
        0,
        (final_last_detail_row - final_first_data_row + 1),
    )

    print(
        f'[VALIDATION] "{worksheet.title}": '
        f"Tableau={tableau_detail_count}; "
        f"Excel={excel_detail_count}"
    )

    if excel_detail_count != tableau_detail_count:

        raise RuntimeError(
            f'Row-count validation failed for "{worksheet.title}". '
            f"Tableau contains {tableau_detail_count} detail row(s), "
            f"but Excel contains {excel_detail_count} detail row(s). "
            "Workbook will not be saved."
        )

    # ========================================================
    # 12. SUCCESS RESULT
    # ========================================================

    print(
        f'[REBUILD COMPLETE] "{worksheet.title}": '
        f"{tableau_detail_count} Tableau row(s) written exactly."
    )

    return {
        "matched_rows": tableau_detail_count,
        "inserted_rows": max(
            row_difference,
            0,
        ),
        "values_updated": values_written,
        "formula_cells_preserved": 0,
        "inserted_row_details": [],
        "unmatched_rows": [],
        "plan": dataset_diagnostics,
        "tableau_detail_rows": tableau_detail_count,
        "excel_detail_rows": excel_detail_count,
    }


def standardize_forecast_font(
    worksheet,
):
    """
    Standardize the processed forecast worksheet to Aptos 12.

    IMPORTANT:
    Only the font NAME and SIZE are changed.

    Existing:
        bold
        italic
        underline
        font color
        fills
        borders
        alignment
        number formats

    are retained.
    """

    for row in worksheet.iter_rows():

        for cell in row:

            # ------------------------------------------------
            # Skip completely unused cells.
            # ------------------------------------------------

            if cell.value is None and not cell.has_style:
                continue

            new_font = copy(cell.font)

            new_font.name = STANDARD_FONT_NAME

            new_font.sz = STANDARD_FONT_SIZE

            cell.font = new_font


def get_display_length(
    value,
) -> int:
    """
    Return a practical character length for column auto-sizing.
    """

    if value is None:
        return 0

    text = str(value)

    # --------------------------------------------------------
    # For multi-line text, use the longest individual line.
    # --------------------------------------------------------

    return max(
        (len(line) for line in text.splitlines()),
        default=0,
    )


def autofit_forecast_columns(
    worksheet,
):

    for column_number in range(
        1,
        worksheet.max_column + 1,
    ):

        maximum_length = 0

        for row_number in range(
            1,
            worksheet.max_row + 1,
        ):

            cell = worksheet.cell(
                row=row_number,
                column=column_number,
            )

            maximum_length = max(
                maximum_length,
                get_display_length(cell.value),
            )

        adjusted_width = min(
            maximum_length + AUTOFIT_PADDING,
            MAX_AUTOFIT_WIDTH,
        )

        # ----------------------------------------------------
        # Avoid creating columns that are unnecessarily narrow.
        # ----------------------------------------------------

        adjusted_width = max(
            adjusted_width,
            8,
        )

        column_letter = get_column_letter(column_number)

        worksheet.column_dimensions[column_letter].width = adjusted_width


def apply_dynamic_july_freeze(
    worksheet,
    structure: dict,
):

    header_row = structure["header_row"]

    month_columns = structure["month_columns"]

    july_column = month_columns.get(7)

    if july_column is None:

        raise RuntimeError(f'July column could not be located in "{worksheet.title}".')

    freeze_row = header_row + 1

    freeze_cell = worksheet.cell(
        row=freeze_row,
        column=july_column,
    )

    worksheet.freeze_panes = freeze_cell.coordinate

    print(f'Freeze panes for "{worksheet.title}": ' f"{freeze_cell.coordinate}")


def apply_total_and_grand_total_borders(worksheet, structure: dict):

    double_side = Side(style="double")
    ledger_left_side = Side(style="thin")
    financial_columns = [
        structure["month_columns"][month] for month in FISCAL_MONTH_ORDER
    ]

    grand_total_column = structure.get("grand_total_column")

    if grand_total_column is not None:
        financial_columns.append(grand_total_column)

    first_data_row = structure["header_row"] + 1

    for row_number in range(first_data_row, worksheet.max_row + 1):
        if not is_excel_total_row(worksheet, row_number, structure):
            continue

        for column_number in financial_columns:
            cell = worksheet.cell(row=row_number, column=column_number)
            old = cell.border
            cell.border = Border(
                left=old.left,
                right=old.right,
                top=double_side,
                bottom=double_side,
                diagonal=old.diagonal,
                diagonal_direction=old.diagonal_direction,
                diagonalUp=old.diagonalUp,
                diagonalDown=old.diagonalDown,
                outline=old.outline,
                vertical=old.vertical,
                horizontal=old.horizontal,
            )

    if grand_total_column is not None:
        for row_number in range(structure["header_row"], worksheet.max_row + 1):
            cell = worksheet.cell(row=row_number, column=grand_total_column)
            old = cell.border
            cell.border = Border(
                left=ledger_left_side,
                right=old.right,
                top=old.top,
                bottom=old.bottom,
                diagonal=old.diagonal,
                diagonal_direction=old.diagonal_direction,
                diagonalUp=old.diagonalUp,
                diagonalDown=old.diagonalDown,
                outline=old.outline,
                vertical=old.vertical,
                horizontal=old.horizontal,
            )


from datetime import datetime
from openpyxl.utils import get_column_letter


def finalize_forecast_sheet(
    worksheet,
    structure: dict,
    fiscal_year,
):

    # ========================================================
    # FISCAL YEAR
    # FY27 = July 2026 through June 2027
    # ========================================================

    fy = parse_fiscal_year(fiscal_year)

    start_year = fy - 1

    fiscal_dates = {
        7: datetime(start_year, 7, 1),
        8: datetime(start_year, 8, 1),
        9: datetime(start_year, 9, 1),
        10: datetime(start_year, 10, 1),
        11: datetime(start_year, 11, 1),
        12: datetime(start_year, 12, 1),
        1: datetime(fy, 1, 1),
        2: datetime(fy, 2, 1),
        3: datetime(fy, 3, 1),
        4: datetime(fy, 4, 1),
        5: datetime(fy, 5, 1),
        6: datetime(fy, 6, 1),
    }

    header_row = structure["header_row"]
    month_columns = structure["month_columns"]

    # ========================================================
    # 1. REPLACE MONTH HEADER FORMULAS WITH REAL DATES
    # ========================================================

    for month_number in FISCAL_MONTH_ORDER:

        column_number = month_columns.get(month_number)

        if not isinstance(column_number, int):
            continue

        cell = worksheet.cell(
            row=header_row,
            column=column_number,
        )

        # Real Excel date -- NOT =EDATE(...)
        cell.value = fiscal_dates[month_number]

        # Display Jul-26, Aug-26, etc.
        cell.number_format = "mmm-yy"

    # ========================================================
    # CREATE / IDENTIFY GRAND TOTAL ROW
    # ========================================================
    #
    # Do NOT depend on a Grand Total row already existing.
    #
    # Tableau has already rebuilt the detail region, so the
    # Grand Total row belongs immediately below the final
    # detail row.
    # ========================================================

    first_detail_row = header_row + 1

    # --------------------------------------------------------
    # Find the last actual detail row using Forecast Type.
    #
    # Forecast Type is populated on every legitimate detail
    # row written from Tableau.
    # --------------------------------------------------------

    forecast_type_column = structure["forecast_type_column"]

    last_detail_row = first_detail_row - 1

    for row_number in range(
        first_detail_row,
        worksheet.max_row + 1,
    ):

        value = worksheet.cell(
            row=row_number,
            column=forecast_type_column,
        ).value

        if value not in (None, ""):
            last_detail_row = row_number

    if last_detail_row < first_detail_row:
        raise RuntimeError(
            f"No forecast detail rows were found in " f'"{worksheet.title}".'
        )

    # --------------------------------------------------------
    # Grand Total belongs directly after the final detail row.
    # --------------------------------------------------------

    total_row = last_detail_row + 1

    month_columns = structure["month_columns"]

    grand_total_column = structure.get("grand_total_column")

    # ========================================================
    # LABEL THE TOTAL ROW
    # ========================================================
    #
    # Put "Grand Total" in the first descriptive area.
    # Forecast Type is the safest consistent column.
    # ========================================================

    # worksheet.cell(
    #     row=total_row,
    #     column=forecast_type_column,
    # ).value = "Grand Total"

    # ========================================================
    # DETAIL-ROW GRAND TOTAL FORMULAS
    # ========================================================

    if isinstance(grand_total_column, int):

        for row_number in range(
            first_detail_row,
            last_detail_row + 1,
        ):

            month_cells = []

            for month_number in FISCAL_MONTH_ORDER:

                column_number = month_columns.get(month_number)

                if isinstance(
                    column_number,
                    int,
                ):

                    month_cells.append(
                        f"{get_column_letter(column_number)}" f"{row_number}"
                    )

            worksheet.cell(
                row=row_number,
                column=grand_total_column,
            ).value = (
                "=SUM(" + ",".join(month_cells) + ")"
            )

    # ========================================================
    # MONTHLY TOTAL FORMULAS
    # ========================================================

    for month_number in FISCAL_MONTH_ORDER:

        column_number = month_columns.get(month_number)

        if not isinstance(
            column_number,
            int,
        ):
            continue

        column_letter = get_column_letter(column_number)

        worksheet.cell(
            row=total_row,
            column=column_number,
        ).value = (
            f"=SUM("
            f"{column_letter}{first_detail_row}:"
            f"{column_letter}{last_detail_row}"
            f")"
        )

    # ========================================================
    # OVERALL GRAND TOTAL FORMULA
    # ========================================================

    if isinstance(grand_total_column, int):

        grand_total_letter = get_column_letter(grand_total_column)

        worksheet.cell(
            row=total_row,
            column=grand_total_column,
        ).value = (
            f"=SUM("
            f"{grand_total_letter}{first_detail_row}:"
            f"{grand_total_letter}{last_detail_row}"
            f")"
        )

    print(
        f'[FINALIZED] "{worksheet.title}": '
        f"month headers converted to real dates; "
        f"detail rows={first_detail_row}-{last_detail_row}; "
        f"Grand Total row created at {total_row}; "
        f"monthly and Grand Total formulas rebuilt."
    )


def get_forecast_automation_log_folders() -> dict:

    from shared.output_root import shared_output_root

    # Logs go to the shared Teams folder, never a personal OneDrive.
    corporate_onedrive = shared_output_root()

    automation_root = corporate_onedrive / "Logs" / "Forecast_Automation"

    # ========================================================
    # RAW TABLEAU COPIES
    # ========================================================

    raw_folder = automation_root / "Tableau_Raw_Copies"

    # ========================================================
    # RUN LOGS
    # ========================================================

    run_folder = automation_root / "Run_Logs"

    # ========================================================
    # CREATE FOLDERS
    # ========================================================
    #
    # exist_ok=True means this is safe when the folders already
    # exist.
    # ========================================================

    raw_folder.mkdir(
        parents=True,
        exist_ok=True,
    )

    run_folder.mkdir(
        parents=True,
        exist_ok=True,
    )

    # ========================================================
    # CONFIRM EXACT LOCATIONS IN TERMINAL
    # ========================================================

    print()
    print("=" * 70)
    print("FORECAST AUTOMATION LOG LOCATIONS")
    print("=" * 70)
    print(f"Shared folder      : {corporate_onedrive}")
    print(f"Automation Logs    : {automation_root}")
    print(f"Tableau Raw Copies : {raw_folder}")
    print(f"Run Logs           : {run_folder}")
    print("=" * 70)

    return {
        "root": automation_root,
        "raw": raw_folder,
        "runs": run_folder,
    }


def write_run_summary_log(log_path: Path, result: dict):
    """Write a concise permanent summary for one forecast automation run."""

    lines = [
        "FORECAST AUTOMATION RUN SUMMARY",
        "=" * 70,
        f"Timestamp: {datetime.now().isoformat(timespec='seconds')}",
        f"Fiscal year: {result.get('fiscal_year')}",
        f"Files selected: {result.get('files_found', 0)}",
        f"Successful: {len(result.get('updated', []))}",
        f"Failed: {len(result.get('failed', []))}",
        f"Missing: {len(result.get('missing', []))}",
        f"Manual action: {len(result.get('manual_required', []))}",
        "",
    ]

    for item in result.get("failed", []):
        lines.append(f"FAILED | {item.get('workbook')} | {item.get('error')}")

    log_path.write_text("\n".join(lines), encoding="utf-8")


def _forecast_total_cell_value(worksheet, coordinate, active_cells=None):
    """Read a number or evaluate the SUM formulas generated by this service.

    openpyxl cannot calculate formulas and newly saved workbooks have no cached
    totals. Only same-sheet SUM references are supported here. Unexpected
    formulas, circular references, and Excel errors are reported explicitly.
    """
    from decimal import Decimal, InvalidOperation

    active_cells = set() if active_cells is None else active_cells
    coordinate = coordinate.replace("$", "").upper()
    if coordinate in active_cells:
        raise ValueError(f"Circular formula at {worksheet.title}!{coordinate}")

    value = worksheet[coordinate].value
    if isinstance(value, str) and value.startswith("="):
        match = re.fullmatch(r"=SUM\(([^()]*)\)", value.strip(), re.IGNORECASE)
        if match is None:
            raise ValueError(
                f"Unsupported total formula at {worksheet.title}!{coordinate}: {value}"
            )
        active_cells.add(coordinate)
        total = Decimal("0")
        try:
            for argument in match.group(1).split(","):
                argument = argument.strip().replace("$", "")
                if re.fullmatch(r"[A-Z]+\d+:[A-Z]+\d+", argument, re.IGNORECASE):
                    for row in worksheet[argument]:
                        for cell in row:
                            total += _forecast_total_cell_value(
                                worksheet, cell.coordinate, active_cells
                            )
                elif re.fullmatch(r"[A-Z]+\d+", argument, re.IGNORECASE):
                    total += _forecast_total_cell_value(
                        worksheet, argument, active_cells
                    )
                else:
                    raise ValueError(f"Unsupported SUM argument: {argument!r}")
            return total
        finally:
            active_cells.remove(coordinate)

    # Blank cells inside a SUM contribute zero, just as they do in Excel.
    if value in (None, ""):
        return Decimal("0")
    if isinstance(value, bool):
        raise ValueError(f"Unexpected Boolean at {coordinate}")
    text = str(value).strip().replace(",", "").replace("$", "")
    if text.startswith("(") and text.endswith(")"):
        text = "-" + text[1:-1]
    try:
        number = Decimal(text)
    except InvalidOperation as exc:
        raise ValueError(f"Invalid amount at {coordinate}: {value!r}") from exc
    if not number.is_finite():
        raise ValueError(f"Invalid amount at {coordinate}: {value!r}")
    return number


def write_forecast_excel_totals_validation(
    forecast_folder,
    forecast_month,
    fiscal_year,
    output_file=None,
    completed_workbooks=None,
):
    """Put saved forecast Grand Totals into an Excel validation worksheet.

    This step does not open Tableau or read financial statements. Each mapped
    Excel sheet has its own row. The Tableau amount is left empty for the next
    stage. Forecast files are opened for reading and are never saved here.

    completed_workbooks is supplied by the update run to prevent a stale file
    from being reported as freshly updated when its update failed. Omit it to
    collect totals separately from an already completed forecast folder.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    folder = Path(forecast_folder)
    if not folder.is_dir():
        raise ValueError("Forecast folder does not exist.")
    groups = build_workbook_groups(forecast_month=forecast_month)
    report_path = (
        Path(output_file)
        if output_file
        else folder
        / (f"Forecast_Workbook_Validation_{fiscal_year}_{forecast_month}.xlsx")
    )
    if report_path.resolve() in {(folder / name).resolve() for name in groups}:
        raise ValueError("The validation report cannot overwrite a forecast workbook.")
    ready = None if completed_workbooks is None else set(completed_workbooks)
    records = []

    for name, routes in groups.items():
        path = folder / name
        # Deduplicate sheets when several Tableau routes feed the same Excel tab.
        sheet_routes = {}
        for route in routes:
            sheet_routes.setdefault(route["excel_sheet"], []).append(
                route["tableau_worksheet"]
            )
        workbook = None
        try:
            if ready is not None and name not in ready:
                raise ValueError("Workbook was not successfully updated in this run.")
            if not path.is_file():
                raise FileNotFoundError("Mapped forecast workbook was not found.")
            workbook = load_workbook(path, data_only=False)
            for configured_sheet, tableau_tabs in sheet_routes.items():
                record = {
                    "workbook": name,
                    "excel_sheet": configured_sheet,
                    "excel_total": None,
                    "total_cell": "",
                    "note": "",
                    "tableau_route": "; ".join(dict.fromkeys(tableau_tabs)),
                }
                try:
                    worksheet = get_configured_excel_sheet(workbook, configured_sheet)
                    record["excel_sheet"] = worksheet.title
                    structure = get_forecast_structure(worksheet, fiscal_year)
                    header = structure["header_row"]
                    type_column = structure["forecast_type_column"]
                    detail_rows = [
                        row
                        for row in range(header + 1, worksheet.max_row + 1)
                        if worksheet.cell(row, type_column).value not in (None, "")
                        and normalize_text(worksheet.cell(row, type_column).value)
                        != normalize_text("Grand Total")
                    ]
                    if not detail_rows:
                        raise ValueError("No forecast detail rows were found.")
                    total_column = structure.get("grand_total_column")
                    if total_column is None:
                        raise ValueError("No Grand Total column was found.")
                    # finalize_forecast_sheet places the unlabelled total just
                    # below the last detail row. No Grand Total text is needed.
                    address = f"{get_column_letter(total_column)}{max(detail_rows) + 1}"
                    record["total_cell"] = address
                    if worksheet[address].value in (None, ""):
                        raise ValueError(f"The Grand Total cell {address} is blank.")
                    total = _forecast_total_cell_value(worksheet, address)
                    record["excel_total"] = float(total)
                except Exception as exc:
                    record["note"] = explain(exc)
                records.append(record)
        except Exception as exc:
            for configured_sheet, tableau_tabs in sheet_routes.items():
                records.append(
                    {
                        "workbook": name,
                        "excel_sheet": configured_sheet,
                        "excel_total": None,
                        "total_cell": "",
                        "tableau_route": "; ".join(dict.fromkeys(tableau_tabs)),
                        "note": explain(exc),
                    }
                )
        finally:
            if workbook is not None:
                workbook.close()

    report = Workbook()
    sheet = report.active
    sheet.title = "Forecast Validation"
    sheet.append(
        ["FORECAST WORKBOOK VALIDATION", str(fiscal_year), str(forecast_month)]
    )
    sheet.append(["Excel totals collected only. Tableau comparison is pending."])
    sheet.append(["Amounts are in forecast workbook dollars; no $000 conversion."])
    sheet.append([])
    sheet.append(
        [
            "Forecast Workbook",
            "Excel Worksheet",
            "Excel Grand Total ($)",
            "Individual Tableau Grand Total ($)",
            "Variance ($)",
            "Status",
            "Excel Total Cell",
            "Current Forecast Route",
            "Exception / Note",
        ]
    )
    for record in records:
        row = sheet.max_row + 1
        sheet.append(
            [
                record["workbook"],
                record["excel_sheet"],
                record["excel_total"],
                None,
                None,
                None,
                record["total_cell"],
                record["tableau_route"],
                record["note"],
            ]
        )
        # Empty Tableau totals must never be interpreted as zero or as PASS.
        sheet.cell(
            row, 5, f'=IF(AND(ISNUMBER(C{row}),ISNUMBER(D{row})),C{row}-D{row},"")'
        )
        sheet.cell(
            row,
            6,
            (
                f'=IF(ISNUMBER(C{row}),IF(ISNUMBER(D{row}),IF(E{row}=0,"PASS","FAIL"),'
                '"PENDING TABLEAU"),"ACTION REQUIRED")'
            ),
        )
        for column in (3, 4, 5):
            sheet.cell(row, column).number_format = "#,##0.00;[Red](#,##0.00)"
        for column in range(1, 10):
            cell = sheet.cell(row, column)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if column == 4:
                cell.fill = PatternFill("solid", fgColor="E2F0D9")
        sheet.row_dimensions[row].height = 42
    for cell in sheet[5]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="703070")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[5].height = 32
    for column, width in {
        "A": 56,
        "B": 30,
        "C": 24,
        "D": 28,
        "E": 22,
        "F": 24,
        "G": 18,
        "H": 44,
        "I": 60,
    }.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "C6"
    sheet.auto_filter.ref = f"A5:I{sheet.max_row}"
    report.calculation.fullCalcOnLoad = True
    report.calculation.forceFullCalc = True
    report.calculation.calcMode = "auto"
    report_path.parent.mkdir(parents=True, exist_ok=True)
    report.save(report_path)
    report.close()
    return {
        "path": str(report_path),
        "totals_collected": sum(item["excel_total"] is not None for item in records),
        "action_required": sum(item["excel_total"] is None for item in records),
        "tableau_status": "NOT RUN",
    }


def _individual_tableau_grand_total(rows):
    """Read the displayed bottom-right Grand Total, without summing detail rows."""
    from decimal import Decimal, InvalidOperation

    totals = []
    for row in rows:
        total_fields = [k for k in row if normalize_text(k) == "grand total"]
        labels = [v for k, v in row.items() if k not in total_fields]
        # Subtotals also contain Total in the account fields. Only an
        # explicit Grand Total label identifies the overall crosstab total.
        if not any(normalize_text(v) == "grand total" for v in labels):
            continue
        if len(total_fields) != 1:
            raise ValueError(
                "Tableau bottom total must contain one Grand Total column."
            )
        value = row[total_fields[0]]
        if value is None or isinstance(value, bool) or str(value).strip() == "":
            raise ValueError("Tableau Grand Total is blank or invalid.")
        text = str(value).strip().replace(",", "").replace("$", "")
        if text.startswith("(") and text.endswith(")"):
            text = "-" + text[1:-1]
        try:
            number = Decimal(text)
        except InvalidOperation as exc:
            raise ValueError(f"Invalid Tableau Grand Total: {value!r}") from exc
        if not number.is_finite():
            raise ValueError("Tableau Grand Total is not a finite number.")
        totals.append(number)
    if len(totals) != 1:
        raise ValueError("Expected exactly one Tableau bottom Grand Total row.")
    return totals[0]


def _prepare_individual_forecast_copy(
    source, destination, tab, company, forecast_start
):
    """Set the period on a copy; retain saved category and account filters."""
    import xml.etree.ElementTree as ET

    tree = ET.parse(source)
    root = tree.getroot()

    def parameter_key(value):
        return re.sub(r"[^a-z0-9]", "", str(value or "").casefold())

    # Tableau versions may place declarations outside the top-level datasource
    # or save captions with spaces/underscores. Match both caption and name.
    parameters = [
        c
        for c in root.iter("column")
        if any(
            parameter_key(c.get(field)) == "forecaststart"
            for field in ("caption", "name")
        )
    ]
    declarations = [c for c in parameters if "param-domain-type" in c.attrib]
    if declarations:
        parameters = declarations
    target_sheets = [
        sheet
        for sheet in root.findall("./worksheets/worksheet")
        if sheet.get("name") == tab
    ]
    if len(target_sheets) != 1:
        raise ValueError(f"Individual Tableau worksheet not found: {tab}")
    # Resolve duplicate captions through the target worksheet's dependencies.
    target_names = {
        c.get("name")
        for c in target_sheets[0].iter("column")
        if parameter_key(c.get("caption")) == "forecaststart"
    }
    if target_names:
        parameters = [c for c in parameters if c.get("name") in target_names]
    identities = {c.get("name") or c.get("caption") for c in parameters}
    if not parameters or len(identities) != 1:
        available = sorted(
            {
                c.get("caption") or c.get("name") or "<unnamed>"
                for c in root.iter("column")
                if "param-domain-type" in c.attrib
            }
        )
        raise ValueError(
            "Cannot identify Forecast Start parameter. Available parameters: "
            + (", ".join(available) or "none found")
            + ". Share the individual TWB file if this name differs."
        )
    value = f"#{forecast_start:%Y-%m-%d}#"
    for parameter in parameters:
        parameter.set("value", value)
        calculation = parameter.find("calculation")
        if calculation is not None:
            calculation.set("formula", value)
        members = parameter.find("members")
        if members is not None and not any(m.get("value") == value for m in members):
            ET.SubElement(members, "member", value=value)
    sheets = [s for s in root.findall("./worksheets/worksheet") if s.get("name") == tab]
    if len(sheets) != 1:
        raise ValueError(f"Individual Tableau worksheet not found: {tab}")
    filters = sheets[0].findall(".//filter")
    company_filters = [
        f for f in filters if f.get("column", "").endswith("[none:Company:nk]")
    ]
    if len(company_filters) != 1:
        raise ValueError("Expected one Company filter on the individual forecast tab.")
    company_filter = company_filters[0]
    if company_filter.get("class") != "categorical":
        raise ValueError("Unsupported Company filter type on individual forecast tab.")
    # Shared HR/Legal workbooks validate each company on its own temporary copy.
    for child in list(company_filter):
        company_filter.remove(child)
    ET.SubElement(
        company_filter,
        "groupfilter",
        function="member",
        level="[none:Company:nk]",
        member=f'"{company}"',
    )
    forecast_filters = [
        f for f in filters if f.get("column", "").endswith("[none:Forecast Type:nk]")
    ]
    if len(forecast_filters) > 1:
        raise ValueError(
            "Multiple Forecast Type filters on the individual forecast tab."
        )
    if forecast_filters:
        selected = forecast_filters[0].findall(".//groupfilter[@function='member']")
        forecast_members = [
            m
            for m in selected
            if re.fullmatch(
                r"FY\d{2} Forecast_\d{6}", m.get("member", "").strip('"'), re.I
            )
        ]
        if len(forecast_members) != 1:
            raise ValueError(
                "Expected one monthly Forecast Type selection in the saved tab."
            )
        fiscal_end = forecast_start.year + (1 if forecast_start.month >= 7 else 0)
        forecast_members[0].set(
            "member", f'"FY{str(fiscal_end)[-2:]} Forecast_{forecast_start:%m%Y}"'
        )
    else:
        # ALIV Monthly Forecast uses Forecast Measure in its calculation,
        # with actuals before Forecast Start and forecasts after that date.
        measure_names = {
            c.get("name")
            for c in sheets[0].iter("column")
            if parameter_key(c.get("caption")) == "forecastmeasure"
        }
        if len(measure_names) != 1:
            raise ValueError(
                "No Forecast Type filter or unique Forecast Measure parameter on the forecast tab."
            )
        measure_name = next(iter(measure_names))
        calculations = [c.get("formula", "") for c in sheets[0].iter("calculation")]
        if not any(
            measure_name in formula and "[Forecast]" in formula
            for formula in calculations
        ):
            raise ValueError(
                "Forecast Measure is not linked to the worksheet forecast calculation."
            )
        for column in root.iter("column"):
            if (
                column.get("name") == measure_name
                and "param-domain-type" in column.attrib
            ):
                column.set("value", '"Forecast"')
                calculation = column.find("calculation")
                if calculation is not None:
                    calculation.set("formula", '"Forecast"')
    # Keep Non Forecast and the other saved selections: actual months are part
    # of the full-year total shown by the department's Monthly Forecast tab.
    tree.write(destination, encoding="utf-8", xml_declaration=True)


def _append_crosstab_audit(
    audit_path,
    rows,
    source,
    opened_copy,
    tab,
    company,
    excel_file,
    excel_sheet,
    sequence,
):
    """Save each exact parsed capture in its own tab within one run workbook."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    audit_path = Path(audit_path)
    audit = load_workbook(audit_path) if audit_path.exists() else Workbook()
    if not audit_path.exists():
        audit.active.title = "Capture Index"
        audit.active.append(
            [
                "Capture Sheet",
                "Source Tableau Workbook",
                "Opened Copy",
                "Tableau Tab",
                "Company",
                "Forecast Workbook",
                "Excel Worksheet",
                "Captured At",
                "Rows",
            ]
        )
    title = f"{sequence:03d}_{company}_{Path(source).stem}"
    title = re.sub(r"[\\/*?:\[\]]", "_", title)[:31]
    if title in audit.sheetnames:
        raise ValueError(f"Duplicate audit sheet: {title}")
    ws = audit.create_sheet(title)
    ws.append(["Source Tableau workbook", str(source)])
    ws.append(["Opened Tableau copy", str(opened_copy)])
    ws.append(["Tableau worksheet", tab])
    ws.append(["Company", company])
    ws.append(["Forecast workbook / worksheet", excel_file, excel_sheet])
    ws.append(["Captured at", datetime.now().isoformat(timespec="seconds")])
    ws.append([])
    headers = list(dict.fromkeys(k for row in rows for k in row))
    ws.append(headers)
    for row in rows:
        ws.append([row.get(header) for header in headers])
    for cell in ws[8]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="703070")
    ws.freeze_panes = "E9"
    if headers:
        ws.auto_filter.ref = f"A8:{get_column_letter(len(headers))}{ws.max_row}"
    for column in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(column)].width = 24
    # Captured strings are evidence, including text beginning with '='.
    for cells in ws.iter_rows():
        for cell in cells:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                cell.data_type = "s"
    audit["Capture Index"].append(
        [
            title,
            str(source),
            str(opened_copy),
            tab,
            company,
            excel_file,
            excel_sheet,
            datetime.now().isoformat(timespec="seconds"),
            len(rows),
        ]
    )
    audit["Capture Index"].freeze_panes = "A2"
    audit.save(audit_path)
    audit.close()
    return title


def populate_individual_forecast_totals(
    report_path,
    forecast_month,
    fiscal_year,
    capture=None,
    opener=None,
    closer=None,
    open_wait_seconds=8,
    raw_output_folder=None,
):
    """Fill the Tableau column for mapped rows, making extraction errors visible."""
    from datetime import date
    from decimal import Decimal
    from uuid import uuid4
    import calendar
    import time

    if capture is None:
        from automations.actuals.tableau_desktop_export_service import (
            get_tableau_crosstab_rows,
        )

        capture = get_tableau_crosstab_rows
    if opener is None:
        from automations.forecast.tableau_service import open_tableau_workbook

        opener = open_tableau_workbook
    if closer is None:
        from automations.actuals.tableau_desktop_close_service import (
            close_tableau_forecast_workbook,
        )

        closer = close_tableau_forecast_workbook
    month = next(
        (
            m
            for m in range(1, 13)
            if calendar.month_abbr[m].lower() == str(forecast_month)[:3].lower()
        ),
        None,
    )
    if month is None:
        raise ValueError("Unrecognized forecast month.")
    fiscal_end = parse_fiscal_year(fiscal_year)
    start = date(fiscal_end - (1 if month >= 7 else 0), month, 1)
    groups = build_workbook_groups(forecast_month)
    lookup = {}
    for filename, routes in groups.items():
        for route in routes:
            if route.get("individual_tableau_workbook"):
                key = (filename, route["excel_sheet"])
                if key in lookup:
                    raise ValueError(f"Duplicate individual Tableau mapping for {key}.")
                lookup[key] = route
    report_path = Path(report_path)
    book = load_workbook(report_path)
    sheet = book["Forecast Validation"]
    raw_base = (
        Path(raw_output_folder)
        if raw_output_folder
        else report_path.parent / "Tableau_Raw_Copies"
    )
    raw = raw_base / f"Forecast_Validation_{uuid4().hex[:8]}"
    raw.mkdir(parents=True, exist_ok=True)
    audit_path = raw / "Individual_Tableau_Crosstabs.xlsx"
    counts = {"passed": 0, "failed": 0, "action_required": 0, "pending": 0}
    try:
        sheet.cell(
            2, 1, "Grand Total comparison against mapped individual Tableau workbooks."
        )
        sheet.cell(5, 10, "Individual Tableau Workbook")
        sheet.cell(5, 11, "Individual Tableau Worksheet")
        for row_number in range(6, sheet.max_row + 1):
            route = lookup.get(
                (sheet.cell(row_number, 1).value, sheet.cell(row_number, 2).value)
            )
            sheet.cell(row_number, 4).value = None
            if route is None:
                counts["pending"] += 1
                continue
            temporary = None
            opened = False
            try:
                if not isinstance(sheet.cell(row_number, 3).value, (int, float)):
                    raise ValueError(
                        "Excel Grand Total unavailable; Tableau comparison was not run."
                    )
                source = (
                    Path(route["individual_tableau_folder"])
                    / route["individual_tableau_workbook"]
                )
                tab = route["individual_tableau_worksheet"]
                company = (
                    route.get("company_filter") or route["tableau_worksheet"].split()[0]
                )
                if company not in ("ALIV", "CBL"):
                    raise ValueError("Individual Tableau company mapping is required.")
                sheet.cell(row_number, 10, str(source))
                sheet.cell(row_number, 11, tab)
                if not source.is_file() or source.suffix.lower() != ".twb":
                    raise FileNotFoundError(
                        f"Individual Tableau TWB not found: {source}"
                    )
                # Same directory keeps any relative data-source paths valid.
                temporary = source.with_name(
                    f"{source.stem}_Validation_{uuid4().hex[:8]}.twb"
                )
                _prepare_individual_forecast_copy(
                    source, temporary, tab, company, start
                )
                opener(tableau_workbook_path=str(temporary))
                opened = True
                if open_wait_seconds:
                    time.sleep(open_wait_seconds)
                rows = capture(
                    worksheet_name=tab,
                    raw_output_folder=raw,
                    sequence_number=row_number,
                    tableau_workbook_path=str(temporary),
                )
                audit_sheet = _append_crosstab_audit(
                    audit_path,
                    rows,
                    source,
                    temporary,
                    tab,
                    company,
                    sheet.cell(row_number, 1).value,
                    sheet.cell(row_number, 2).value,
                    row_number,
                )
                sheet.cell(row_number, 12, audit_sheet)
                total = _individual_tableau_grand_total(rows)
                # Mirror the approved Excel revenue sign convention.
                if "REVENUE" in route["workbook"].upper():
                    total = -total
                sheet.cell(row_number, 4, float(total))
                variance = Decimal(str(sheet.cell(row_number, 3).value)) - total
                sheet.cell(row_number, 5, float(variance))
                status = "PASS" if variance == 0 else "FAIL"
                sheet.cell(row_number, 6, status)
                sheet.cell(
                    row_number,
                    9,
                    (
                        "Grand Totals agree."
                        if status == "PASS"
                        else "Grand Totals differ."
                    ),
                )
                counts["passed" if status == "PASS" else "failed"] += 1
            except Exception as exc:
                sheet.cell(row_number, 4).value = None
                sheet.cell(row_number, 5).value = None
                sheet.cell(row_number, 6, "ACTION REQUIRED")
                sheet.cell(row_number, 9, explain(exc))
                counts["action_required"] += 1
            finally:
                close_ok = not opened
                if opened:
                    try:
                        closer(tableau_workbook_path=str(temporary))
                        close_ok = True
                    except Exception as exc:
                        sheet.cell(
                            row_number,
                            9,
                            str(sheet.cell(row_number, 9).value)
                            + f" Close warning: {exc}",
                        )
                if temporary is not None and close_ok:
                    temporary.unlink(missing_ok=True)
        sheet.cell(5, 12, "Captured Crosstab Sheet")
        sheet.auto_filter.ref = f"A5:L{sheet.max_row}"
        for col in (10, 11, 12):
            from copy import copy as clone

            sheet.cell(5, col)._style = clone(sheet.cell(5, 8)._style)
            sheet.column_dimensions[get_column_letter(col)].width = (
                48 if col == 10 else 28
            )
        book.save(report_path)
    finally:
        book.close()
    return {
        "path": str(report_path),
        "tableau_status": (
            "REVIEW REQUIRED"
            if counts["failed"] or counts["action_required"] or counts["pending"]
            else "PASS"
        ),
        "individual_raw_output_folder": str(raw),
        "individual_crosstab_audit": str(audit_path),
        **counts,
    }


def run_standalone_forecast_folder(
    forecast_folder,
    output_folder,
    forecast_month,
    fiscal_year,
    tableau_workbook_path=None,
    validation_base_folder=None,
):

    from automations.actuals.tableau_desktop_export_service import (
        get_tableau_crosstab_rows,
    )

    # Validation follows the selected forecast month in the shared FY tree.
    from automations.actuals.validation_output_folders import (
        create_validation_output_folders,
        find_actuals_validation_results_folder,
    )
    from datetime import date
    import calendar

    validation_month = next(
        (
            m
            for m in range(1, 13)
            if calendar.month_abbr[m].lower() == str(forecast_month)[:3].lower()
        ),
        None,
    )
    if validation_month is None:
        raise ValueError("Unrecognized forecast month.")
    validation_fiscal_end = parse_fiscal_year(fiscal_year)
    validation_period = date(
        validation_fiscal_end - (1 if validation_month >= 7 else 0), validation_month, 1
    )
    if validation_base_folder is None:
        validation_base_folder = find_actuals_validation_results_folder()
    if (
        not str(validation_base_folder).strip()
        or not Path(validation_base_folder).is_dir()
    ):
        raise ValueError(
            "Select the synced shared validation root used for Actuals Validation."
        )
    validation_paths = create_validation_output_folders(
        validation_base_folder, validation_period
    )
    validation_report_path = validation_paths["forecast_validation"] / (
        f"Forecast_Workbook_Validation_{fiscal_year}_{forecast_month}.xlsx"
    )

    fiscal_end_year = parse_fiscal_year(fiscal_year)

    fiscal_headers = build_fiscal_month_headers(fiscal_year)

    forecast_folder = Path(forecast_folder)

    output_folder = Path(output_folder)

    if not forecast_folder.exists():
        raise RuntimeError(
            "Forecast workbook folder was not found:\n" f"{forecast_folder}"
        )

    output_folder.mkdir(
        parents=True,
        exist_ok=True,
    )

    log_folders = get_forecast_automation_log_folders()
    raw_output_folder = log_folders["raw"]
    run_log_path = (
        log_folders["runs"]
        / f"forecast_run_{datetime.now().strftime('%Y%m%d_%H%M%S')}.txt"
    )

    workbook_groups = build_workbook_groups(
        forecast_month=forecast_month,
    )

    expected_workbook_names = set(workbook_groups.keys())

    forecast_files = sorted(
        [
            file
            for file in forecast_folder.glob("*.xlsx")
            if (not file.name.startswith("~$") and file.name in expected_workbook_names)
        ],
        key=lambda path: path.name.lower(),
    )

    existing_workbook_names = {file.name for file in forecast_files}

    missing_workbooks = sorted(expected_workbook_names - existing_workbook_names)

    # ========================================================
    # STARTUP DIAGNOSTICS
    # ========================================================

    print()
    print("=" * 70)
    print("MONTHLY TABLEAU FORECAST UPDATE")
    print("=" * 70)

    print(f"Forecast workbook folder : " f"{forecast_folder}")

    print(f"Completed output folder   : " f"{output_folder}")

    print(f"Tableau workbook          : " f"{tableau_workbook_path}")

    print(f"Forecast month            : " f"{forecast_month}")

    print(f"Fiscal year               : " f"{fiscal_year}")

    print(f"Fiscal end year           : " f"{fiscal_end_year}")

    print(
        "Fiscal calendar           : "
        + ", ".join(str(fiscal_headers[month]) for month in FISCAL_MONTH_ORDER)
    )

    print(f"Forecast workbooks found  : " f"{len(forecast_files)}")

    print("=" * 70)

    # ========================================================
    # RESULT COLLECTIONS
    # ========================================================

    updated = []

    failed = []

    missing = []

    manual_required = []

    for workbook_name in missing_workbooks:

        missing.append(
            {
                "workbook": workbook_name,
                "error": (
                    "Mapped current-month forecast workbook "
                    "was not found in the forecast workbook folder."
                ),
            }
        )

    sequence_number = 1

    # ========================================================
    # PROCESS EACH FORECAST WORKBOOK
    # ========================================================

    for workbook_index, source_path in enumerate(
        forecast_files,
        start=1,
    ):

        print()
        print("=" * 70)

        print(f"[{workbook_index}/{len(forecast_files)}] " f"{source_path.name}")

        print("=" * 70)

        routes = workbook_groups.get(source_path.name)

        if not routes:

            # This should normally never occur because forecast_files
            # was already restricted to mapped workbook names.
            failed.append(
                {
                    "workbook": source_path.name,
                    "error": ("No Tableau route is configured " "for this workbook."),
                }
            )

            continue

        workbook = None

        try:

            # =================================================
            # OPEN THE NEWLY CREATED FORECAST WORKBOOK
            # =================================================

            workbook = load_workbook(
                source_path,
                data_only=False,
            )

            # =================================================
            # IDENTIFY UNIQUE TABLEAU WORKSHEETS
            # =================================================

            unique_tableau_names = []

            for route in routes:

                tableau_name = route["tableau_worksheet"]

                if tableau_name not in unique_tableau_names:

                    unique_tableau_names.append(tableau_name)

            # =================================================
            # EXPORT TABLEAU CROSSTABS
            # =================================================

            extracted = {}

            for tableau_name in unique_tableau_names:

                print()
                print(f"Reading Tableau: " f"{tableau_name}")

                extracted[tableau_name] = get_tableau_crosstab_rows(
                    worksheet_name=tableau_name,
                    raw_output_folder=(raw_output_folder),
                    sequence_number=(sequence_number),
                    tableau_workbook_path=(tableau_workbook_path),
                )

                sequence_number += 1
            # =================================================
            # HANDLE COMPANY-SPECIFIC ROUTES
            # ========================================================

            special_tableau_names = {
                route["tableau_worksheet"]
                for route in routes
                if route.get("company_filter")
            }

            skip_workbook = False

            for tableau_name in special_tableau_names:

                requested_companies = {
                    normalize_text(route["company_filter"])
                    for route in routes
                    if (
                        route["tableau_worksheet"] == tableau_name
                        and route.get("company_filter")
                    )
                }

                available_companies = {
                    normalize_text(get_tableau_company_value(row))
                    for row in (get_tableau_detail_rows(extracted[tableau_name]))
                    if normalize_text(get_tableau_company_value(row))
                }

                missing_companies = requested_companies - available_companies

                if missing_companies:

                    message = (
                        f"Tableau worksheet "
                        f'"{tableau_name}" '
                        "does not currently contain "
                        "all required company data. "
                        f"Available="
                        f"{sorted(available_companies)}; "
                        f"Required but missing="
                        f"{sorted(missing_companies)}."
                    )

                    print("[MANUAL FILTER REQUIRED] " + message)

                    manual_required.append(
                        {
                            "workbook": (source_path.name),
                            "tableau_worksheet": (tableau_name),
                            "error": message,
                        }
                    )

                    skip_workbook = True

            if skip_workbook:

                workbook.close()

                workbook = None

                continue

            # =================================================
            # ROUTE TABLEAU DATA TO THE CORRECT EXCEL SHEETS
            # =================================================

            datasets_by_excel_sheet = defaultdict(list)

            for route in routes:

                tableau_name = route["tableau_worksheet"]

                tableau_rows = extracted[tableau_name]

                company_filter = route.get("company_filter")

                if company_filter:

                    tableau_rows = filter_rows_for_company(
                        tableau_rows=(tableau_rows),
                        company=(company_filter),
                    )

                    if not tableau_rows:

                        raise RuntimeError(
                            "No Tableau rows remained "
                            "after applying "
                            f'Company="{company_filter}" '
                            f'for "{tableau_name}".'
                        )

                target_sheet = get_configured_excel_sheet(
                    workbook=workbook,
                    configured_sheet_name=(route["excel_sheet"]),
                )

                print(
                    f'Tableau "{tableau_name}" ' f"-> Excel " f'"{target_sheet.title}"'
                )

                datasets_by_excel_sheet[target_sheet.title].append(
                    (
                        tableau_name,
                        tableau_rows,
                    )
                )

            # =================================================
            # VALIDATE EACH EXCEL SHEET STRUCTURE
            # =================================================

            structures_by_sheet = {}

            for sheet_name in datasets_by_excel_sheet:

                worksheet = workbook[sheet_name]

                structure = get_forecast_structure(
                    worksheet=worksheet,
                    fiscal_year=fiscal_year,
                )

                structures_by_sheet[sheet_name] = structure

                print(
                    f'Validated "{sheet_name}": '
                    f"header row="
                    f'{structure["header_row"]}'
                )

            # =================================================
            # REBUILD EACH FORECAST SHEET FROM TABLEAU
            # =================================================

            workbook_sheet_count = 0

            workbook_tableau_rows = 0

            workbook_values_written = 0

            for (
                sheet_name,
                sheet_datasets,
            ) in datasets_by_excel_sheet.items():

                worksheet = workbook[sheet_name]

                structure = structures_by_sheet[sheet_name]

                print()
                print("Rebuilding Excel sheet " f'from Tableau: "{sheet_name}"')

                # ---------------------------------------------
                # Tableau becomes the source of truth.
                #
                # The function:
                # - keeps Excel headers
                # - keeps Forecast Type
                # - resizes the detail region
                # - writes all Tableau detail rows
                # - writes Jul-Jun values
                # - rebuilds Grand Total formulas
                # ---------------------------------------------

                update_result = update_existing_forecast_rows(
                    worksheet=worksheet,
                    tableau_datasets=(sheet_datasets),
                    structure=structure,
                    fiscal_year=fiscal_year,
                    workbook_name=source_path.name,
                )

                # ---------------------------------------------
                # Convert fiscal headers to real Excel dates and
                # apply the approved workbook presentation.
                # ---------------------------------------------

                finalize_forecast_sheet(
                    worksheet=worksheet,
                    structure=structure,
                    fiscal_year=fiscal_year,
                )

                workbook_tableau_rows += update_result["matched_rows"]

                workbook_values_written += update_result["values_updated"]

                workbook_sheet_count += 1

            # =================================================
            # FORCE EXCEL RECALCULATION WHEN OPENED
            # =================================================

            try:

                workbook.calculation.fullCalcOnLoad = True

                workbook.calculation.forceFullCalc = True

                workbook.calculation.calcMode = "auto"

            except Exception:

                # Recalculation properties can vary by
                # openpyxl/workbook version. Failure here should
                # not invalidate an otherwise completed workbook.
                pass

            # =================================================
            # SAVE COMPLETED WORKBOOK
            # =================================================

            output_path = output_folder / source_path.name

            workbook.save(output_path)

            workbook.close()

            workbook = None

            print()
            print("[SUCCESS]")

            print(f"Excel sheets rebuilt: " f"{workbook_sheet_count}")

            print(f"Tableau detail rows written: " f"{workbook_tableau_rows}")

            print(f"Values written: " f"{workbook_values_written}")

            print(f"Saved: " f"{output_path}")

            updated.append(
                {
                    "workbook": (source_path.name),
                    "sheets_updated": (workbook_sheet_count),
                    "rows": (workbook_tableau_rows),
                    "values_updated": (workbook_values_written),
                    "tableau_worksheets": (unique_tableau_names),
                    "output_path": str(output_path),
                }
            )

        except Exception as exc:

            # =================================================
            # ALWAYS CLOSE THE WORKBOOK AFTER A FAILURE
            # =================================================

            if workbook is not None:

                try:
                    workbook.close()

                except Exception:
                    pass

            print()
            print(f"[FAILED] " f"{source_path.name}")

            printexplain(exc)

            failed.append(
                {
                    "workbook": (source_path.name),
                    "error": explain(exc),
                }
            )

    # ========================================================
    # FINAL SUMMARY
    # ========================================================

    print()
    print("=" * 70)
    print("MONTHLY FORECAST UPDATE SUMMARY")
    print("=" * 70)

    print(f"Workbooks found:       " f"{len(forecast_files)}")

    print(f"Successfully updated:  " f"{len(updated)}")

    print(f"Missing workbooks:      " f"{len(missing)}")

    print(f"Manual filter needed:  " f"{len(manual_required)}")

    print(f"Failed:                 " f"{len(failed)}")

    print("=" * 70)

    result = {
        "files_found": len(forecast_files),
        "updated": updated,
        "missing": missing,
        "manual_required": (manual_required),
        "failed": failed,
        "output_folder": str(output_folder),
        "fiscal_year": str(fiscal_year),
        "fiscal_end_year": (fiscal_end_year),
        "run_log": str(run_log_path),
        "raw_output_folder": str(raw_output_folder),
    }

    # Collect saved Excel totals, then compare the mapped individual Tableau totals.
    try:
        result["forecast_validation"] = write_forecast_excel_totals_validation(
            forecast_folder=output_folder,
            forecast_month=forecast_month,
            fiscal_year=fiscal_year,
            completed_workbooks=[item["workbook"] for item in updated],
            output_file=validation_report_path,
        )
        individual_result = populate_individual_forecast_totals(
            result["forecast_validation"]["path"],
            forecast_month,
            fiscal_year,
            raw_output_folder=validation_paths["tableau_raw_copies"],
        )
        result["forecast_validation"].update(individual_result)
        print("Forecast totals validation: " + result["forecast_validation"]["path"])
    except Exception as exc:
        # Keep workbook-update success separate from a failed report export.
        result["forecast_validation"] = {
            "error": explain(exc),
            "tableau_status": "NOT RUN",
        }
        print("[VALIDATION REPORT ERROR] " + result["forecast_validation"]["error"])

    write_run_summary_log(
        run_log_path,
        result,
    )

    print(f"Run log: " f"{run_log_path}")

    print(f"Raw Tableau copies: " f"{raw_output_folder}")

    return result
