"""Individual crosstab workbook: one tab per Tableau capture plus a Capture Index.

This is where the team looks to see exactly what was copied from Tableau.
Every capture tab starts with a header block that lists the source workbook,
the worksheet and the filters that were set before the crosstab was copied.
"""

from __future__ import annotations

import re
from datetime import datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

INDEX_SHEET = "Capture Index"
PURPLE = "703070"
NAVY = "243F5C"
TINT = "F3EDF6"
BAND = "F7F9FB"
THIN = Side(style="thin", color="D0D7DE")

INDEX_HEADERS = (
    "Capture Sheet",
    "Area",
    "Basis",
    "Company",
    "Source Tableau Workbook",
    "Tableau Worksheet",
    "Filters Used",
    "Captured At",
    "Rows",
)
INDEX_WIDTHS = (30, 12, 10, 10, 48, 48, 70, 20, 8)


def format_filters(filters: dict | None) -> str:
    """One readable line, e.g. 'Forecast Start = 09/01/2026; Fiscal Year = 2027'."""
    if not filters:
        return "No filters recorded"
    return "; ".join(f"{k} = {v}" for k, v in filters.items() if v not in (None, ""))


def _typed(value):
    """Numbers become numbers; everything else stays text (never a formula)."""
    text = str(value if value is not None else "").strip()
    if text and not text.endswith("%") and not text.startswith("="):
        try:
            return float(Decimal(text.replace(",", "").replace("$", "")))
        except Exception:
            pass
    return text


def _table(rows_or_table):
    """Accept a list of dicts (parsed rows) or a list of lists (raw table)."""
    rows = list(rows_or_table)
    if rows and isinstance(rows[0], dict):
        headers = list(dict.fromkeys(k for r in rows for k in r if not str(k).startswith("_")))
        return [headers] + [[r.get(h) for h in headers] for r in rows], 1
    return rows, None


def _new_index(workbook) -> None:
    ws = workbook.active
    ws.title = INDEX_SHEET
    ws.sheet_properties.tabColor = NAVY
    ws.append(list(INDEX_HEADERS))
    for col, width in enumerate(INDEX_WIDTHS, 1):
        cell = ws.cell(1, col)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"


def append_crosstab(
    audit_path,
    rows_or_table,
    *,
    source,
    tab,
    company,
    area,
    basis,
    sequence,
    filters: dict | None = None,
    header_rows: int | None = None,
):
    """Add one capture as its own tab and register it in the Capture Index.

    ``rows_or_table`` is either parsed rows (list of dicts) or the raw table
    (list of lists). ``header_rows`` is how many table rows are headings
    (1 by default; 2 for Tableau's month / measure double heading).
    Returns the new sheet's title.
    """
    table, inferred = _table(rows_or_table)
    if not table:
        raise ValueError("The Tableau crosstab was empty; nothing to save.")
    header_rows = header_rows or inferred or 1

    audit_path = Path(audit_path)
    audit_path.parent.mkdir(parents=True, exist_ok=True)
    existed = audit_path.exists()
    workbook = load_workbook(audit_path) if existed else Workbook()
    try:
        if not existed:
            _new_index(workbook)
        title = re.sub(
            r"[\\/*?:\[\]]", "_", f"{sequence:03d}_{company}_{Path(source).stem}"
        )[:31]
        if title in workbook.sheetnames:
            raise ValueError(f"Duplicate capture sheet: {title}")
        ws = workbook.create_sheet(title)
        ws.sheet_properties.tabColor = PURPLE

        captured_at = datetime.now().isoformat(timespec="seconds")
        filter_text = format_filters(filters)
        info = [
            ("Source Tableau workbook", str(source)),
            ("Tableau worksheet", tab),
            ("Area / Basis / Company", f"{area} / {basis} / {company}"),
            ("Captured at", captured_at),
            ("Filters used", filter_text),
        ]
        for r, (label, value) in enumerate(info, 1):
            ws.cell(r, 1, label).font = Font(bold=True, color="44546A")
            ws.cell(r, 1).fill = PatternFill("solid", fgColor=TINT)
            ws.cell(r, 2, value).alignment = Alignment(wrap_text=False, vertical="center")
        ws.cell(5, 2).font = Font(bold=True, color="1F3FBF")

        first = len(info) + 2  # one blank row after the header block
        for offset, row in enumerate(table):
            for col, value in enumerate(row, 1):
                ws.cell(first + offset, col, _typed(value) if offset >= header_rows else str(value or "").strip())
        for offset in range(header_rows):
            for cell in ws[first + offset]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor=PURPLE)
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        width = max((len(r) for r in table), default=1)
        for col in range(1, width + 1):
            ws.column_dimensions[get_column_letter(col)].width = 27 if col <= 4 else 19
        ws.column_dimensions["A"].width = 27
        last_header = first + header_rows - 1
        ws.freeze_panes = ws.cell(last_header + 1, 5)
        ws.auto_filter.ref = f"A{last_header}:{get_column_letter(width)}{ws.max_row}"
        for row in ws.iter_rows(min_row=last_header + 1):
            for cell in row:
                if isinstance(cell.value, float):
                    cell.number_format = "#,##0.00;[Red](#,##0.00)"
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    cell.data_type = "s"  # captured text is evidence, never a formula

        index = workbook[INDEX_SHEET]
        row = index.max_row + 1
        values = (
            title, area, basis, company, str(source), tab,
            filter_text, captured_at, len(table) - header_rows,
        )
        for col, value in enumerate(values, 1):
            cell = index.cell(row, col, value)
            cell.alignment = Alignment(vertical="center", wrap_text=col in (5, 6, 7))
            cell.border = Border(bottom=THIN)
            if row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor=BAND)
        link = index.cell(row, 1)
        link.hyperlink = f"#'{title}'!A1"
        link.font = Font(color="1F5FA8", underline="single")
        index.auto_filter.ref = f"A1:{get_column_letter(len(INDEX_HEADERS))}{row}"
        workbook.save(audit_path)
        return title
    finally:
        workbook.close()
