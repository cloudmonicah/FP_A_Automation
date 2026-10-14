from __future__ import annotations

from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP
import csv
from pathlib import Path
import time
from typing import Any
import getpass

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from automations.actuals.opex_financial_statement_reader import (
    read_aliv_opex,
)

from automations.actuals.validation_config import (
    OPEX_MAPPINGS,
    OPEX_TABLEAU_WORKBOOK,
    OPEX_TABLEAU_WORKSHEET,
)

# ============================================================
# STATUS CONSTANTS
# ============================================================

STATUS_PASS = "PASS"
STATUS_FAIL = "FAIL"
STATUS_ACTION_REQUIRED = "ACTION REQUIRED"
STATUS_REVIEW_REQUIRED = "REVIEW REQUIRED"

# ============================================================
# WORKBOOK STYLE CONSTANTS
# ============================================================

TITLE_FILL = PatternFill(
    fill_type="solid",
    fgColor="5B1760",
)

HEADER_FILL = PatternFill(
    fill_type="solid",
    fgColor="7A1F73",
)

ACCENT_FILL = PatternFill(
    fill_type="solid",
    fgColor="EADCF0",
)

PASS_FILL = PatternFill(
    fill_type="solid",
    fgColor="E2F0D9",
)

FAIL_FILL = PatternFill(
    fill_type="solid",
    fgColor="F4CCCC",
)

ACTION_FILL = PatternFill(
    fill_type="solid",
    fgColor="FFF2CC",
)
NEUTRAL_FILL = PatternFill(
    fill_type="solid",
    fgColor="F3F3F3",
)

WHITE_FONT = Font(
    color="FFFFFF",
    bold=True,
)
THIN_BORDER = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)

# ============================================================
# SOURCE NORMALIZATION
# ============================================================


def _round_excel(value: float | int | Decimal) -> int:
    """Round to a whole number using Excel-style ROUND_HALF_UP."""
    return int(
        Decimal(str(value)).quantize(
            Decimal("1"),
            rounding=ROUND_HALF_UP,
        )
    )


def _tableau_to_000(value: float | int | Decimal) -> int:
    """
    Tableau stores Actual YTD in whole dollars.

    Example:
        883551 -> 883.551 -> 884
    """
    return _round_excel(Decimal(str(value)) / Decimal("1000"))


def _fs_to_000(value: float | int | Decimal) -> int:
    """
    Financial statements are already stored in $000.

    Example:
        883.551 -> 884
    """
    return _round_excel(value)


def _normalize(value: Any) -> str:
    if value is None:
        return ""
    return " ".join(str(value).strip().upper().split())


def _money(value: Any) -> float | None:
    """
    Convert common Tableau crosstab money text into a float.

    Supports:
        1,234
        $1,234
        (1,234)
        1234.56
    """
    if value is None:
        return None

    if isinstance(value, (int, float)):
        return float(value)

    text = str(value).strip()

    if not text:
        return None

    negative = text.startswith("(") and text.endswith(")")

    text = text.strip("()").replace("$", "").replace(",", "").replace(" ", "")

    try:
        amount = float(text)
    except ValueError:
        return None

    return -amount if negative else amount


# ============================================================
# TABLEAU CROSSTAB COLUMN DETECTION
# ============================================================


def _find_field(
    headers: list[str],
    exact_candidates: list[str],
    contains_candidates: list[str] | None = None,
) -> str | None:
    """
    Resolve a crosstab FIELD/header.

    This is only flexible for header names.
    Business/category mappings remain explicit in validation_config.py.
    """

    normalized = {_normalize(header): header for header in headers}

    for candidate in exact_candidates:
        key = _normalize(candidate)

        if key in normalized:
            return normalized[key]

    contains_candidates = contains_candidates or []

    for header in headers:
        header_norm = _normalize(header)

        for candidate in contains_candidates:
            if _normalize(candidate) in header_norm:
                return header

    return None


def _resolve_tableau_columns(
    tableau_rows: list[dict],
    closed_period: date | datetime | None = None,
) -> dict[
    str,
    str | None,
]:
    """
    Detect the copied Tableau OPEX Crosstab layout.

    CURRENT LIVE TABLEAU OUTPUT
    ---------------------------
    The confirmed OPEX worksheet returns:

        Company
        GL FS Category
        July 2026
        August 2026
        ...
        June 2027
        Grand Total

    Because Company and GL FS Category are explicit fields,
    this is a FLAT layout.

    The closed-period month column is the amount used for the
    actuals validation.
    """

    if not tableau_rows:

        raise RuntimeError("Tableau returned no OPEX rows.")

    headers = list(tableau_rows[0].keys())

    # ========================================================
    # COMPANY FIELD
    # ========================================================

    company_field = _find_field(
        headers,
        exact_candidates=[
            "Company",
        ],
        contains_candidates=[
            "Company",
        ],
    )

    # ========================================================
    # GL FS CATEGORY FIELD
    # ========================================================

    category_field = _find_field(
        headers,
        exact_candidates=[
            "GL FS Category",
            "FS Category",
        ],
        contains_candidates=[
            "FS Category",
        ],
    )

    # ========================================================
    # AMOUNT FIELD
    # ========================================================

    # --------------------------------------------------------
    # Some Tableau views may expose Actual YTD directly.
    # --------------------------------------------------------

    actual_field = _find_field(
        headers,
        exact_candidates=[
            "Actual YTD",
        ],
        contains_candidates=[
            "Actual YTD",
        ],
    )

    # --------------------------------------------------------
    # Current OPEX Tableau view exposes one column per month.
    #
    # For July 2026 validation, use:
    #
    #     July 2026
    #
    # NOT Grand Total.
    # --------------------------------------------------------

    if actual_field is None and closed_period is not None:

        month_header = closed_period.strftime("%B %Y")

        actual_field = _find_field(
            headers,
            exact_candidates=[
                month_header,
            ],
        )

    # ========================================================
    # FALLBACK MONTH RESOLUTION
    # ========================================================

    if actual_field is None:

        month_names = (
            "JANUARY",
            "FEBRUARY",
            "MARCH",
            "APRIL",
            "MAY",
            "JUNE",
            "JULY",
            "AUGUST",
            "SEPTEMBER",
            "OCTOBER",
            "NOVEMBER",
            "DECEMBER",
        )

        for header in headers:

            normalized = _normalize(header)

            # Never use Grand Total for the
            # closed-period comparison.
            if normalized == "GRAND TOTAL":
                continue

            if any(month in normalized for month in month_names):

                actual_field = header

                break

    # ========================================================
    # CURRENT LIVE FLAT LAYOUT
    # ========================================================

    if (
        company_field is not None
        and category_field is not None
        and actual_field is not None
    ):

        return {
            "layout": "FLAT",
            "company_field": (company_field),
            "category_field": (category_field),
            "actual_field": (actual_field),
            "label_field": None,
        }

    # ========================================================
    # LEGACY / DIAGNOSTIC MATRIX FALLBACK
    # ========================================================

    label_field = None

    for header in headers:

        if _normalize(header) == "":

            label_field = header

            break

    if label_field is None:

        label_field = headers[0]

    return {
        "layout": "MATRIX",
        "company_field": None,
        "category_field": None,
        "actual_field": (actual_field),
        "label_field": (label_field),
    }


# ============================================================
# TABLEAU OPEX EXTRACTION
# ============================================================


def build_opex_tableau_values(tableau_rows, closed_period=None, basis="MONTH"):
    """Read category-level monthly actuals without adding subtotal rows twice."""
    if not tableau_rows or closed_period is None:
        raise ValueError("OPEX capture and selected closed period are required.")
    if basis not in {"MONTH", "YTD"}:
        raise ValueError("Expected MONTH or YTD basis.")
    headers = list(tableau_rows[0])

    def field(label):
        matches = [h for h in headers if _normalize(h) == _normalize(label)]
        if len(matches) != 1:
            raise ValueError(f"Expected one OPEX column: {label}")
        return matches[0]

    company_key = field("Company")
    category_key = field("GL FS Category")
    amount_key = field(
        closed_period.strftime("%B %Y") + " Actual YTD"
        if basis == "MONTH"
        else "Grand Total Actual YTD"
    )
    if any(_normalize(h) == "GL ACCOUNT MAPPING" for h in headers):
        raise ValueError(
            "Expected the category-level Forecast Summary worksheet, but received GL mapping details."
        )
    values = {"ALIV": {}, "CBL": {}}
    totals = {}
    company = None
    for row in tableau_rows:
        company_label = _normalize(row.get(company_key))
        category = _normalize(row.get(category_key))
        if company_label == "GRAND TOTAL" or category == "GRAND TOTAL":
            continue
        if company_label in values:
            company = company_label
        elif company_label and company_label not in {
            "TOTAL",
            "ALIV TOTAL",
            "CBL TOTAL",
        }:
            raise ValueError(f"Unexpected OPEX company: {company_label}")
        if company_label in {"ALIV TOTAL", "CBL TOTAL"}:
            company = company_label.split()[0]
        if company not in values:
            raise ValueError("OPEX category appears before a recognized company.")
        amount = _money(row.get(amount_key))
        is_total = category == "TOTAL" or company_label in {
            "TOTAL",
            "ALIV TOTAL",
            "CBL TOTAL",
        }
        if is_total:
            if company in totals or amount is None:
                raise ValueError(f"Expected one numeric {company} Total row.")
            totals[company] = amount
            continue
        if not category or category == "NULL":
            continue
        if category in values[company]:
            raise ValueError(
                f"Duplicate category in summary capture: {company} / {category}"
            )
        # Tableau renders empty category cells for no activity; preserve them as zero.
        if (
            amount is None
            and row.get(amount_key) is not None
            and str(row.get(amount_key)).strip()
        ):
            raise ValueError(f"Invalid {amount_key} amount for {company} / {category}")
        values[company][category] = amount if amount is not None else 0.0
    for company in values:
        if not values[company] or company not in totals:
            raise ValueError(
                f"Missing {company} category rows or company Total in summary capture."
            )
        values[company][f"{company} TOTAL"] = totals[company]
    return values


def _variance_note(
    variance_000: int,
) -> str:
    if variance_000 == 0:
        return "Reconciled."

    amount = abs(variance_000)

    direction = "above" if variance_000 > 0 else "below"

    return f"Tableau is ${amount:,}K " f"{direction} Financial Statement."


def _canonical_fs_line(value: Any) -> str:
    """
    Canonicalize Financial Statement labels for explicit matching.

    This is intentionally limited to harmless naming differences such as
    "&" versus "and". It does not perform fuzzy matching.
    """

    if value is None:
        return ""

    text = str(value).strip().upper()
    text = text.replace("&", " AND ")
    return " ".join(text.split())


def _resolve_fs_value(
    *,
    fs_line: str,
    fs_values: dict[str, float],
) -> float | None:

    if fs_line in fs_values:
        return fs_values[fs_line]

    target = _canonical_fs_line(fs_line)

    for extracted_line, value in fs_values.items():
        if _canonical_fs_line(extracted_line) == target:
            return value

    return None


def compare_opex_company(
    *,
    company: str,
    fs_values: dict[str, float],
    tableau_values: dict[str, float],
) -> dict[str, Any]:

    details: list[dict[str, Any]] = []

    for mapping in OPEX_MAPPINGS:

        if _normalize(mapping.get("company")) != company:
            continue

        fs_line = mapping["fs_line"]
        tableau_category = mapping["tableau_category"]

        # ----------------------------------------------------
        # 1. RESOLVE FINANCIAL STATEMENT VALUE FIRST
        # ----------------------------------------------------

        fs_raw = _resolve_fs_value(
            fs_line=fs_line,
            fs_values=fs_values,
        )

        fs_000 = _fs_to_000(fs_raw) if fs_raw is not None else None

        # ----------------------------------------------------
        # 2. RESOLVE TABLEAU VALUE
        # ----------------------------------------------------
        tableau_raw = tableau_values.get(_normalize(tableau_category))

        result = {
            "company": company,
            "statement_line": fs_line,
            "tableau_category": tableau_category,
            "fs_raw_000": fs_raw,
            "fs_rounded_000": fs_000,
            "tableau_raw_dollars": tableau_raw,
            "tableau_rounded_000": None,
            "variance_000": None,
            "status": None,
            "note": "",
        }

        # ----------------------------------------------------
        # 3. FINANCIAL STATEMENT SOURCE CHECK
        # ----------------------------------------------------
        if fs_raw is None:
            result["status"] = STATUS_ACTION_REQUIRED
            result["note"] = (
                "Financial Statement line was not found in the "
                "extracted finalized statement."
            )
            details.append(result)
            continue

        # ----------------------------------------------------
        # 4. TABLEAU SOURCE CHECK
        # ----------------------------------------------------
        # Financial Statement amount remains populated here.
        if tableau_raw is None:
            result["status"] = STATUS_ACTION_REQUIRED
            result["note"] = (
                "Mapped Tableau category " f"'{tableau_category}' was not found."
            )
            details.append(result)
            continue

        # ----------------------------------------------------
        # 5. NORMALIZE / COMPARE
        # ----------------------------------------------------
        tableau_000 = _tableau_to_000(tableau_raw)
        variance_000 = tableau_000 - fs_000

        status = STATUS_PASS if variance_000 == 0 else STATUS_FAIL

        note = _variance_note(variance_000)

        if company == "CBL" and fs_line == "Corporate Sales" and status == STATUS_FAIL:
            note += " Underlying detailed Tableau investigation required."

        result.update(
            {
                "tableau_rounded_000": tableau_000,
                "variance_000": variance_000,
                "status": status,
                "note": note,
            }
        )

        details.append(result)

    failed = sum(row["status"] == STATUS_FAIL for row in details)

    action_required = sum(row["status"] == STATUS_ACTION_REQUIRED for row in details)

    passed = sum(row["status"] == STATUS_PASS for row in details)

    section_status = (
        STATUS_PASS if failed == 0 and action_required == 0 else STATUS_REVIEW_REQUIRED
    )

    return {
        "company": company,
        "status": section_status,
        "passed": passed,
        "failed": failed,
        "action_required": action_required,
        "summary": {
            "passed": (passed),
            "failed": (failed),
            "action_required": (action_required),
        },
        "details": details,
    }


def _style_sheet(
    worksheet,
) -> None:
    """Apply simple FP&A validation formatting."""

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    for cell in worksheet[1]:
        cell.fill = HEADER_FILL
        cell.font = Font(
            color="FFFFFF",
            bold=True,
        )
        cell.alignment = Alignment(
            vertical="center",
            wrap_text=True,
        )

    for column_cells in worksheet.columns:
        max_length = 0
        column_letter = get_column_letter(column_cells[0].column)

        for cell in column_cells:
            if cell.value is None:
                continue

            max_length = max(
                max_length,
                len(str(cell.value)),
            )

        worksheet.column_dimensions[column_letter].width = min(
            max(
                max_length + 2,
                12,
            ),
            45,
        )


def _write_detail_sheet(
    workbook: Workbook,
    *,
    sheet_name: str,
    result: dict[str, Any],
    closed_period: date | datetime,
    forecast_start: date,
) -> None:

    ws = workbook.create_sheet(title=sheet_name)

    # --------------------------------------------------------
    # TITLE / CONTROL AREA
    # --------------------------------------------------------

    ws.merge_cells("A1:G1")

    ws["A1"] = f"{sheet_name} ACTUALS VALIDATION"

    ws["A1"].fill = HEADER_FILL

    ws["A1"].font = Font(
        color="FFFFFF",
        bold=True,
        size=16,
    )

    ws["A1"].alignment = Alignment(
        vertical="center",
    )

    ws.row_dimensions[1].height = 28

    # --------------------------------------------------------
    # CLOSED PERIOD
    # --------------------------------------------------------

    ws.merge_cells("A2:B2")

    ws["A2"] = "Closed Period"

    ws["C2"] = closed_period.strftime("%B %Y")

    # --------------------------------------------------------
    # OVERALL STATUS
    # --------------------------------------------------------

    ws["E2"] = "Overall Status"

    ws.merge_cells("F2:G2")

    ws["F2"] = result["status"]

    ws["F2"].font = Font(
        bold=True,
    )

    if result["status"] == STATUS_PASS:

        ws["F2"].fill = PASS_FILL

    else:

        ws["F2"].fill = FAIL_FILL

    # --------------------------------------------------------
    # TABLEAU FORECAST START
    # --------------------------------------------------------

    ws.merge_cells("A3:B3")

    ws["A3"] = "Tableau Forecast Start"

    ws["C3"] = forecast_start.strftime("%m/%d/%Y")

    ws["C3"].font = Font(
        bold=True,
        color="5B1760",
    )

    # --------------------------------------------------------
    # VALIDATION RULE
    # --------------------------------------------------------

    ws["A4"] = (
        "PASS = rounded Tableau actual agrees with "
        "the finalized financial statement ($000)."
    )

    ws.merge_cells("A4:G4")

    ws["A4"].font = Font(
        italic=True,
        color="666666",
    )

    # --------------------------------------------------------
    # COMPARISON TABLE
    # --------------------------------------------------------

    header_row = 6

    headers = [
        "Financial Statement Line",
        "Financial Statement ($000)",
        "Tableau Actual ($000)",
        "Variance ($000)",
        "Status",
        "Exception / Note",
        "Company",
    ]

    for column_number, header in enumerate(
        headers,
        start=1,
    ):
        cell = ws.cell(
            row=header_row,
            column=column_number,
            value=header,
        )
        cell.fill = HEADER_FILL
        cell.font = Font(
            color="FFFFFF",
            bold=True,
        )
        cell.alignment = Alignment(
            vertical="center",
            wrap_text=True,
        )

    ws.row_dimensions[header_row].height = 30

    for row in result["details"]:

        ws.append(
            [
                row["statement_line"],
                row["fs_rounded_000"],
                row["tableau_rounded_000"],
                row["variance_000"],
                row["status"],
                row["note"],
                row["company"],
            ]
        )

        current_row = ws.max_row

        status_cell = ws.cell(
            row=current_row,
            column=5,
        )

        variance_cell = ws.cell(
            row=current_row,
            column=4,
        )

        if row["status"] == STATUS_PASS:
            status_cell.fill = PASS_FILL
            status_cell.font = Font(
                color="006100",
                bold=True,
            )

        elif row["status"] == STATUS_FAIL:
            status_cell.fill = FAIL_FILL
            status_cell.font = Font(
                color="9C0006",
                bold=True,
            )
            variance_cell.fill = FAIL_FILL
            variance_cell.font = Font(
                color="9C0006",
                bold=True,
            )

        else:
            status_cell.fill = ACTION_FILL
            status_cell.font = Font(
                color="9C6500",
                bold=True,
            )

        # Make total lines visually distinct.
        if "TOTAL" in _normalize(row["tableau_category"]) or row["statement_line"] in {
            "Operating Expenses",
            "Total Operating Costs",
        }:
            for column_number in range(
                1,
                8,
            ):
                ws.cell(
                    row=current_row,
                    column=column_number,
                ).font = Font(
                    bold=True,
                )

    # --------------------------------------------------------
    # FORMATTING
    # --------------------------------------------------------

    for row_number in range(
        header_row + 1,
        ws.max_row + 1,
    ):
        for column_number in (
            2,
            3,
        ):
            ws.cell(
                row=row_number,
                column=column_number,
            ).number_format = "#,##0"

        ws.cell(
            row=row_number,
            column=4,
        ).number_format = "+#,##0;-#,##0;0"

    ws.freeze_panes = f"A{header_row + 1}"

    ws.auto_filter.ref = f"A{header_row}:G{ws.max_row}"

    # Clean widths instead of exposing every technical field.
    widths = {
        "A": 38,
        "B": 24,
        "C": 22,
        "D": 18,
        "E": 20,
        "F": 52,
        "G": 12,
    }

    for column, width in widths.items():
        ws.column_dimensions[column].width = width

    for row in ws.iter_rows(
        min_row=header_row + 1,
        max_row=ws.max_row,
    ):
        for cell in row:
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=(cell.column == 6),
            )

    # Company is useful for filtering but redundant on company-specific tabs.
    # Keep it hidden rather than deleting it from the underlying table.
    ws.column_dimensions["G"].hidden = True

    ws.sheet_view.showGridLines = False


def _write_executive_summary_sheet(
    workbook: Workbook,
    *,
    closed_period: date | datetime,
    forecast_start: date,
    aliv_result: dict[str, Any],
    cbl_result: dict[str, Any],
) -> None:
    """Create a polished executive landing page with company scorecards."""

    ws = workbook.active
    ws.title = "Executive Summary"
    ws.sheet_view.showGridLines = False

    # Title banner
    ws.merge_cells("A1:H2")
    ws["A1"] = "FP&A ACTUALS VALIDATION | EXECUTIVE SUMMARY"
    ws["A1"].fill = TITLE_FILL
    ws["A1"].font = Font(color="FFFFFF", bold=True, size=18)
    ws["A1"].alignment = Alignment(vertical="center", horizontal="left")

    ws.merge_cells("A3:H3")
    ws["A3"] = (
        f"Closed Period: "
        f"{closed_period.strftime('%B %Y')}"
        f"  |  "
        f"Tableau Forecast Start: "
        f"{forecast_start.strftime('%m/%d/%Y')}"
        f"  |  "
        f"Financial Statements vs Tableau"
    )
    ws["A3"].fill = ACCENT_FILL
    ws["A3"].font = Font(bold=True, color="5B1760")

    def add_scorecard(start_col: int, company: str, result: dict[str, Any]) -> None:
        end_col = start_col + 2

        ws.merge_cells(
            start_row=5,
            start_column=start_col,
            end_row=5,
            end_column=end_col,
        )
        cell = ws.cell(row=5, column=start_col, value=f"{company} OPEX")
        cell.fill = HEADER_FILL
        cell.font = Font(color="FFFFFF", bold=True, size=13)
        cell.alignment = Alignment(horizontal="center")

        ws.merge_cells(
            start_row=6,
            start_column=start_col,
            end_row=6,
            end_column=end_col,
        )
        status_cell = ws.cell(row=6, column=start_col, value=result["status"])
        status_cell.fill = PASS_FILL if result["status"] == STATUS_PASS else FAIL_FILL
        status_cell.font = Font(
            bold=True,
            size=14,
            color="006100" if result["status"] == STATUS_PASS else "9C0006",
        )
        status_cell.alignment = Alignment(horizontal="center")

        metrics = [
            ("Passed", result["summary"]["passed"]),
            ("Failed", result["summary"]["failed"]),
            ("Action Required", result["summary"]["action_required"]),
        ]

        for row_number, (label, value) in enumerate(metrics, start=8):
            label_cell = ws.cell(row=row_number, column=start_col, value=label)
            label_cell.fill = NEUTRAL_FILL
            label_cell.font = Font(bold=True, color="5B1760")

            ws.merge_cells(
                start_row=row_number,
                start_column=start_col + 1,
                end_row=row_number,
                end_column=end_col,
            )
            value_cell = ws.cell(
                row=row_number,
                column=start_col + 1,
                value=value,
            )
            value_cell.font = Font(bold=True, size=12)
            value_cell.alignment = Alignment(horizontal="center")

        for row_number in range(5, 11):
            for column_number in range(start_col, end_col + 1):
                ws.cell(
                    row=row_number,
                    column=column_number,
                ).border = THIN_BORDER

    add_scorecard(1, "ALIV", aliv_result)
    add_scorecard(5, "CBL", cbl_result)

    overall_status = (
        STATUS_PASS
        if (
            aliv_result["status"] == STATUS_PASS and cbl_result["status"] == STATUS_PASS
        )
        else STATUS_REVIEW_REQUIRED
    )

    ws.merge_cells("A12:H12")
    ws["A12"] = f"OVERALL VALIDATION STATUS: {overall_status}"
    ws["A12"].fill = PASS_FILL if overall_status == STATUS_PASS else ACTION_FILL
    ws["A12"].font = Font(
        bold=True,
        size=13,
        color="006100" if overall_status == STATUS_PASS else "9C6500",
    )
    ws["A12"].alignment = Alignment(horizontal="center")

    # Exception table
    ws.merge_cells("A14:H14")
    ws["A14"] = "EXCEPTIONS REQUIRING REVIEW"
    ws["A14"].fill = HEADER_FILL
    ws["A14"].font = WHITE_FONT

    headers = [
        "Company",
        "Financial Statement Line",
        "FS ($000)",
        "Tableau ($000)",
        "Variance ($000)",
        "Status",
        "Exception / Note",
        "Owner / Action",
    ]

    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=15, column=col, value=header)
        cell.fill = ACCENT_FILL
        cell.font = Font(bold=True, color="5B1760")
        cell.border = THIN_BORDER
        cell.alignment = Alignment(wrap_text=True)

    row_number = 16

    for company_result in (aliv_result, cbl_result):
        for item in company_result["details"]:
            if item["status"] == STATUS_PASS:
                continue

            values = [
                item["company"],
                item["statement_line"],
                item["fs_rounded_000"],
                item["tableau_rounded_000"],
                item["variance_000"],
                item["status"],
                item["note"],
                "",
            ]

            for col, value in enumerate(values, start=1):
                cell = ws.cell(row=row_number, column=col, value=value)
                cell.border = THIN_BORDER
                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=(col in {2, 7, 8}),
                )

            status_cell = ws.cell(row=row_number, column=6)
            if item["status"] == STATUS_FAIL:
                status_cell.fill = FAIL_FILL
                status_cell.font = Font(color="9C0006", bold=True)
            else:
                status_cell.fill = ACTION_FILL
                status_cell.font = Font(color="9C6500", bold=True)

            ws.cell(
                row=row_number,
                column=5,
            ).number_format = "+#,##0;-#,##0;0"

            row_number += 1

    if row_number == 16:
        ws.merge_cells("A16:H17")
        ws["A16"] = "No validation exceptions identified."
        ws["A16"].fill = PASS_FILL
        ws["A16"].font = Font(bold=True, color="006100")
        ws["A16"].alignment = Alignment(horizontal="center", vertical="center")

    widths = {
        "A": 14,
        "B": 34,
        "C": 16,
        "D": 18,
        "E": 18,
        "F": 20,
        "G": 48,
        "H": 24,
    }

    for column, width in widths.items():
        ws.column_dimensions[column].width = width

    ws.freeze_panes = "A15"


def write_opex_validation_workbook(
    *,
    output_file: str | Path,
    closed_period: date | datetime,
    forecast_start: date,
    aliv_result: dict[str, Any],
    cbl_result: dict[str, Any],
) -> Path:
    """
    Write the validation workbook into the OPEX output folder.
    """

    output_file = Path(output_file)

    output_file.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    workbook = Workbook()

    # Executive Summary is the first sheet users see.
    _write_executive_summary_sheet(
        workbook,
        closed_period=(closed_period),
        forecast_start=(forecast_start),
        aliv_result=(aliv_result),
        cbl_result=(cbl_result),
    )

    _write_detail_sheet(
        workbook,
        sheet_name="ALIV OPEX",
        result=aliv_result,
        closed_period=closed_period,
        forecast_start=forecast_start,
    )

    _write_detail_sheet(
        workbook,
        sheet_name="CBL OPEX",
        result=cbl_result,
        closed_period=closed_period,
        forecast_start=forecast_start,
    )

    workbook["Executive Summary"].sheet_properties.tabColor = "7A1F73"
    workbook["ALIV OPEX"].sheet_properties.tabColor = "D9EAD3"
    workbook["CBL OPEX"].sheet_properties.tabColor = "FCE5CD"
    workbook.active = 0

    workbook.save(output_file)

    return output_file


# ============================================================
# RAW TABLEAU AUDIT COPY
# ============================================================


def _save_tableau_rows_csv(
    *,
    rows: list[dict],
    output_file: str | Path,
) -> Path:

    output_file = Path(output_file)

    output_file.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    if not rows:
        raise RuntimeError("No Tableau rows were available " "for raw audit output.")

    headers = list(rows[0].keys())

    with output_file.open(
        "w",
        encoding="utf-8-sig",
        newline="",
    ) as handle:

        writer = csv.DictWriter(
            handle,
            fieldnames=headers,
        )

        writer.writeheader()
        writer.writerows(rows)

    return output_file


def _save_tableau_rows_workbook(
    *,
    rows: list[dict],
    output_file: str | Path,
) -> Path:
    """
    Save the exact parsed Tableau crosstab into an Excel workbook.

    This workbook is audit evidence only. The comparison still uses the
    SAME in-memory Tableau rows, so Tableau is not extracted twice.
    """

    output_file = Path(output_file)

    output_file.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    if not rows:
        raise RuntimeError("No Tableau rows were available " "for Excel audit output.")

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "OPEX Tableau Extract"

    headers = list(rows[0].keys())

    worksheet.append(headers)

    for row in rows:
        worksheet.append([row.get(header) for header in headers])

    _style_sheet(worksheet)

    workbook.save(output_file)

    return output_file


def _get_forecast_start_from_closed_period(
    closed_period: date | datetime,
) -> date:

    # --------------------------------------------------------
    # Handle December -> January of the next calendar year.
    # --------------------------------------------------------

    if closed_period.month == 12:

        return date(
            closed_period.year + 1,
            1,
            1,
        )

    # --------------------------------------------------------
    # All other months simply move forward one month.
    # --------------------------------------------------------

    return date(
        closed_period.year,
        closed_period.month + 1,
        1,
    )


# ============================================================
# LIVE OPEX VALIDATION
# ===========================================================


def run_live_opex_validation(
    *,
    aliv_fs_path: str | Path,
    cbl_fs_path: str | Path,
    output_paths: dict[str, Path],
    closed_period: date | datetime,
    tableau_workbook_path: str | Path,
    open_tableau_first: bool = False,
    tableau_open_wait_seconds: int = 8,
    run_individual_validation: bool = True,
) -> dict[str, Any]:

    # --------------------------------------------------------
    # 1. READ FINANCIAL STATEMENTS
    # --------------------------------------------------------

    fs_results = {"ALIV": read_aliv_opex(aliv_fs_path, closed_period)}

    # ========================================================
    # 2. DETERMINE TABLEAU FORECAST START
    # ========================================================

    from automations.actuals.cbl_statement_periods import read_cbl_area

    fs_results["CBL"] = read_cbl_area(cbl_fs_path, "OPEX", closed_period)["MONTH"]

    forecast_start = _get_forecast_start_from_closed_period(closed_period)

    print()
    print("=" * 70)
    print("TABLEAU ACTUALS PERIOD")
    print("=" * 70)

    print(
        "Actuals Month  :",
        closed_period.strftime("%B %Y"),
    )

    print(
        "Forecast Start :",
        forecast_start.strftime("%m/%d/%Y"),
    )

    print("=" * 70)

    # ========================================================
    # 3. UPDATE TABLEAU FORECAST START
    # ========================================================

    from uuid import uuid4
    from automations.actuals.actuals_report_layout import prepare_group_copy

    source_tableau_path = Path(tableau_workbook_path)
    from shared.tableau_temp_copies import temp_twb_path

    tableau_workbook_path = temp_twb_path(source_tableau_path, "GroupActuals")
    prepare_group_copy(
        source_tableau_path,
        tableau_workbook_path,
        OPEX_TABLEAU_WORKSHEET,
        closed_period,
    )

    # ========================================================
    # 4. OPTIONAL - OPEN MAPPED TABLEAU WORKBOOK
    # ========================================================

    if open_tableau_first:

        from automations.actuals.tableau_forecast_service import (
            open_tableau_workbook,
        )

        open_tableau_workbook(tableau_workbook_path=(tableau_workbook_path))

        print("Mapped OPEX Tableau workbook opened:")

        print(tableau_workbook_path)

        print("Waiting for Tableau Desktop...")

        time.sleep(tableau_open_wait_seconds)
    # --------------------------------------------------------
    # 5. EXTRACT TABLEAU CROSSTAB
    # --------------------------------------------------------

    from automations.actuals.tableau_desktop_export_service import get_tableau_crosstab_text
    from automations.actuals.tableau_desktop_close_service import close_tableau_forecast_workbook
    from automations.actuals.direct_cost_validation_service import _parse_direct_cost_crosstab_text

    raw_folder = Path(output_paths["tableau_raw_copies"])
    try:
        text = get_tableau_crosstab_text(
            worksheet_name=OPEX_TABLEAU_WORKSHEET,
            tableau_workbook_path=str(tableau_workbook_path),
        )
    finally:
        close_tableau_forecast_workbook(
            tableau_workbook_path=str(tableau_workbook_path)
        )
        Path(tableau_workbook_path).unlink(missing_ok=True)
    tableau_rows = _parse_direct_cost_crosstab_text(text)

    # Save a parsed CSV copy as a second audit artifact.
    period_tag = closed_period.strftime("%Y-%m")

    raw_csv = raw_folder / f"OPEX_Tableau_Raw_{period_tag}.csv"

    _save_tableau_rows_csv(
        rows=tableau_rows,
        output_file=raw_csv,
    )

    # Save the SAME Tableau extract to Excel for easy human review.
    raw_excel = raw_folder / f"OPEX_Tableau_Extract_{period_tag}.xlsx"

    _save_tableau_rows_workbook(
        rows=tableau_rows,
        output_file=raw_excel,
    )

    # --------------------------------------------------------
    # 6. BUILD TABLEAU VALUES BY COMPANY/CATEGORY
    # --------------------------------------------------------

    tableau_values = build_opex_tableau_values(
        tableau_rows,
        closed_period=(closed_period),
    )

    # --------------------------------------------------------
    # 7. COMPARE ALIV AND CBL
    # --------------------------------------------------------

    aliv_result = compare_opex_company(
        company="ALIV",
        fs_values=(fs_results["ALIV"]),
        tableau_values=(tableau_values["ALIV"]),
    )

    cbl_result = compare_opex_company(
        company="CBL",
        fs_values=(fs_results["CBL"]),
        tableau_values=(tableau_values["CBL"]),
    )

    # --------------------------------------------------------
    # 8. WRITE OPEX VALIDATION WORKBOOK
    # --------------------------------------------------------

    validation_file = Path(output_paths["opex"]) / (
        "OPEX_Actuals_Validation_" f"{period_tag}.xlsx"
    )

    write_opex_validation_workbook(
        output_file=(validation_file),
        closed_period=(closed_period),
        forecast_start=(forecast_start),
        aliv_result=(aliv_result),
        cbl_result=(cbl_result),
    )

    _validate_group_period_headers(tableau_rows, closed_period)
    tableau_ytd_values = build_opex_tableau_values(
        tableau_rows, closed_period=closed_period, basis="YTD"
    )
    _write_group_actuals_sheet(
        validation_file,
        tableau_values,
        tableau_ytd_values,
        closed_period,
        aliv_fs_path=aliv_fs_path,
        cbl_fs_path=cbl_fs_path,
    )
    individual_result = {"status": "PENDING"}
    if run_individual_validation:
        from automations.actuals.individual_actuals_validation_service import (
            run_individual_opex_pilot,
        )

        individual_result = run_individual_opex_pilot(
            report_path=validation_file,
            aliv_fs_path=aliv_fs_path,
            output_paths=output_paths,
            closed_period=closed_period,
        )

    # Final report must be created AFTER all individual captures have been merged.
    # The intermediate A:S workbook is never the user-facing deliverable.
    from automations.actuals.actuals_report_layout import (
        finalize_actuals_report_in_place,
    )
    finalize_actuals_report_in_place(
        validation_file,
        forecast_start=forecast_start,
    )

    # ========================================================
    # 9. DETERMINE OVERALL VALIDATION STATUS
    # ========================================================

    overall_status = (
        STATUS_PASS
        if (
            aliv_result["status"] == STATUS_PASS
            and cbl_result["status"] == STATUS_PASS
            and (not run_individual_validation or individual_result["status"] == "PASS")
        )
        else STATUS_REVIEW_REQUIRED
    )

    # ========================================================
    # BUILD RESULT OBJECT
    # ========================================================

    from automations.actuals.actuals_report_layout import report_status

    overall_status = report_status(validation_file)

    result = {
        "individual_actuals": individual_result,
        "area": "OPEX",
        "closed_period": (closed_period.strftime("%Y-%m-%d")),
        "forecast_start": (forecast_start.strftime("%Y-%m-%d")),
        "overall_status": (overall_status),
        "tableau_workbook": str(source_tableau_path),
        "tableau_worksheet": (OPEX_TABLEAU_WORKSHEET),
        "financial_statements": {
            "ALIV": str(aliv_fs_path),
            "CBL": str(cbl_fs_path),
        },
        "validation_workbook": str(validation_file),
        "raw_tableau_csv": str(raw_csv),
        "raw_tableau_workbook": str(raw_excel),
        "companies": {
            "ALIV": aliv_result,
            "CBL": cbl_result,
        },
    }
    # ========================================================
    # CREATE UNIQUE TIME-STAMPED LOG FILE
    # ========================================================

    run_timestamp = datetime.now()

    run_by = getpass.getuser()

    log_file = Path(output_paths["logs"]) / (
        "OPEX_Validation_"
        f"{closed_period:%Y-%m}_"
        f"{run_timestamp:%Y%m%d_%H%M%S}"
        ".log"
    )
    # ========================================================
    # CAPTURE RUN AUDIT INFORMATION
    # ========================================================

    run_timestamp = datetime.now()

    run_by = getpass.getuser()

    # ========================================================
    # WRITE VALIDATION RUN LOG
    # ========================================================

    log_file.write_text(
        "\n".join(
            [
                "OPEX ACTUALS VALIDATION",
                "=" * 70,
                "",
                # ------------------------------------------------
                # RUN INFORMATION
                # ------------------------------------------------
                "RUN INFORMATION",
                (f"Run Date: " f"{run_timestamp:%m/%d/%Y}"),
                (f"Run Time: " f"{run_timestamp:%I:%M:%S %p}"),
                (f"Run By: " f"{run_by}"),
                "",
                # ------------------------------------------------
                # VALIDATION PERIOD
                # ------------------------------------------------
                "VALIDATION PERIOD",
                (f"Closed Period: " f"{closed_period:%B %Y}"),
                (f"Forecast Start: " f"{forecast_start:%m/%d/%Y}"),
                "",
                # ------------------------------------------------
                # SOURCE FILES
                # ------------------------------------------------
                "SOURCE FILES",
                (f"ALIV Financial Statement: " f"{Path(aliv_fs_path).resolve()}"),
                (f"CBL Financial Statement: " f"{Path(cbl_fs_path).resolve()}"),
                (f"Tableau Workbook: " f"{Path(tableau_workbook_path).resolve()}"),
                (f"Tableau Worksheet: " f"{OPEX_TABLEAU_WORKSHEET}"),
                "",
                # ------------------------------------------------
                # VALIDATION RESULTS
                # ------------------------------------------------
                "VALIDATION RESULTS",
                (f"Overall Status: " f"{overall_status}"),
                (f"ALIV Status: " f"{aliv_result['status']}"),
                (f"CBL Status: " f"{cbl_result['status']}"),
                "",
                # ------------------------------------------------
                # OUTPUT FILES
                # ------------------------------------------------
                "OUTPUT FILES",
                (f"Validation Workbook: " f"{validation_file}"),
                (f"Raw Tableau CSV: " f"{raw_csv}"),
                (f"Raw Tableau Workbook: " f"{raw_excel}"),
                "",
                "=" * 70,
            ]
        ),
        encoding="utf-8",
    )
    print()
    print("=" * 70)
    print("OPEX ACTUALS VALIDATION COMPLETE")
    print("=" * 70)
    print(
        "Overall:",
        overall_status,
    )
    print(
        "ALIV:",
        aliv_result["status"],
    )
    print(
        "CBL:",
        cbl_result["status"],
    )
    print(
        "Workbook:",
        validation_file,
    )
    print(
        "Raw Tableau CSV:",
        raw_csv,
    )
    print(
        "Raw Tableau Workbook:",
        raw_excel,
    )
    print("=" * 70)

    return result


def _write_group_actuals_sheet(
    report_path,
    monthly,
    ytd,
    closed_period,
    area="OPEX",
    aliv_fs_path=None,
    cbl_fs_path=None,
):
    from automations.actuals.actuals_report_layout import merge_group_values, refresh_summary
    from openpyxl import load_workbook

    merge_group_values(
        report_path, monthly, ytd, closed_period, area, aliv_fs_path, cbl_fs_path
    )
    report = load_workbook(report_path)
    try:
        refresh_summary(report)
        report.save(report_path)
    finally:
        report.close()


def _validate_group_period_headers(rows, closed_period):
    """Require fiscal months through close before treating Grand Total as YTD."""
    import re

    month_names = "|".join(date(2000, m, 1).strftime("%B") for m in range(1, 13))
    seen = set()
    for header in rows[0]:
        match = re.match(
            rf"^({month_names}) (\d{{4}})(?: |$)", str(header).strip(), re.I
        )
        if match:
            seen.add(datetime.strptime(" ".join(match.groups()), "%B %Y").date())
    current = date(closed_period.year - (closed_period.month < 7), 7, 1)
    expected = set()
    while current <= closed_period:
        expected.add(current)
        current = date(current.year + (current.month == 12), current.month % 12 + 1, 1)
    if seen != expected:
        raise ValueError(
            "Group worksheet posting-month filters must include only fiscal July through the selected closed month before Grand Total can be treated as YTD."
        )
