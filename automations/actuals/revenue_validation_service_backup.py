"""Capture selected Revenue TWBs and validate monthly and YTD separately."""

from datetime import date, datetime
from pathlib import Path
from decimal import Decimal
from io import StringIO
from uuid import uuid4
import csv
import re
import time
import xml.etree.ElementTree as ET
from openpyxl import load_workbook

SOURCES = [
    (
        "consumer",
        "ALIV Consumer",
        "ALIV RESIDENTIAL REVENUE Detail - Monthly Actual v Budget YTD",
    ),
    ("consumer", "CBL Consumer", "CBL RESIDENTIAL REVENUE SUMMARY - Monthly Actuals"),
    (
        "corporate",
        "Corporate",
        "CBL & ALIV REVENUE SUMMARY - Month YTD Actual v Budget",
    ),
    ("group", "CBL Group", "CBL FS Rollup View Group Monthly Revenue YTD_SUMMARY"),
]


def canonical(value):
    return " ".join(str(value or "").strip().upper().split())


def number(value):
    text = str(value or "").strip().replace(",", "").replace("$", "")
    if not text:
        raise ValueError(
            "A mapped Revenue actual is blank; it cannot be treated as zero."
        )
    if text.startswith("(") and text.endswith(")"):
        text = "-" + text[1:-1]
    result = Decimal(text)
    if not result.is_finite():
        raise ValueError("Invalid Revenue actual.")
    return float(result)


def prepare_revenue_copy(source, destination, tabs, closed_period, basis):
    root = ET.parse(source).getroot()
    start = date(
        closed_period.year + (closed_period.month == 12),
        closed_period.month % 12 + 1,
        1,
    )
    fiscal_start = date(closed_period.year - (closed_period.month < 7), 7, 1)
    lower = (
        fiscal_start
        if basis == "YTD"
        else date(closed_period.year, closed_period.month, 1)
    )
    actual_names = set()
    for tab in tabs:
        sheet = root.find(f'./worksheets/worksheet[@name="{tab}"]')
        if sheet is None:
            raise ValueError(f"Revenue worksheet not found: {tab}")
        date_filters = [
            f
            for f in sheet.iter("filter")
            if any(
                f.get("column", "").endswith(f"[{level}:GL Posting Date:ok]")
                for level in ("my", "tmn")
            )
        ]
        if not date_filters:
            raise ValueError(
                f"No mapped posting-date filter on Revenue worksheet: {tab}"
            )
        actual_names.update(
            c.get("name")
            for c in sheet.iter("column")
            if c.get("caption") == "Actual YTD"
        )
    parameters = [
        c
        for c in root.iter("column")
        if canonical(c.get("caption")) == "FORECAST START"
        and "param-domain-type" in c.attrib
    ]
    if not parameters or len({c.get("name") for c in parameters}) != 1:
        raise ValueError("Expected one Revenue Forecast Start parameter name.")
    for parameter in parameters:
        parameter.set("value", f"#{start:%Y-%m-%d}#")
        if parameter.find("calculation") is not None:
            parameter.find("calculation").set("formula", f"#{start:%Y-%m-%d}#")
    formula = f"IF [GL Posting Date] >= #{lower:%Y-%m-%d}# AND [GL Posting Date] < #{start:%Y-%m-%d}# THEN [Actual] ELSE 0 END"
    for c in root.iter("column"):
        if c.get("name") in actual_names and c.find("calculation") is not None:
            c.find("calculation").set("formula", formula)
    months = []
    current = lower
    while current < start:
        months.append(current)
        current = date(current.year + (current.month == 12), current.month % 12 + 1, 1)
    for f in root.iter("filter"):
        column = f.get("column", "")
        for level in ("my", "tmn"):
            if column.endswith(f"[{level}:GL Posting Date:ok]"):
                for child in list(f):
                    f.remove(child)
                union = ET.SubElement(f, "groupfilter", function="union")
                for month in months:
                    member = (
                        f"{month:%Y%m}"
                        if level == "my"
                        else f"#{month:%Y-%m-%d} 00:00:00#"
                    )
                    ET.SubElement(
                        union,
                        "groupfilter",
                        function="member",
                        level=f"[{level}:GL Posting Date:ok]",
                        member=member,
                    )
        if column.endswith("[none:Fiscal Year:ok]"):
            for child in list(f):
                f.remove(child)
            ET.SubElement(
                f,
                "groupfilter",
                function="member",
                level="[none:Fiscal Year:ok]",
                member=str(fiscal_start.year + 1),
            )
        if column.endswith("[none:Forecast Type:nk]"):
            for member in f.findall('.//groupfilter[@function="member"]'):
                if re.fullmatch(
                    r"FY\d{2} Forecast_\d{6}", member.get("member", "").strip('"'), re.I
                ):
                    member.set(
                        "member",
                        f'"FY{(fiscal_start.year+1)%100:02} Forecast_{start:%m%Y}"',
                    )
    ET.ElementTree(root).write(destination, encoding="utf-8", xml_declaration=True)


def parse_capture(text, basis, closed_period):
    """Flatten one or two header rows and restore suppressed hierarchy labels."""
    raw = [
        r
        for r in csv.reader(StringIO(text), delimiter="\t")
        if any(v.strip() for v in r)
    ]
    if len(raw) < 2:
        raise ValueError("Revenue capture is empty.")
    first = raw[0]
    two_headers = any(
        canonical(v) in ("ACTUAL YTD", "BUDGET YTD", "VAR YTD $", "VAR YTD %")
        for v in raw[1]
    )
    if two_headers:
        from automations.actuals.direct_cost_validation_service import (
            _parse_direct_cost_crosstab_text,
        )

        rows = _parse_direct_cost_crosstab_text(text)
    else:
        if len(first) != len(set(first)):
            raise ValueError("Revenue capture has duplicate headers.")
        rows = [dict(zip(first, r + [""] * (len(first) - len(r)))) for r in raw[1:]]
    headers = list(rows[0])
    normalized = {canonical(h): h for h in headers}
    target = "GRAND TOTAL ACTUAL YTD" if two_headers else "GRAND TOTAL"
    if target not in normalized:
        if "ACTUAL YTD" in normalized:
            target = "ACTUAL YTD"
        elif basis == "MONTH":
            labels = [
                closed_period.strftime("%b-%y").upper(),
                closed_period.strftime("%B %Y").upper(),
            ]
            target = next((label for label in labels if label in normalized), target)
    if target not in normalized:
        raise ValueError(f"Revenue actual column is missing: {target}")
    # Period filters were set on the temporary TWB before capture. Check visible month headers too.
    visible = set()
    for h in headers:
        for pattern, fmt in [
            (r"^([A-Za-z]+ \d{4})(?: |$)", "%B %Y"),
            (r"^([A-Za-z]{3}-\d{2})(?: |$)", "%b-%y"),
        ]:
            match = re.match(pattern, str(h))
            if match:
                try:
                    visible.add(datetime.strptime(match[1], fmt).date())
                except ValueError:
                    pass
    expected = set()
    current = (
        date(closed_period.year - (closed_period.month < 7), 7, 1)
        if basis == "YTD"
        else date(closed_period.year, closed_period.month, 1)
    )
    while current <= closed_period:
        expected.add(current)
        current = date(current.year + (current.month == 12), current.month % 12 + 1, 1)
    if visible and visible != expected:
        raise ValueError(
            "Revenue capture month headers do not match the requested period."
        )
    dimensions = [h for h in headers if h.startswith("GL ") or h == "Company"]
    last = {h: "" for h in dimensions}
    for row in rows:
        for h in dimensions:
            value = str(row.get(h) or "").strip()
            if value:
                if value != last[h]:
                    for lower in dimensions[dimensions.index(h) + 1 :]:
                        last[lower] = ""
                last[h] = value
            else:
                row[h] = last[h]
        row["_actual"] = row.get(normalized[target])
    return rows


def unique(rows, **dimensions):
    matches = [
        r
        for r in rows
        if all(canonical(r.get(k)) == canonical(v) for k, v in dimensions.items())
    ]
    if len(matches) != 1:
        raise ValueError(
            f"Expected one Revenue subtotal for {dimensions}; found {len(matches)}."
        )
    return number(matches[0]["_actual"])


def source_values(label, rows):
    if label == "ALIV Consumer":
        types = {
            "Consumer Prepaid": "PAY IN ADVANCE",
            "Consumer Postpaid": "POSTPAID",
            "Visitor Roaming": "VISITOR ROAMING",
            "Interconnect": "INTERCONNECT REVENUE",
        }
        return {
            "ALIV": {
                line: unique(
                    rows,
                    **{
                        "Company": "ALIV",
                        "GL FS Type": kind,
                        "GL FS IS Breakdown": "Total",
                    },
                )
                for line, kind in types.items()
            }
        }
    if label == "CBL Consumer":
        return {
            "CBL": {
                "Residential": unique(
                    rows,
                    **{
                        "Company": "CBL",
                        "GL Bus Code Name": "Total",
                        "GL Account Mapping": "Total",
                    },
                )
            }
        }
    if label == "Corporate":
        aliv = unique(rows, **{"Company": "ALIV", "GL Account Mapping": "Total"})
        total = unique(rows, **{"Company": "CBL", "GL Account Mapping": "Total"})
        oneoff = unique(
            rows, **{"Company": "CBL", "GL Account Mapping": "CORPORATE ONE OFF"}
        )
        return {
            "ALIV": {"Corporate": aliv},
            "CBL": {
                "Corporate - Recurring": total - oneoff,
                "Corporate - One off": oneoff,
            },
        }
    routes = {
        "Government": ("ALIV BUSINESS", "CORPORATE GOVERNMENT", "Total"),
        "Other": ("HFC FIXED NETWORK CONSUMER", "OTHER SUBSCRIBER", "Total"),
        "MSA-ALIV": ("INTERCOMPANY", "OTHER SUBSCRIBER", "Total"),
        "Media": ("MEDIA", "CORPORATE POST PAID RECURRING", "Total"),
        "Gross Revenue": ("Grand Total", "Total", "Total"),
    }
    return {
        "CBL": {
            line: unique(
                rows,
                **dict(
                    zip(
                        ["GL Bus Code Name", "GL FS Category", "GL Account Mapping"],
                        route,
                    )
                ),
            )
            for line, route in routes.items()
        }
    }


def aliv_statement_periods(path, closed_period):
    from automations.actuals.revenue_financial_statement_reader import (
        ALIV_REVENUE_LINES,
        _canonical_statement_line,
    )

    report = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = report["P & L_"]
        period = sheet["D1"].value
        if not isinstance(period, (date, datetime)) or (period.year, period.month) != (
            closed_period.year,
            closed_period.month,
        ):
            raise ValueError("ALIV statement month does not match selected close.")
        if (
            str(sheet["C5"].value).strip().lower() != "actual"
            or str(sheet["G5"].value).strip().lower() != "actual"
            or "YTD" not in str(sheet["H3"].value).upper()
        ):
            raise ValueError("ALIV statement monthly/YTD headers do not match.")
        result = {"MONTH": {}, "YTD": {}}
        for row in sheet.iter_rows(values_only=True):
            key = _canonical_statement_line(row[1])
            if key in ALIV_REVENUE_LINES and isinstance(row[2], (int, float)):
                label = ALIV_REVENUE_LINES[key]
                if label in result["MONTH"]:
                    raise ValueError(f"Duplicate ALIV Revenue statement label: {label}")
                result["MONTH"][label] = number(row[2])
                result["YTD"][label] = number(row[6])
                if canonical(label) == "SERVICE REVENUES":
                    break
        missing = set(ALIV_REVENUE_LINES.values()) - set(result["MONTH"])
        if missing:
            raise ValueError(f"ALIV Revenue statement lines missing: {sorted(missing)}")
        return result
    finally:
        report.close()


def run_live_revenue_validation(
    *,
    aliv_fs_path,
    cbl_fs_path,
    tableau_sources,
    output_paths,
    closed_period,
    capture=None,
    opener=None,
    closer=None,
    wait_seconds=8,
):
    from automations.actuals.cbl_statement_periods import read_cbl_area
    from automations.actuals.revenue_validation_service import compare_revenue_company
    from automations.actuals.actuals_report_layout import write_area_workbook
    from automations.actuals.actuals_report_layout import (
        add_company_columns,
        status_cell,
        refresh_summary,
        report_status,
    )
    from automations.actuals.tableau_standalone_forecast_service import (
        _append_crosstab_audit,
    )

    if capture is None:
        from automations.actuals.tableau_desktop_export_service import (
            get_tableau_crosstab_text,
        )

        capture = get_tableau_crosstab_text
    if opener is None:
        from automations.actuals.tableau_forecast_service import open_tableau_workbook

        opener = open_tableau_workbook
    if closer is None:
        from automations.actuals.tableau_desktop_close_service import (
            close_tableau_forecast_workbook,
        )

        closer = close_tableau_forecast_workbook
    # Validate all selected files/tabs before opening Tableau.
    for key, _, tab in SOURCES:
        path = Path(tableau_sources[key])
        if not path.is_file() or path.suffix.lower() != ".twb":
            raise ValueError(f"Select the Revenue {key} TWB.")
        if (
            ET.parse(path).getroot().find(f'./worksheets/worksheet[@name="{tab}"]')
            is None
        ):
            raise ValueError(
                f"{path.name} does not contain the mapped Revenue tab: {tab}"
            )
    statements = {
        "ALIV": aliv_statement_periods(aliv_fs_path, closed_period),
        "CBL": read_cbl_area(cbl_fs_path, "Revenue", closed_period),
    }
    tag = closed_period.strftime("%Y-%m")
    folder = Path(output_paths["revenue"])
    folder.mkdir(parents=True, exist_ok=True)
    raw = (
        Path(output_paths["tableau_raw_copies"]) / f"Revenue_Actuals_{uuid4().hex[:8]}"
    )
    raw.mkdir(parents=True, exist_ok=True)
    audit = raw / "Revenue_Actuals_Crosstabs.xlsx"
    values = {basis: {"ALIV": {}, "CBL": {}} for basis in ("MONTH", "YTD")}
    notes = {basis: {"ALIV": [], "CBL": []} for basis in values}
    sequence = 0
    for basis in ("MONTH", "YTD"):
        for key in ("consumer", "corporate", "group"):
            source = Path(tableau_sources[key])
            targets = [(label, tab) for k, label, tab in SOURCES if k == key]
            temporary = source.with_name(
                f"{source.stem}_RevenueActuals_{uuid4().hex[:8]}.twb"
            )
            opened = False
            try:
                prepare_revenue_copy(
                    source, temporary, [tab for _, tab in targets], closed_period, basis
                )
                opener(tableau_workbook_path=str(temporary))
                opened = True
                if wait_seconds:
                    time.sleep(wait_seconds)
                for label, tab in targets:
                    sequence += 1
                    print(f"Revenue {basis}: {label} | {tab}", flush=True)
                    try:
                        text = capture(
                            worksheet_name=tab, tableau_workbook_path=str(temporary)
                        )
                        (
                            raw
                            / f'{sequence:03}_{key}_{label.replace(" ","_")}_{basis}.txt'
                        ).write_text(text, encoding="utf-8")
                        raw_rows = list(csv.reader(StringIO(text), delimiter="\t"))
                        audit_rows = [
                            {f"Column {i+1}": value for i, value in enumerate(row)}
                            for row in raw_rows
                        ]
                        _append_crosstab_audit(
                            audit,
                            audit_rows,
                            source,
                            temporary,
                            tab,
                            label,
                            "Revenue",
                            basis,
                            sequence,
                        )
                        rows = parse_capture(text, basis, closed_period)
                        parsed = source_values(label, rows)
                        for company, items in parsed.items():
                            if set(items) & set(values[basis][company]):
                                raise ValueError(
                                    "Duplicate Revenue mapping across sources."
                                )
                            values[basis][company].update(items)
                    except Exception as exc:
                        affected = (
                            ("ALIV",)
                            if label == "ALIV Consumer"
                            else ("ALIV", "CBL") if label == "Corporate" else ("CBL",)
                        )
                        for company in affected:
                            notes[basis][company].append(
                                f"{label}: {type(exc).__name__}: {exc}"
                            )
            finally:
                if opened:
                    closer(tableau_workbook_path=str(temporary))
                temporary.unlink(missing_ok=True)
    compared = {
        basis: {
            company: compare_revenue_company(
                company=company,
                fs_values=statements[company][basis],
                tableau_values=values[basis][company],
            )
            for company in ("ALIV", "CBL")
        }
        for basis in values
    }
    for basis in values:
        for company in ("ALIV", "CBL"):
            for item in compared[basis][company]["details"]:
                if item["status"] == "ACTION REQUIRED":
                    item["note"] = (
                        "; ".join(notes[basis][company])
                        or "Revenue subtotal mapping is not confirmed for this statement line."
                    )
    path = folder / f"Revenue_Actuals_Validation_{tag}.xlsx"
    start = date(
        closed_period.year + (closed_period.month == 12),
        closed_period.month % 12 + 1,
        1,
    )
    write_area_workbook(
        area="Revenue",
        output_file=path,
        closed_period=closed_period,
        forecast_start=start,
        aliv_result=compared["MONTH"]["ALIV"],
        cbl_result=compared["MONTH"]["CBL"],
    )
    report = load_workbook(path)
    try:
        for company in ("ALIV", "CBL"):
            sheet = report[f"{company} Revenue"]
            add_company_columns(sheet)
            lookup = {
                item["statement_line"]: item
                for item in compared["YTD"][company]["details"]
            }
            for row in range(7, sheet.max_row + 1):
                item = lookup[sheet.cell(row, 1).value]
                for col, key in [
                    (8, "fs_rounded_000"),
                    (9, "tableau_rounded_000"),
                    (10, "variance_000"),
                ]:
                    sheet.cell(row, col, item[key])
                status_cell(sheet.cell(row, 11), item["status"])
                if item["status"] != "PASS":
                    from openpyxl.comments import Comment

                    sheet.cell(row, 11).comment = Comment(item["note"], "Validation")
        refresh_summary(report)
        report.save(path)
    finally:
        report.close()
    status = "REVIEW REQUIRED"  # Revenue mappings are not yet fully verified.
    log = (
        Path(output_paths["logs"])
        / f"Revenue_Validation_{tag}_{datetime.now():%Y%m%d_%H%M%S}.log"
    )
    log.parent.mkdir(parents=True, exist_ok=True)
    log.write_text(
        "\n".join(
            [
                f"Revenue actuals: {tag}",
                f"Overall: {status}",
                f"Report: {path}",
                f"Crosstabs: {audit}",
            ]
            + [
                f"{basis} {company}: {note}"
                for basis in notes
                for company in notes[basis]
                for note in notes[basis][company]
            ]
        ),
        encoding="utf-8",
    )
    return {
        "area": "Revenue",
        "overall_status": status,
        "validation_workbook": str(path),
        "raw_tableau_workbook": str(audit),
        "log_file": str(log),
        "companies": compared["MONTH"],
        "ytd_companies": compared["YTD"],
    }


def compare_revenue_company(*, company, fs_values, tableau_values):
    """Conservative comparison: never certify incomplete Revenue mappings."""
    from automations.actuals.mapping import REVENUE_MAPPINGS
    from decimal import Decimal, ROUND_HALF_UP

    def as_000(value):
        if value is None:
            return None
        return int(Decimal(str(value)).quantize(Decimal("1"), rounding=ROUND_HALF_UP))

    details = []
    for mapping in REVENUE_MAPPINGS:
        if canonical(mapping.get("company")) != canonical(company):
            continue
        line = mapping["fs_line"]
        fs_raw = fs_values.get(line)
        tableau_raw = tableau_values.get(line)
        fs_rounded = as_000(fs_raw)
        tableau_rounded = as_000(Decimal(str(tableau_raw)) / 1000) if tableau_raw is not None else None
        difference = (tableau_rounded - fs_rounded) if fs_rounded is not None and tableau_rounded is not None else None
        # The current Revenue mapping remains incomplete; even zero variances need review.
        details.append({
            "company": company,
            "statement_line": line,
            "tableau_category": line,
            "fs_raw_000": fs_raw,
            "fs_rounded_000": fs_rounded,
            "tableau_raw_dollars": tableau_raw,
            "tableau_rounded_000": tableau_rounded,
            "variance_000": difference,
            "status": "ACTION REQUIRED" if fs_raw is None else "REVIEW REQUIRED",
            "note": "Revenue mapping and period amounts require confirmation; no automatic PASS.",
        })
    return {"company": company, "details": details, "overall_status": "REVIEW REQUIRED"}


def validate_revenue_workbook_path(revenue_workbook_path):
    if not revenue_workbook_path:
        raise ValueError("Please select the consolidated Revenue Tableau workbook.")

    workbook = Path(revenue_workbook_path)

    if not workbook.is_file():
        raise FileNotFoundError(f"Revenue Tableau workbook not found: {workbook}")

    if workbook.suffix.lower() not in (".twb", ".twbx"):
        raise ValueError("Revenue Tableau workbook must be a .twb or .twbx file.")

    return str(workbook)
