"""Shared actuals report layout and period preparation for group captures."""

from pathlib import Path
from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP
from copy import copy
import re
import xml.etree.ElementTree as ET
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.comments import Comment

EXTRA_HEADERS = [
    "YTD Statement ($000)",
    "YTD Group Tableau ($000)",
    "YTD Group - Statement ($000)",
    "YTD Group Status",
    "Monthly Individual ($000)",
    "Monthly Individual - Statement ($000)",
    "Monthly Individual - Group ($000)",
    "Monthly Individual Status",
    "YTD Individual ($000)",
    "YTD Individual - Statement ($000)",
    "YTD Individual - Group ($000)",
    "YTD Individual Status",
]


def canonical(value):
    return " ".join(str(value or "").upper().replace("&", " AND ").split())


def rounded(value):
    return int(Decimal(str(value)).quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def add_company_columns(sheet):
    from openpyxl.utils import get_column_letter

    for column, label in enumerate(EXTRA_HEADERS, 8):
        cell = sheet.cell(6, column, label)
        cell.fill = PatternFill("solid", fgColor="145A70")
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
        sheet.column_dimensions[get_column_letter(column)].width = 21
    for cell in sheet[6]:
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
    sheet.row_dimensions[6].height = 55
    sheet.auto_filter.ref = f"A6:S{sheet.max_row}"
    sheet.freeze_panes = "B7"
    for row in range(7, sheet.max_row + 1):
        for column in (11, 15, 19):
            if sheet.cell(row, column).value is None:
                sheet.cell(row, column, "NOT VALIDATED")
        for column in (8, 9, 10, 12, 13, 14, 16, 17, 18):
            sheet.cell(row, column).number_format = "#,##0;[Red](#,##0);–"


def status_cell(cell, status):
    cell.value = status
    cell.font = Font(
        bold=True,
        color=(
            "006100"
            if status == "PASS"
            else "9C0006" if status in ("FAIL", "ACTION REQUIRED") else "666666"
        ),
    )
    cell.fill = PatternFill(
        "solid",
        fgColor=(
            "E2EFDA"
            if status == "PASS"
            else "F4CCCC" if status in ("FAIL", "ACTION REQUIRED") else "EEEEEE"
        ),
    )


def write_area_workbook(
    *, area, output_file, closed_period, forecast_start, aliv_result, cbl_result
):
    # All areas use the same OPEX formatter and its company/detail result schema.
    from automations.actuals.opex_validation_service import (
        write_opex_validation_workbook,
    )

    for result in (aliv_result, cbl_result):
        for item in result["details"]:
            item.setdefault("tableau_category", item["statement_line"])
    path = write_opex_validation_workbook(
        output_file=output_file,
        closed_period=closed_period,
        forecast_start=forecast_start,
        aliv_result=aliv_result,
        cbl_result=cbl_result,
    )
    report = load_workbook(path)
    try:
        for company in ("ALIV", "CBL"):
            sheet = report[f"{company} OPEX"]
            sheet.title = f"{company} {area}"
            add_company_columns(sheet)
        for sheet in report:
            for cells in sheet:
                for cell in cells:
                    if isinstance(cell.value, str) and "OPEX" in cell.value:
                        cell.value = cell.value.replace("OPEX", area)
        refresh_summary(report)
        report.save(path)
    finally:
        report.close()
    return path


def statement_ytd(path, closed_period, area=None):
    if not path:
        return {}
    from automations.actuals.adaptive_aliv_statement_reader import read_aliv_period

    return {
        k: rounded(v[1])
        for k, v in read_aliv_period(path, closed_period, area).items()
        if v is not None
    }


def merge_group_values(
    report_path,
    monthly,
    ytd,
    closed_period,
    area="OPEX",
    aliv_fs_path=None,
    cbl_fs_path=None,
):
    from automations.actuals.validation_config import (
        OPEX_MAPPINGS,
        DIRECT_COST_MAPPINGS,
        AGG_TOTAL,
    )

    mappings = OPEX_MAPPINGS if area == "OPEX" else DIRECT_COST_MAPPINGS
    fs_ytd = statement_ytd(aliv_fs_path, closed_period, area) if aliv_fs_path else {}
    from automations.actuals.cbl_statement_periods import read_cbl_area

    cbl_ytd = (
        {
            canonical(k): rounded(v)
            for k, v in read_cbl_area(cbl_fs_path, area, closed_period)["YTD"].items()
        }
        if cbl_fs_path
        else {}
    )
    report = load_workbook(report_path)
    try:
        for company in ("ALIV", "CBL"):
            sheet = report[f"{company} {area}"]
            add_company_columns(sheet)
            for row in range(7, sheet.max_row + 1):
                label = canonical(sheet.cell(row, 1).value)
                routes = [
                    m
                    for m in mappings
                    if m["company"] == company and canonical(m["fs_line"]) == label
                ]
                if len(routes) != 1:
                    status_cell(sheet.cell(row, 11), "ACTION REQUIRED")
                    continue
                mapping = routes[0]
                if area == "OPEX":
                    keys = [canonical(mapping["tableau_category"])]
                elif mapping["aggregation"] == AGG_TOTAL:
                    keys = [f"{company} TOTAL"]
                else:
                    keys = [canonical(k) for k in mapping.get("tableau_categories", [])]
                normalized = {canonical(k): v for k, v in ytd[company].items()}
                # The statement value does not depend on Tableau, so it is
                # written even when the consolidated Tableau value is missing.
                _fs = (fs_ytd if company == "ALIV" else cbl_ytd).get(label)
                if _fs is not None:
                    sheet.cell(row, 8, _fs)
                if not keys or any(k not in normalized for k in keys):
                    status_cell(sheet.cell(row, 11), "ACTION REQUIRED")
                    continue
                sheet.cell(
                    row,
                    9,
                    rounded(sum(Decimal(str(normalized[k])) for k in keys) / 1000),
                )
                statement_values = fs_ytd if company == "ALIV" else cbl_ytd
                if statement_values.get(label) is not None:
                    sheet.cell(row, 8, statement_values[label])
                    difference = sheet.cell(row, 9).value - statement_values[label]
                    sheet.cell(row, 10, difference)
                    status_cell(
                        sheet.cell(row, 11), "PASS" if difference == 0 else "FAIL"
                    )
                else:
                    status_cell(sheet.cell(row, 11), "NOT VALIDATED")
                    sheet.cell(row, 11).comment = Comment(
                        "Financial statement YTD layout is not confirmed for this company.",
                        "Validation",
                    )
        for name in ("Group Actuals", "Validation Summary"):
            if name in report.sheetnames:
                del report[name]
        report.save(report_path)
    finally:
        report.close()


def merge_individual_values(report, groups):
    """Write individual captures into the capture (A:S) workbook."""
    for mapping, pair in groups:
        sheet = report[mapping["output_company_sheet"]]
        if str(sheet["A6"].value or "").strip() != "Financial Statement Line":
            raise ValueError(f"{sheet.title}: unrecognized validation layout")
        add_company_columns(sheet)
        matches = [
            r
            for r in range(7, sheet.max_row + 1)
            if canonical(sheet.cell(r, 1).value) == canonical(mapping["fs_line"])
        ]
        if len(matches) != 1:
            raise ValueError(
                f"Expected one row for {mapping['company']} / {mapping['fs_line']}; found {len(matches)}"
            )
        row = matches[0]
        month, ytd = pair
        values = {
            8: ytd[1],
            9: ytd[2],
            12: month[4],
            13: month[5],
            14: month[6],
            16: ytd[4],
            17: ytd[5],
            18: ytd[6],
        }
        for column, value in values.items():
            sheet.cell(row, column).value = value
        status_cell(sheet.cell(row, 15), month[7])
        status_cell(sheet.cell(row, 19), ytd[7])
        for col, result in ((15, month), (19, ytd)):
            if result[13]:
                sheet.cell(row, col).comment = Comment(str(result[13]), "Validation")
    for name in ("Individual Actuals", "Validation Summary", "Group Actuals"):
        if name in report.sheetnames:
            del report[name]
    refresh_summary(report)


def refresh_summary(report):
    sheet = report["Executive Summary"]
    anchor = next(
        (
            r
            for r in range(1, sheet.max_row + 1)
            if sheet.cell(r, 1).value == "ADDITIONAL VALIDATION CHECKS"
        ),
        sheet.max_row + 3,
    )
    if anchor <= sheet.max_row:
        sheet.delete_rows(anchor, sheet.max_row - anchor + 1)
    sheet.cell(anchor, 1, "ADDITIONAL VALIDATION CHECKS")
    for col, label in enumerate(
        ["Company", "Check", "Passed", "Failed", "Action Required", "Not Validated"], 1
    ):
        sheet.cell(anchor + 1, col, label)
    position = anchor + 2
    for company in ("ALIV", "CBL"):
        company_sheets = [
            s
            for s in report
            if s.title.startswith(company + " ") and s.title != "Executive Summary"
        ]
        for label, column in [
            ("Group YTD", 11),
            ("Individual Monthly", 15),
            ("Individual YTD", 19),
        ]:
            statuses = [
                s.cell(r, column).value
                for s in company_sheets
                for r in range(7, s.max_row + 1)
            ]
            for col, val in enumerate(
                [
                    company,
                    label,
                    statuses.count("PASS"),
                    statuses.count("FAIL"),
                    sum(v in ("ACTION REQUIRED", "ERROR") for v in statuses),
                    sum(
                        v not in ("PASS", "FAIL", "ACTION REQUIRED", "ERROR")
                        for v in statuses
                    ),
                ],
                1,
            ):
                sheet.cell(position, col, val)
            position += 1
    for row in (anchor, anchor + 1):
        for cell in sheet[row][:6]:
            cell.fill = PatternFill("solid", fgColor="145A70")
            cell.font = Font(bold=True, color="FFFFFF")
    exceptions = []
    pending = False
    for detail in report:
        if not detail.title.startswith(("ALIV ", "CBL ")):
            continue
        for row in range(7, detail.max_row + 1):
            for label, column in [
                ("Group YTD", 11),
                ("Individual Monthly", 15),
                ("Individual YTD", 19),
            ]:
                cell = detail.cell(row, column)
                if cell.value in ("FAIL", "ACTION REQUIRED", "ERROR"):
                    exceptions.append(
                        [
                            detail.title,
                            detail.cell(row, 1).value,
                            label,
                            cell.value,
                            (
                                cell.comment.text
                                if cell.comment
                                else "Review the variances on the company sheet."
                            ),
                        ]
                    )
                elif cell.value != "PASS":
                    pending = True
    sheet.cell(position + 1, 1, "ADDITIONAL EXCEPTIONS REQUIRING REVIEW")
    for col, label in enumerate(
        ["Company Sheet", "Statement Line", "Check", "Status", "Action / Note"], 1
    ):
        sheet.cell(position + 2, col, label)
    for offset, values in enumerate(
        exceptions
        or [
            [
                "",
                "",
                "",
                "",
                "No additional exceptions identified; unvalidated checks remain pending.",
            ]
        ]
    ):
        for col, value in enumerate(values, 1):
            cell = sheet.cell(position + 3 + offset, col, value)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for cells in sheet:
        for cell in cells:
            if isinstance(cell.value, str) and cell.value.startswith(
                "OVERALL VALIDATION STATUS:"
            ):
                if exceptions:
                    cell.value = "OVERALL VALIDATION STATUS: REVIEW REQUIRED"
                elif pending and cell.value.endswith("PASS"):
                    cell.value = "OVERALL VALIDATION STATUS: PENDING"
                if not cell.value.endswith(": PASS"):
                    cell.fill = PatternFill("solid", fgColor="FFF2CC")


def report_status(path):
    from automations.actuals.report_builder import report_status as _report_status

    return _report_status(path)


def prepare_group_copy(source, destination, tab, closed_period):
    root = ET.parse(source).getroot()
    sheet = root.find(f"./worksheets/worksheet[@name='{tab}']")
    if sheet is None:
        raise ValueError(f"Group worksheet not found: {tab}")
    start = date(
        closed_period.year + (closed_period.month == 12),
        closed_period.month % 12 + 1,
        1,
    )
    fiscal_start = date(closed_period.year - (closed_period.month < 7), 7, 1)
    fiscal_end = fiscal_start.year + 1
    names = {
        c.get("name")
        for c in sheet.iter("column")
        if c.get("caption") == "Forecast Start"
    }
    if len(names) != 1:
        raise ValueError("Expected one group Forecast Start dependency.")
    for c in root.iter("column"):
        if c.get("name") in names and "param-domain-type" in c.attrib:
            c.set("value", f"#{start:%Y-%m-%d}#")
            if c.find("calculation") is not None:
                c.find("calculation").set("formula", f"#{start:%Y-%m-%d}#")
    months = []
    current = fiscal_start
    while current < start:
        months.append(current)
        current = date(current.year + (current.month == 12), current.month % 12 + 1, 1)
    for f in root.iter("filter"):
        column = f.get("column", "")
        if column.endswith("[my:GL Posting Date:ok]"):
            for child in list(f):
                f.remove(child)
            union = ET.SubElement(f, "groupfilter", function="union")
            for month in months:
                ET.SubElement(
                    union,
                    "groupfilter",
                    function="member",
                    level="[my:GL Posting Date:ok]",
                    member=f"{month:%Y%m}",
                )
        if column.endswith("[none:Fiscal Year:ok]"):
            for child in list(f):
                f.remove(child)
            ET.SubElement(
                f,
                "groupfilter",
                function="member",
                level="[none:Fiscal Year:ok]",
                member=str(fiscal_end),
            )
        if column.endswith("[none:Forecast Type:nk]"):
            for member in f.findall('.//groupfilter[@function="member"]'):
                if re.fullmatch(
                    r"FY\d{2} Forecast_\d{6}", member.get("member", "").strip('"'), re.I
                ):
                    forecast_end = start.year + (start.month >= 7)
                    member.set(
                        "member", f'"FY{forecast_end%100:02} Forecast_{start:%m%Y}"'
                    )
    # Preserve all department, account, and Excl Media filters from the supplied TWB.
    ET.ElementTree(root).write(destination, encoding="utf-8", xml_declaration=True)


def finalize_actuals_report_in_place(
    report_path, forecast_type=None, forecast_start=None
):
    """Rebuild the capture workbook as the final report layout; safe to repeat."""
    from automations.actuals.report_builder import finalize_report

    return finalize_report(
        report_path, forecast_type=forecast_type, forecast_start=forecast_start
    )
