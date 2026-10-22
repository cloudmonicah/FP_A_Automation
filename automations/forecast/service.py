from __future__ import annotations
from shared.friendly_errors import GENERIC, PlainLanguageError, explain, friendly

from datetime import date
from pathlib import Path
import os
import re
import shutil
import xml.etree.ElementTree as ET

from openpyxl import Workbook

from .excel_service import (
    create_forecast_refresh,
    forecast_type_for,
    populate_forecast_sheet_from_audit,
)
from .mapping import FORECAST_CATEGORIES, routes_for
from .tableau_service import (
    extract_forecast_worksheets,
    find_tableau_parameter,
    open_tableau_workbook,
    set_forecast_start_parameter,
)
from shared.forecast_output_paths import (
    get_forecast_validation_output_folder,
)

from shared.forecast_validation_report import (
    merge_saved_results,
    write_forecast_validation_report,
)
from .validation_service import run_grouped_individual_validation

MONTH_ABBR = {
    1: "Jan",
    2: "Feb",
    3: "Mar",
    4: "Apr",
    5: "May",
    6: "Jun",
    7: "Jul",
    8: "Aug",
    9: "Sep",
    10: "Oct",
    11: "Nov",
    12: "Dec",
}


def default_forecast_start(forecast_month: int, forecast_year: int) -> date:
    """The 1st of the forecast month (October 2026 -> 10/1/2026); editable on the page."""
    return date(forecast_year, forecast_month, 1)


def build_forecast_period(
    fiscal_year: str,
    forecast_month: int,
    forecast_year: int,
    forecast_start: date | None = None,
) -> dict:
    """Build the single monthly period used by Excel and every Tableau source."""
    if forecast_month not in range(1, 13):
        raise ValueError("Forecast month must be between 1 and 12.")

    fiscal_year = str(fiscal_year or "").strip().upper()
    if not re.fullmatch(r"FY\d{2}", fiscal_year):
        raise ValueError("Fiscal year must look like FY27.")

    month_start = date(forecast_year, forecast_month, 1)
    # Forecast Start is entered on the page; it defaults to the forecast month.
    start = forecast_start or default_forecast_start(forecast_month, forecast_year)
    if start.day != 1:
        raise PlainLanguageError("Forecast Start must be the 1st of a month, for example 10/01/2026.")
    # Forecast Type is the selected month (October -> FY27 Forecast_102026),
    # whatever Forecast Start is set to.
    forecast_type = forecast_type_for(fiscal_year, forecast_month, forecast_year)

    return {
        "fiscal_year": fiscal_year,
        "forecast_month": forecast_month,
        "forecast_year": forecast_year,
        "forecast_start": start,
        "forecast_start_tableau": f"#{start:%Y-%m-%d}#",
        "forecast_type": forecast_type,
        "target_month_abbr": MONTH_ABBR[forecast_month],
        "display": f"{month_start:%B %Y}",
    }


def check_tableau_forecast_type(text: str, fiscal_year: str) -> str:
    """The Forecast Type chosen for Tableau, e.g. "FY27 Forecast_102026", tidied up."""
    value = " ".join(str(text or "").split())
    match = re.fullmatch(r"(FY\d{2}) Forecast_(\d{2})(\d{4})", value, re.I)
    if not match or not 1 <= int(match.group(2)) <= 12:
        raise PlainLanguageError(
            f'Tableau Forecast Type "{value}" doesn\'t look right. Type it as it appears in '
            "Tableau's Forecast Type filter, for example FY27 Forecast_102026."
        )
    return f"{match.group(1).upper()} Forecast_{match.group(2)}{match.group(3)}"


def set_consolidated_forecast_period(
    tableau_workbook_path: str | Path,
    fiscal_year: str,
    forecast_month: int,
    forecast_year: int,
    forecast_start: date | None = None,
    tableau_forecast_type: str | None = None,
) -> dict:
    """
    Set Forecast Start and the Forecast Type filter in the consolidated workbook.

    The Forecast Type is the one entered on the page as "Tableau Forecast
    Type" (default: the selected month's). It is used for every .twb -
    consolidated and individual - and never for the Excel workbooks' Forecast
    Type column, which always follows the selected month (Step 1).
    """
    path = Path(tableau_workbook_path)
    period = build_forecast_period(
        fiscal_year=fiscal_year,
        forecast_month=forecast_month,
        forecast_year=forecast_year,
        forecast_start=forecast_start,
    )
    target = check_tableau_forecast_type(
        tableau_forecast_type or period["forecast_type"], period["fiscal_year"]
    )
    period["tableau_forecast_type"] = target

    # Proven Forecast Start updater + disk verification.
    set_forecast_start_parameter(
        tableau_workbook_path=path,
        forecast_month=period["forecast_start"].month,
        forecast_year=period["forecast_start"].year,
    )

    tree = ET.parse(path)
    root = tree.getroot()
    changed = 0

    # Saved Forecast Type filter selections.
    for groupfilter in root.iter("groupfilter"):
        member = str(groupfilter.get("member") or "").strip('"')
        if re.fullmatch(r"FY\d{2} Forecast_\d{6}", member, re.I):
            groupfilter.set("member", f'"{target}"')
            changed += 1

    # Parameter-style Forecast Type, if a workbook uses one.
    for column in root.iter("column"):
        caption = re.sub(r"[^a-z0-9]+", "", str(column.get("caption") or "").casefold())
        if caption != "forecasttype":
            continue
        value = str(column.get("value") or "").strip('"')
        if re.fullmatch(r"FY\d{2} Forecast_\d{6}", value, re.I):
            column.set("value", f'"{target}"')
            calculation = column.find("calculation")
            if calculation is not None and re.search(
                r"FY\d{2} Forecast_\d{6}", str(calculation.get("formula") or ""), re.I
            ):
                calculation.set("formula", f'"{target}"')
            changed += 1

    if changed:
        tree.write(path, encoding="utf-8", xml_declaration=True)
        saved = read_consolidated_forecast_period(path)["forecast_types"]
        if saved and saved != [target]:
            raise PlainLanguageError(
                f"The Forecast Type filter in {path.name} could not be set to {target} "
                f"(it shows {', '.join(saved)}). Close Tableau and run Step 1 again."
            )

    period["forecast_type_updates"] = changed
    return period


def read_consolidated_forecast_period(tableau_workbook_path: str | Path) -> dict:
    """The Forecast Start and Forecast Type saved in the consolidated .twb."""
    path = Path(tableau_workbook_path)
    try:
        root = ET.parse(path).getroot()
    except FileNotFoundError as exc:
        raise PlainLanguageError(f"The Tableau Forecast Workbook was not found: {path}") from exc
    except ET.ParseError as exc:
        raise PlainLanguageError(
            f'The Tableau Forecast Workbook "{path.name}" could not be read; it may be '
            "damaged. Run Step 1 again, or pick the right .twb file."
        ) from exc
    parameter = find_tableau_parameter(root=root, parameter_caption="Forecast Start")
    start = parameter.attrib.get("value") if parameter is not None else None
    types = sorted(
        {
            str(item.get("member") or "").strip('"')
            for item in root.iter("groupfilter")
            if re.fullmatch(r"FY\d{2} Forecast_\d{6}", str(item.get("member") or "").strip('"'), re.I)
        }
    )
    return {"forecast_start_tableau": start, "forecast_types": types}


def _tab_key(name) -> str:
    """Tab names compare the way Tableau's menu search does: case and spacing ignored."""
    return " ".join(str(name or "").lower().split())


def consolidated_tab_problems(tableau_workbook_path: str | Path, tab_names) -> dict[str, str]:
    """{tab: plain reason} for mapped tabs the saved .twb doesn't have, or hides.

    Checked before copying, so a renamed or hidden tab is named precisely
    (with the closest real tab names) instead of failing inside Tableau.
    """
    import difflib

    root = ET.parse(Path(tableau_workbook_path)).getroot()
    sheets = {}
    for sheet in root.iter("worksheet"):
        name = sheet.get("name")
        if name:
            sheets.setdefault(_tab_key(name), name)
    hidden = {
        _tab_key(window.get("name"))
        for window in root.iter("window")
        if window.get("class") == "worksheet" and str(window.get("hidden", "")).lower() == "true"
    }
    workbook_name = Path(tableau_workbook_path).name
    problems = {}
    for tab in tab_names:
        key = _tab_key(tab)
        if key in sheets:
            if key in hidden:
                problems[tab] = (
                    f'The tab "{sheets[key]}" is hidden in {workbook_name}, so the app can\'t '
                    "open it. In Tableau, right-click the tab (or the dashboard that uses it) and "
                    "choose Unhide, save the workbook, then run Step 1 and Update again."
                )
            continue
        close = difflib.get_close_matches(key, list(sheets), n=3, cutoff=0.6)
        hint = (
            " The closest tab names are: " + ", ".join(f'"{sheets[c]}"' for c in close) + "."
            if close
            else ""
        )
        problems[tab] = (
            f'{workbook_name} has no tab called "{tab}".{hint} Rename the tab in Tableau to '
            "match, or tell support to change the mapping."
        )
    return problems


def check_consolidated_forecast_period(tableau_workbook_path: str | Path, period: dict) -> None:
    """Step 2 reads the workbook Step 1 prepared; stop if its period differs from the page.

    Step 2 does not rewrite the .twb: Tableau already has it open, and a change
    on disk would not reach the open window, so the copy would use old values.
    """
    saved = read_consolidated_forecast_period(tableau_workbook_path)
    problems = []
    if saved["forecast_start_tableau"] != period["forecast_start_tableau"]:
        shown = (saved["forecast_start_tableau"] or "not set").strip("#")
        problems.append(
            f"Forecast Start in Tableau is {shown}, but this page has "
            f"{period['forecast_start']:%Y-%m-%d}"
        )
    expected_type = period.get("tableau_forecast_type") or period["forecast_type"]
    if saved["forecast_types"] and saved["forecast_types"] != [expected_type]:
        problems.append(
            f"Forecast Type in Tableau is {', '.join(saved['forecast_types'])}, "
            f"but the page's Tableau Forecast Type is {expected_type}"
        )
    if problems:
        raise PlainLanguageError(
            "The Tableau workbook is not set to the period on this page:\n  - "
            + "\n  - ".join(problems)
            + "\n\nNothing was updated. Close Tableau, click "
            "\"1. Generate Forecast & Open Tableau\" again for this period, then Update."
        )


def prepare_monthly_forecast(
    source_folder: str | Path,
    output_folder: str | Path,
    tableau_workbook_path: str | Path,
    fiscal_year: str,
    forecast_month: int,
    forecast_year: int,
    forecast_start: date | None = None,
    overwrite_mode: str = "skip",
    open_tableau: bool = True,
    tableau_forecast_type: str | None = None,
) -> dict:
    """Stage 1 + set the consolidated Tableau period + optionally open Tableau."""
    period = build_forecast_period(
        fiscal_year=fiscal_year,
        forecast_month=forecast_month,
        forecast_year=forecast_year,
        forecast_start=forecast_start,
    )

    excel_result = create_forecast_refresh(
        source_folder=source_folder,
        output_folder=output_folder,
        fiscal_year=fiscal_year,
        forecast_month=forecast_month,
        forecast_year=forecast_year,
        overwrite_mode=overwrite_mode,
    )

    period = set_consolidated_forecast_period(
        tableau_workbook_path=tableau_workbook_path,
        fiscal_year=fiscal_year,
        forecast_month=forecast_month,
        forecast_year=forecast_year,
        forecast_start=forecast_start,
        tableau_forecast_type=tableau_forecast_type,
    )

    if open_tableau:
        open_tableau_workbook(tableau_workbook_path=tableau_workbook_path)

    return {
        "period": period,
        "excel": excel_result,
        "output_folder": str(Path(output_folder)),
    }


def _safe_sheet_title(name: str, used: set[str]) -> str:
    title = re.sub(r"[\\/*?:\[\]]", "_", str(name))[:31] or "Tableau"
    base = title
    suffix = 1
    while title.casefold() in used:
        suffix += 1
        tag = f"_{suffix}"
        title = base[: 31 - len(tag)] + tag
    used.add(title.casefold())
    return title


def write_tableau_audit_workbook(
    extracts: dict[str, list[dict]],
    audit_path: str | Path,
) -> Path:
    """
    Save the exact consolidated Tableau crosstabs into one audit workbook.
    Stage 4 already knows how to consume this workbook.
    """
    audit_path = Path(audit_path)
    audit_path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    wb.remove(wb.active)
    used: set[str] = set()

    for tableau_name, rows in extracts.items():
        ws = wb.create_sheet(_safe_sheet_title(tableau_name, used))

        headers = []
        for row in rows:
            for key in row:
                if key not in headers:
                    headers.append(key)

        if headers:
            ws.append(headers)
            for row in rows:
                ws.append([row.get(header) for header in headers])

    if not wb.sheetnames:
        wb.create_sheet("No Data")

    wb.save(audit_path)
    wb.close()
    return audit_path


def _resolve_target_filename(
    baseline_name: str,
    target_month_abbr: str,
) -> str:
    resolved, count = re.subn(
        r"_[A-Za-z]{3}\s+Forecast\.xlsx$",
        f"_{target_month_abbr} Forecast.xlsx",
        str(baseline_name),
        count=1,
        flags=re.I,
    )
    if count != 1:
        raise ValueError(f"Could not resolve mapped workbook name: {baseline_name}")
    return resolved


def _make_sure_tableau_is_open(tableau_workbook_path) -> None:
    """Open the consolidated workbook if step 1 didn't (or it was closed)."""
    from .tableau_service import _get_tableau_window, wait_for_tableau_workbook

    try:
        _get_tableau_window(
            max_attempts=1,
            retry_wait_seconds=0,
            tableau_workbook_path=tableau_workbook_path,
        )
        return
    except Exception:
        pass
    print("The consolidated Tableau workbook is not open; opening it now...")
    open_tableau_workbook(tableau_workbook_path=tableau_workbook_path)
    wait_for_tableau_workbook(tableau_workbook_path)


def _reason(error) -> str:
    """A specific, plain reason. Never the bare "something went wrong"."""
    text = friendly(error)
    if text != GENERIC:
        return text
    raw = str(error).strip().splitlines()[0] if str(error).strip() else type(error).__name__
    raw = re.sub(r"^\w+(Error|Exception): ", "", raw)
    return f"an unexpected problem stopped it ({raw[:200]}). The details are in the log."


def _paste_route_safely(target: Path, excel_sheet: str, audit_path: Path,
                        tableau_worksheet: str, fiscal_year: str) -> tuple[dict, bool]:
    """Paste one Tableau tab into one sheet without ever leaving it half-updated.

    The paste is done on a copy, which replaces the workbook only when it is
    complete. An unexpected failure is tried once more; returns (result, retried).
    Problems the user has to fix (file open, sheet missing, layout) are not retried.
    """
    if not target.is_file():
        raise PlainLanguageError(
            f'The workbook "{target.name}" is not in the Forecast Output Folder. '
            "Run Step 1 first, or check that the folder is this month's."
        )
    working = target.with_name(f"~updating {target.name}")
    last_error = None
    for attempt in (1, 2):
        try:
            shutil.copy2(target, working)
            result = populate_forecast_sheet_from_audit(
                workbook_path=working,
                excel_sheet_name=excel_sheet,
                audit_workbook_path=audit_path,
                tableau_worksheet_name=tableau_worksheet,
                fiscal_year=fiscal_year,
            )
            try:
                os.replace(working, target)
            except PermissionError as exc:
                raise PlainLanguageError(
                    f'"{target.name}" is open in Excel (or another program), so it '
                    "could not be updated. Close it and click Update again."
                ) from exc
            return result, attempt > 1
        except (PlainLanguageError, PermissionError) as exc:
            if isinstance(exc, PermissionError):
                raise PlainLanguageError(
                    f'"{target.name}" is open in Excel (or another program), so it '
                    "could not be updated. Close it and click Update again."
                ) from exc
            raise
        except Exception as exc:
            last_error = exc
            print(f"[STEP 2] {target.name} / {excel_sheet}: attempt {attempt} failed: "
                  f"{type(exc).__name__}: {exc}", flush=True)
        finally:
            try:
                working.unlink(missing_ok=True)
            except OSError:
                pass
    raise last_error


def _populate_monthly_forecast_from_tableau(
    forecast_folder: str | Path,
    tableau_workbook_path: str | Path,
    fiscal_year: str,
    forecast_month: int,
    forecast_year: int,
    forecast_start: date | None = None,
    categories: list[str] | None = None,
    tableau_forecast_type: str | None = None,
) -> dict:
    """
    Step 2 for one category (OPEX, Direct Cost, Revenue, Other) or all.

    Capture the consolidated Tableau crosstabs once and run the locked Stage 4
    population/finalization logic against every production mapping route.
    """
    folder = Path(forecast_folder)
    period = build_forecast_period(
        fiscal_year=fiscal_year,
        forecast_month=forecast_month,
        forecast_year=forecast_year,
        forecast_start=forecast_start,
    )

    # Use the workbook Step 1 prepared and opened; never rewrite it here.
    period["tableau_forecast_type"] = check_tableau_forecast_type(
        tableau_forecast_type or period["forecast_type"], period["fiscal_year"]
    )
    check_consolidated_forecast_period(tableau_workbook_path, period)

    routes = routes_for(categories)
    worksheet_names = [route["tableau_worksheet"] for route in routes]

    # Tabs missing from (or hidden in) the saved workbook are reported now,
    # with the closest real names, and not tried in Tableau.
    tab_problems = consolidated_tab_problems(tableau_workbook_path, worksheet_names)
    worksheet_names = [name for name in worksheet_names if name not in tab_problems]

    _make_sure_tableau_is_open(tableau_workbook_path)

    # Everything Step 2 produces besides the workbooks goes to the month's
    # Forecast Validation folder, never the Forecast Output Folder (that one
    # is published to the departments and holds only their workbooks).
    tag = "All" if not categories else "_".join(c.replace(" ", "") for c in categories)
    validation_folder = get_forecast_validation_output_folder(
        fiscal_year=period["fiscal_year"],
        forecast_month=period["forecast_month"],
        forecast_year=period["forecast_year"],
        create=True,
    )
    raw_folder = validation_folder / "Tableau Raw Copies"
    extract_folder = validation_folder / "Tableau Extracts"

    extracts, extract_errors = (
        extract_forecast_worksheets(
            tableau_workbook_path=tableau_workbook_path,
            worksheet_names=worksheet_names,
            raw_output_folder=raw_folder,
        )
        if worksheet_names
        else ({}, {})
    )

    # A tab that failed to copy is tried once more, after the others.
    retried_tabs = set()
    if extract_errors:
        print(f"Trying {len(extract_errors)} Tableau tab(s) again: {', '.join(extract_errors)}")
        again, still_failing = extract_forecast_worksheets(
            tableau_workbook_path=tableau_workbook_path,
            worksheet_names=list(extract_errors),
            raw_output_folder=raw_folder,
        )
        extracts.update(again)
        retried_tabs.update(again)
        extract_errors = {name: still_failing.get(name, error) for name, error in extract_errors.items() if name not in again}

    # One extract file per category, so running Revenue never overwrites the
    # OPEX extract.
    audit_path = extract_folder / (
        f"Tableau_Forecast_Extract_{forecast_year}-{forecast_month:02d}_{tag}.xlsx"
    )
    write_tableau_audit_workbook(extracts, audit_path)

    passed = []
    failed = []

    for route in routes:
        target_name = _resolve_target_filename(
            route["workbook"],
            period["target_month_abbr"],
        )
        target = folder / target_name
        tab = route["tableau_worksheet"]
        base = {"workbook": target_name, "worksheet": route["excel_sheet"], "tableau_worksheet": tab}

        if tab in tab_problems:
            failed.append({**base, "error": tab_problems[tab]})
            continue

        # Do not touch the Excel sheet if its Tableau source failed.
        if tab in extract_errors:
            failed.append({**base, "error": f'Tableau tab "{tab}" could not be copied: '
                           + _reason(extract_errors[tab])})
            continue

        try:
            result, retried = _paste_route_safely(
                target=target,
                excel_sheet=route["excel_sheet"],
                audit_path=audit_path,
                tableau_worksheet=tab,
                fiscal_year=fiscal_year,
            )
        except Exception as exc:
            failed.append({**base, "error": explain(exc) if isinstance(exc, PlainLanguageError) else _reason(exc)})
            continue

        passed.append(
            {
                **base,
                "rows": result.get("tableau_detail_rows", 0),
                "months_pasted": result.get("months_pasted", []),
                "months_blank_in_tableau": result.get("months_blank_in_tableau", []),
                "comments_dropped": result.get("comments_dropped", []),
                "unmerged": result.get("unmerged", []),
                "date_formats_fixed": result.get("date_formats_fixed", 0),
                # Something went wrong once; it worked on the second try. Worth a look.
                "check_manually": retried or tab in retried_tabs,
            }
        )

    pasted = []
    for item in passed:
        for month in item.get("months_pasted", []):
            if month not in pasted:
                pasted.append(month)

    return {
        "period": period,
        "categories": list(categories or FORECAST_CATEGORIES),
        "months_pasted": ", ".join(pasted),
        "audit_workbook": str(audit_path),
        "validation_folder": str(validation_folder),
        "passed": passed,
        "failed": failed,
        "routes_processed": len(passed) + len(failed),
        "routes_passed": len(passed),
        "routes_failed": len(failed),
        "workbooks_updated": len({item["workbook"] for item in passed}),
        "check_manually": [item for item in passed if item["check_manually"]],
        "worksheets_updated": len(passed),
    }


def validate_monthly_forecast(
    forecast_folder: str | Path,
    fiscal_year: str,
    forecast_month: int,
    forecast_year: int,
    forecast_start: date | None = None,
    tolerance: float = 1000.00,
    categories: list[str] | None = None,
    tableau_forecast_type: str | None = None,
) -> dict:
    """Step 3 for one category (OPEX, Direct Cost, Revenue, Other) or all.

    The individual Tableau workbooks use ``tableau_forecast_type`` - the same
    Forecast Type chosen in the consolidated workbook (default: the month's).

    The month's report keeps the latest results of every category: running
    Revenue replaces Revenue's rows and leaves the other categories alone.
    """
    period = build_forecast_period(
        fiscal_year=fiscal_year,
        forecast_month=forecast_month,
        forecast_year=forecast_year,
        forecast_start=forecast_start,
    )

    if tableau_forecast_type:
        period["forecast_type"] = check_tableau_forecast_type(tableau_forecast_type, period["fiscal_year"])

    validation_output_folder = get_forecast_validation_output_folder(
        fiscal_year=period["fiscal_year"],
        forecast_month=period["forecast_month"],
        forecast_year=period["forecast_year"],
        create=True,
    )

    result = run_grouped_individual_validation(
        routes=routes_for(categories),
        forecast_folder=forecast_folder,
        target_month_abbr=period["target_month_abbr"],
        forecast_month=forecast_month,
        forecast_year=forecast_year,
        tolerance=tolerance,
        forecast_start=period["forecast_start"],
        forecast_type=period["forecast_type"],
        raw_output_folder=validation_output_folder / "Tableau Raw Copies" / "Step 3",
    )

    from datetime import datetime as _dt

    stamp = _dt.now().strftime("%m/%d/%Y %I:%M %p")
    for item in result["results"]:
        item["validated_at"] = stamp

    all_results = merge_saved_results(
        output_folder=validation_output_folder,
        forecast_month=period["forecast_month"],
        forecast_year=period["forecast_year"],
        forecast_type=period["forecast_type"],
        categories_run=list(categories or FORECAST_CATEGORIES),
        results=result["results"],
    )

    report_path = write_forecast_validation_report(
        validation_result={**result, "results": all_results},
        output_folder=validation_output_folder,
        fiscal_year=period["fiscal_year"],
        forecast_month=period["forecast_month"],
        forecast_year=period["forecast_year"],
        forecast_type=period["forecast_type"],
        tolerance=tolerance,
        forecast_start=period["forecast_start"],
    )

    result["categories"] = list(categories or FORECAST_CATEGORIES)
    result["period"] = period
    result["validation_output_folder"] = str(validation_output_folder)
    result["validation_report"] = str(report_path)
    # OneDrive keeps a second copy ("...-<PC name>.xlsx") when the report was
    # open elsewhere while it was saved; Teams may then show the old version.
    result["conflict_copies"] = sorted(
        item.name for item in Path(report_path).parent.glob(f"{Path(report_path).stem}-*.xlsx")
    )
    return result


def populate_monthly_forecast_from_tableau(
    forecast_folder: str | Path,
    tableau_workbook_path: str | Path,
    fiscal_year: str,
    forecast_month: int,
    forecast_year: int,
    forecast_start: date | None = None,
    categories: list[str] | None = None,
    tableau_forecast_type: str | None = None,
) -> dict:
    """Step 2, with a run log saved in the month's Forecast Validation\\Logs folder."""
    import logging
    from datetime import datetime

    tag = "All" if not categories else ", ".join(categories)
    handler = None
    try:
        log_folder = get_forecast_validation_output_folder(
            fiscal_year=fiscal_year,
            forecast_month=forecast_month,
            forecast_year=forecast_year,
            create=True,
        ) / "Logs"
        log_folder.mkdir(parents=True, exist_ok=True)
        log_path = log_folder / f"Step 2 - {tag} - {datetime.now():%Y-%m-%d %H%M%S}.log"
        handler = logging.FileHandler(log_path, encoding="utf-8")
        handler.setFormatter(logging.Formatter("%(asctime)s %(levelname)-7s %(name)s: %(message)s"))
        handler.setLevel(logging.INFO)
        logging.getLogger().addHandler(handler)
    except Exception:
        handler = None  # the inner run reports a missing folder in plain words

    root = logging.getLogger()
    previous_level = root.level
    if handler is not None and (previous_level == logging.NOTSET or previous_level > logging.INFO):
        root.setLevel(logging.INFO)
    try:
        result = _populate_monthly_forecast_from_tableau(
            forecast_folder=forecast_folder,
            tableau_workbook_path=tableau_workbook_path,
            fiscal_year=fiscal_year,
            forecast_month=forecast_month,
            forecast_year=forecast_year,
            forecast_start=forecast_start,
            categories=categories,
            tableau_forecast_type=tableau_forecast_type,
        )
        logging.getLogger("step2").info(
            "%s: %s sheets updated, %s not updated",
            tag, result["routes_passed"], result["routes_failed"],
        )
        for item in result["failed"]:
            logging.getLogger("step2").warning("Not updated: %s / %s: %s",
                                               item["workbook"], item["worksheet"], item["error"])
        for item in result["passed"]:
            if item.get("unmerged"):
                logging.getLogger("step2").info("Unmerged cells: %s / %s: %s",
                                                item["workbook"], item["worksheet"], ", ".join(item["unmerged"]))
            if item.get("months_blank_in_tableau"):
                logging.getLogger("step2").info("No figures in Tableau (left empty): %s / %s: %s",
                                                item["workbook"], item["worksheet"],
                                                ", ".join(item["months_blank_in_tableau"]))
            if item.get("date_formats_fixed"):
                logging.getLogger("step2").info("Amounts shown as dates fixed: %s / %s: %s",
                                                item["workbook"], item["worksheet"], item["date_formats_fixed"])
        for item in result.get("check_manually", []):
            logging.getLogger("step2").warning("Check by hand (retried): %s / %s",
                                               item["workbook"], item["worksheet"])
        if handler is not None:
            result["run_log"] = str(log_path)
        return result
    finally:
        if handler is not None:
            root.removeHandler(handler)
            root.setLevel(previous_level)
            handler.close()
