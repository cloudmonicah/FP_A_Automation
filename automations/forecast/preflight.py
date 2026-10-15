"""Quick checks before steps 2 and 3 open Tableau.

Reading the files takes seconds; finding the same problem through Tableau
takes minutes. Nothing here changes any file. Each problem names the route and
says what to fix; the person decides whether to continue (problem routes are
then reported as failed or not validated by the real run).
"""

from __future__ import annotations

import xml.etree.ElementTree as ET
from datetime import date, datetime
from functools import lru_cache
from pathlib import Path

from openpyxl import load_workbook

from .validation_service import _norm, describe_forecast_type, resolve_forecast_workbook_name


def _problem(route, problem):
    return {
        "category": route.get("category", ""),
        "workbook": route.get("workbook", ""),
        "worksheet": route.get("excel_sheet", ""),
        "problem": problem,
    }


@lru_cache(maxsize=64)
def _twb_info(path: str, mtime: float):
    """Worksheet names and parameter captions of a .twb (cached per file version)."""
    root = ET.parse(path).getroot()
    sheets = {_norm(sheet.get("name")) for sheet in root.findall("./worksheets/worksheet")}
    parameters = {
        _norm(column.get("caption") or column.get("name"))
        for column in root.iter("column")
        if "param-domain-type" in column.attrib
    }
    return sheets, parameters


def _read_twb(path: Path):
    return _twb_info(str(path), path.stat().st_mtime)


def _sheet_facts(path: Path) -> dict:
    """Per sheet (normalized name): Forecast Type values and count of month-date headers."""
    facts = {}
    wb = load_workbook(path, read_only=True, data_only=False)
    try:
        for ws in wb.worksheets:
            info = {"title": ws.title, "forecast_types": set(), "month_columns": 0, "header": False}
            header_col = None
            for row in ws.iter_rows(values_only=True):
                if header_col is None:
                    for index, value in enumerate(row):
                        if _norm(value) == "forecast type":
                            header_col = index
                            info["header"] = True
                            info["month_columns"] = sum(
                                isinstance(item, (datetime, date)) for item in row
                            )
                            break
                    continue
                if header_col < len(row) and row[header_col] not in (None, ""):
                    info["forecast_types"].add(str(row[header_col]).strip())
            facts[_norm(ws.title)] = info
    finally:
        wb.close()
    return facts


def _check_excel(route, folder: Path, period: dict, need_months: bool, cache: dict) -> str | None:
    name = resolve_forecast_workbook_name(route["workbook"], period["target_month_abbr"])
    path = folder / name
    if not path.is_file():
        return f"{name} is not in the Forecast Output Folder. Run step 1 for this month."
    if path not in cache:
        try:
            cache[path] = _sheet_facts(path)
        except Exception:
            cache[path] = None
    facts = cache[path]
    if facts is None:
        return f"{name} could not be opened. It may be open in Excel with unsaved changes, or damaged."
    sheet = facts.get(_norm(route["excel_sheet"]))
    if sheet is None:
        return f'{name} has no sheet named "{route["excel_sheet"]}".'
    if not sheet["header"]:
        return f'{name} / {sheet["title"]} has no "Forecast Type" column.'
    wrong = sorted(t for t in sheet["forecast_types"] if _norm(t) != _norm(period["forecast_type"]))
    if wrong:
        return (
            f'{name} / {sheet["title"]} is for {describe_forecast_type(wrong[0])}, not '
            f'{period["display"]}. Run step 1 with Existing Files = replace.'
        )
    if need_months and sheet["month_columns"] != 12:
        return (
            f'{name} / {sheet["title"]} does not have its 12 month dates yet. Run step 2 '
            "(Update from Tableau) for this category first."
        )
    return None


def check_before_update(routes, output_folder, tableau_workbook, period) -> list[dict]:
    """Problems that would make step 2 fail for some routes."""
    problems = []
    folder = Path(output_folder)
    workbook = Path(tableau_workbook)
    try:
        sheets, parameters = _read_twb(workbook)
    except Exception:
        return [_problem({}, f"The Tableau workbook {workbook.name} could not be read.")]
    if "forecast start" not in parameters:
        problems.append(
            _problem({}, f'{workbook.name} has no "Forecast Start" parameter. Check that the '
                         "FY Forecast Tableau workbook is selected.")
        )
    cache: dict = {}
    for route in routes:
        if _norm(route["tableau_worksheet"]) not in sheets:
            problems.append(
                _problem(route, f'{workbook.name} has no tab called "{route["tableau_worksheet"]}".')
            )
            continue
        issue = _check_excel(route, folder, period, need_months=False, cache=cache)
        if issue:
            problems.append(_problem(route, issue))
    return problems


def check_before_validation(routes, output_folder, period) -> list[dict]:
    """Problems that would make step 3 report routes as not validated or action required.

    Routes not set up for validation are skipped here; the report lists them.
    """
    problems = []
    folder = Path(output_folder)
    cache: dict = {}
    for route in routes:
        if route.get("validation_status") != "confirmed":
            continue
        issue = _check_excel(route, folder, period, need_months=True, cache=cache)
        if issue:
            problems.append(_problem(route, issue))
            continue
        source = Path(route.get("individual_tableau_folder") or "") / (route.get("individual_tableau_workbook") or "")
        if not source.is_file():
            problems.append(_problem(route, f"The individual Tableau workbook was not found: {source}"))
            continue
        try:
            sheets, _ = _read_twb(source)
        except Exception:
            problems.append(_problem(route, f"The individual Tableau workbook {source.name} could not be read."))
            continue
        if _norm(route.get("individual_tableau_worksheet")) not in sheets:
            problems.append(
                _problem(route, f'{source.name} has no tab called "{route["individual_tableau_worksheet"].strip()}".')
            )
    return problems
