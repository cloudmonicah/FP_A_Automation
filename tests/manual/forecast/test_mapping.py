from pathlib import Path
from collections import Counter

from openpyxl import load_workbook

from automations.forecast.mapping import FORECAST_ROUTES

# ============================================================
# CONFIG
# ============================================================

# Change this only if your October test folder is somewhere else.
FORECAST_FOLDER = Path(r"C:\Users\moncloud\OneDrive - Aliv\Monicah\ForecastTest")

TARGET_MONTH = "Oct"


# ============================================================
# HELPERS
# ============================================================

MONTHS = (
    "Jan",
    "Feb",
    "Mar",
    "Apr",
    "May",
    "Jun",
    "Jul",
    "Aug",
    "Sep",
    "Oct",
    "Nov",
    "Dec",
)


def normalize(value):
    return " ".join(str(value or "").strip().lower().split())


def baseline_to_target_filename(filename, target_month):
    """
    Converts:
        FY27 OPEX Legal_Aug Forecast.xlsx

    to:
        FY27 OPEX Legal_Oct Forecast.xlsx
    """
    result = filename

    for month in MONTHS:
        marker = f"_{month} Forecast"
        if marker.lower() in result.lower():
            index = result.lower().index(marker.lower())

            result = (
                result[:index]
                + f"_{target_month} Forecast"
                + result[index + len(marker) :]
            )

            break

    return result


# ============================================================
# TEST
# ============================================================


def main():

    print()
    print("=" * 72)
    print("FORECAST MAPPING VALIDATION")
    print("=" * 72)

    if not FORECAST_FOLDER.exists():
        raise FileNotFoundError(f"Forecast folder does not exist:\n{FORECAST_FOLDER}")

    actual_files = {
        file.name.lower(): file
        for file in FORECAST_FOLDER.glob("*.xlsx")
        if not file.name.startswith("~$")
    }

    print(f"\nForecast folder: {FORECAST_FOLDER}")
    print(f"Excel workbooks found: {len(actual_files)}")
    print(f"Mapping routes found: {len(FORECAST_ROUTES)}")

    errors = []
    warnings = []

    mapped_target_files = set()

    confirmed_count = 0
    needs_filter_count = 0
    pending_count = 0

    # --------------------------------------------------------
    # Validate every route
    # --------------------------------------------------------

    print()
    print("-" * 72)
    print("CHECKING ROUTES")
    print("-" * 72)

    for index, route in enumerate(FORECAST_ROUTES, start=1):

        tableau_worksheet = route.get("tableau_worksheet")
        baseline_workbook = route.get("workbook")
        excel_sheet = route.get("excel_sheet")
        validation_status = route.get(
            "validation_status",
            "pending",
        )

        if validation_status == "confirmed":
            confirmed_count += 1

        elif validation_status == "needs_filter":
            needs_filter_count += 1

        else:
            pending_count += 1

        if not tableau_worksheet:
            errors.append(f"Route {index}: missing tableau_worksheet")

        if not baseline_workbook:
            errors.append(f"Route {index}: missing workbook")
            continue

        if not excel_sheet:
            errors.append(f"Route {index}: missing excel_sheet")
            continue

        target_filename = baseline_to_target_filename(
            baseline_workbook,
            TARGET_MONTH,
        )

        mapped_target_files.add(target_filename.lower())

        actual_path = actual_files.get(target_filename.lower())

        if actual_path is None:

            errors.append(
                "WORKBOOK NOT FOUND | "
                f"{target_filename} | "
                f"Tableau: {tableau_worksheet}"
            )

            print(f"[FAIL] {target_filename}" f" -> workbook not found")

            continue

        try:

            workbook = load_workbook(
                actual_path,
                read_only=True,
                data_only=False,
            )

            actual_sheet_lookup = {
                normalize(sheet_name): sheet_name for sheet_name in workbook.sheetnames
            }

            expected_sheet = normalize(excel_sheet)

            if expected_sheet not in actual_sheet_lookup:

                errors.append(
                    "SHEET NOT FOUND | "
                    f"{target_filename} | "
                    f"Expected: {excel_sheet} | "
                    f"Available: {workbook.sheetnames}"
                )

                print(
                    f"[FAIL] {target_filename}" f" -> sheet '{excel_sheet}' not found"
                )

            else:

                actual_sheet = actual_sheet_lookup[expected_sheet]

                print(
                    f"[PASS] {target_filename}"
                    f" -> {actual_sheet}"
                    f" [{validation_status}]"
                )

            workbook.close()

        except Exception as exc:

            errors.append("OPEN ERROR | " f"{target_filename} | " f"{exc}")

    # --------------------------------------------------------
    # Find Excel workbooks with no mapping
    # --------------------------------------------------------

    print()
    print("-" * 72)
    print("CHECKING FOR UNMAPPED EXCEL WORKBOOKS")
    print("-" * 72)

    unmapped_files = []

    for filename_lower, path in sorted(actual_files.items()):

        if filename_lower not in mapped_target_files:

            unmapped_files.append(path.name)

            print(f"[UNMAPPED] {path.name}")

    if not unmapped_files:
        print("[PASS] Every Excel workbook has a mapping.")

    # --------------------------------------------------------
    # Duplicate exact route check
    # --------------------------------------------------------

    print()
    print("-" * 72)
    print("CHECKING DUPLICATE ROUTES")
    print("-" * 72)

    route_keys = []

    for route in FORECAST_ROUTES:

        route_keys.append(
            (
                normalize(route.get("tableau_worksheet")),
                normalize(route.get("workbook")),
                normalize(route.get("excel_sheet")),
            )
        )

    duplicate_keys = [key for key, count in Counter(route_keys).items() if count > 1]

    if duplicate_keys:

        for key in duplicate_keys:

            errors.append(f"DUPLICATE ROUTE | {key}")

            print(f"[FAIL] Duplicate route: {key}")

    else:

        print("[PASS] No exact duplicate routes.")

    # --------------------------------------------------------
    # Check duplicate Tableau worksheet names
    # --------------------------------------------------------

    print()
    print("-" * 72)
    print("CHECKING SHARED CONSOLIDATED TABLEAU WORKSHEETS")
    print("-" * 72)

    tableau_counts = Counter(
        normalize(route.get("tableau_worksheet")) for route in FORECAST_ROUTES
    )

    shared_tableau = {
        name: count for name, count in tableau_counts.items() if name and count > 1
    }

    if shared_tableau:

        for worksheet, count in shared_tableau.items():

            warnings.append(
                "SHARED TABLEAU WORKSHEET | " f"{worksheet} | " f"{count} routes"
            )

            print(f"[WARNING] {worksheet}" f" -> used by {count} routes")

    else:

        print("[PASS] Every consolidated Tableau worksheet " "maps to one route.")

    # --------------------------------------------------------
    # Validation configuration
    # --------------------------------------------------------

    print()
    print("-" * 72)
    print("VALIDATION CONFIGURATION")
    print("-" * 72)

    for route in FORECAST_ROUTES:

        status = route.get(
            "validation_status",
            "pending",
        )

        if status == "needs_filter":

            print(
                "[NEEDS FILTER] "
                f"{route['workbook']} | "
                f"{route['excel_sheet']} | "
                f"{route.get('individual_tableau_workbook')} | "
                f"{route.get('individual_tableau_worksheet')} | "
                f"{route.get('validation_filter_note', '')}"
            )

        elif status == "pending":

            print("[PENDING] " f"{route['workbook']} | " f"{route['excel_sheet']}")

    # --------------------------------------------------------
    # Summary
    # --------------------------------------------------------

    print()
    print("=" * 72)
    print("SUMMARY")
    print("=" * 72)

    print(f"Excel workbooks found:       {len(actual_files)}")

    print(f"Forecast routes:             {len(FORECAST_ROUTES)}")

    print(f"Confirmed validation routes: {confirmed_count}")

    print(f"Needs-filter routes:         {needs_filter_count}")

    print(f"Pending validation routes:   {pending_count}")

    print(f"Unmapped Excel workbooks:    {len(unmapped_files)}")

    print(f"Errors:                      {len(errors)}")

    print(f"Warnings:                    {len(warnings)}")

    if errors:

        print()
        print("=" * 72)
        print("ERROR DETAILS")
        print("=" * 72)

        for error in errors:
            print(f"- {error}")

    if warnings:

        print()
        print("=" * 72)
        print("WARNING DETAILS")
        print("=" * 72)

        for warning in warnings:
            print(f"- {warning}")

    print()

    if errors:

        print("STAGE 2 RESULT: FAIL")

        raise AssertionError(
            f"Mapping validation failed with " f"{len(errors)} error(s)."
        )

    print("STAGE 2 RESULT: PASS")


if __name__ == "__main__":
    main()
