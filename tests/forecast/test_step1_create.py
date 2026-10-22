"""Step 1: create the month's forecast workbooks from last month's files."""

import pytest
from openpyxl import Workbook

from automations.forecast.excel_service import ForecastFilesExistError, create_forecast_refresh
from shared.friendly_errors import GENERIC, friendly


def _workbook(path, forecast_type="FY27 Forecast_092026"):
    wb = Workbook()
    ws = wb.active
    ws.title = "ALIV"
    ws.append(["Forecast Type", "GL Account", "Jul", "Grand Total"])
    ws.append([forecast_type, "6000", 5, "=SUM(C2:C2)"])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


@pytest.fixture
def folders(tmp_path):
    source, out = tmp_path / "Sep", tmp_path / "Oct"
    _workbook(source / "FY27 OPEX Legal_Sep Forecast.xlsx")
    _workbook(source / "FY27 OPEX HR_Sep Forecast.xlsx")
    _workbook(source / "FY27 OPEX IT_Aug Forecast.xlsx", "FY27 Forecast_082026")
    return source, out


def _run(source, out, mode="skip"):
    return create_forecast_refresh(source, out, "FY27", 10, 2026, overwrite_mode=mode)


def test_creates_every_family_and_reports_older_sources(folders):
    source, out = folders
    result = _run(source, out)
    assert result["files_created"] == 3 and not result["failures"]
    assert sorted(p.name for p in out.iterdir()) == [
        "FY27 OPEX HR_Oct Forecast.xlsx",
        "FY27 OPEX IT_Oct Forecast.xlsx",
        "FY27 OPEX Legal_Oct Forecast.xlsx",
    ]
    assert [item["family"] for item in result["older_sources"]] == ["FY27 OPEX IT"]


def test_cancel_changes_nothing_when_files_exist(folders):
    source, out = folders
    _run(source, out)
    legal = out / "FY27 OPEX Legal_Oct Forecast.xlsx"
    (out / "FY27 OPEX HR_Oct Forecast.xlsx").unlink()
    before = legal.read_bytes()
    with pytest.raises(ForecastFilesExistError) as error:
        _run(source, out, "cancel")
    assert "2 of this month's forecast files already exist" in str(error.value)
    assert not (out / "FY27 OPEX HR_Oct Forecast.xlsx").exists()  # nothing was created
    assert legal.read_bytes() == before
    assert friendly(error.value) == str(error.value)  # shown as written


def test_skip_keeps_existing_files(folders):
    source, out = folders
    _run(source, out)
    result = _run(source, out, "skip")
    assert result["files_created"] == 0 and len(result["skipped_files"]) == 3


@pytest.mark.parametrize(
    "error, expected",
    [
        (ValueError("July is the initial fiscal-year forecast month. Use Initial Forecast Month Setup to create July."),
         "Initial Budget Setup"),
        (FileNotFoundError(r"No valid FY27 forecast workbooks were found before Oct in C:\Forecast\Sep"),
         r"C:\Forecast\Sep"),
        (FileNotFoundError(r"Forecast source folder not found: C:\Forecast\Sep"),
         r"Previous Forecast Folder was not found: C:\Forecast\Sep"),
    ],
)
def test_step1_errors_are_explained(error, expected):
    message = friendly(error)
    assert message != GENERIC and expected in message


# --- Hardening: no half-made files, macros kept, problems reported ----------

import zipfile

from openpyxl import load_workbook

from automations.forecast.excel_service import find_header


def _leftovers(folder):
    return [p.name for p in folder.iterdir() if p.name.startswith("~building")]


def test_a_failed_workbook_leaves_nothing_behind(folders):
    source, out = folders
    (source / "FY27 OPEX Bad_Sep Forecast.xlsx").write_bytes(b"not a workbook")
    result = _run(source, out)
    assert [item["file"] for item in result["failures"]] == ["FY27 OPEX Bad_Sep Forecast.xlsx"]
    assert not (out / "FY27 OPEX Bad_Oct Forecast.xlsx").exists()
    assert _leftovers(out) == []
    # A later "skip" run must not find a broken file to keep.
    assert _run(source, out)["files_skipped"] == 3


def test_replace_keeps_last_runs_file_if_the_new_one_fails(folders):
    source, out = folders
    _run(source, out)
    kept = out / "FY27 OPEX Legal_Oct Forecast.xlsx"
    before = kept.read_bytes()
    (source / "FY27 OPEX Legal_Sep Forecast.xlsx").write_bytes(b"broken")
    result = _run(source, out, "replace")
    assert [item["file"] for item in result["failures"]] == ["FY27 OPEX Legal_Sep Forecast.xlsx"]
    assert kept.read_bytes() == before
    assert _leftovers(out) == []


def test_replace_recreates_the_file(folders):
    source, out = folders
    _run(source, out)
    target = out / "FY27 OPEX HR_Oct Forecast.xlsx"
    wb = load_workbook(target)
    wb.active["B2"] = "edited"
    wb.save(target)
    result = _run(source, out, "replace")
    assert result["files_created"] == 3
    assert load_workbook(target).active["B2"].value == "6000"


def test_forecast_type_is_set_to_the_previous_month(folders):
    source, out = folders
    _run(source, out)
    ws = load_workbook(out / "FY27 OPEX Legal_Oct Forecast.xlsx").active
    assert ws["A2"].value == "FY27 Forecast_092026"
    assert ws["D2"].value == "=SUM(C2:C2)"


def _xlsm(path):
    """A .xlsm whose package carries a (dummy) VBA project."""
    _workbook(path.with_suffix(".xlsx"))
    with zipfile.ZipFile(path.with_suffix(".xlsx")) as src, zipfile.ZipFile(path, "w") as dst:
        for item in src.infolist():
            data = src.read(item)
            if item.filename == "[Content_Types].xml":
                data = data.replace(
                    b"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml",
                    b"application/vnd.ms-excel.sheet.macroEnabled.main+xml",
                ).replace(
                    b"</Types>",
                    b'<Default Extension="bin" ContentType="application/vnd.ms-office.vbaProject"/></Types>',
                )
            dst.writestr(item, data)
        dst.writestr("xl/vbaProject.bin", b"VBA-PROJECT")
    path.with_suffix(".xlsx").unlink()
    return path


def test_macro_workbooks_keep_their_macros(tmp_path):
    source, out = tmp_path / "Sep", tmp_path / "Oct"
    _xlsm(source / "FY27 OPEX Macro_Sep Forecast.xlsm")
    result = create_forecast_refresh(source, out, "FY27", 10, 2026)
    assert result["files_created"] == 1 and not result["failures"]
    with zipfile.ZipFile(out / "FY27 OPEX Macro_Oct Forecast.xlsm") as archive:
        assert archive.read("xl/vbaProject.bin") == b"VBA-PROJECT"


def test_xlsb_files_fail_with_a_plain_message(folders):
    source, out = folders
    (source / "FY27 OPEX Binary_Sep Forecast.xlsb").write_bytes(b"binary workbook")
    result = _run(source, out)
    (failure,) = result["failures"]
    assert "save it as .xlsx" in failure["error"]
    assert result["files_created"] == 3


def test_a_workbook_without_forecast_type_is_flagged(folders):
    source, out = folders
    wb = Workbook()
    wb.active.append(["GL Account", "Jul"])
    wb.save(source / "FY27 OPEX Plain_Sep Forecast.xlsx")
    result = _run(source, out)
    assert result["files_created"] == 4
    assert result["warnings"] == [
        {"file": "FY27 OPEX Plain_Oct Forecast.xlsx",
         "problems": ["no Forecast Type cells were found to update"]}
    ]


def test_clean_run_has_no_warnings(folders):
    source, out = folders
    assert _run(source, out)["warnings"] == []


def test_header_search_stops_near_the_top():
    wb = Workbook()
    ws = wb.active
    ws.cell(row=5, column=3, value=" Forecast Type ")
    assert find_header(ws, {"forecast type"}) == (5, 3)
    ws.cell(row=500, column=1, value="Forecast Type")
    ws["C5"] = None
    assert find_header(ws, {"forecast type"}) == (None, None)
