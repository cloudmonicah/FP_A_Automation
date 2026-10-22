"""Step 3: validate one category at a time; the month's report keeps the others."""

import xml.etree.ElementTree as ET
from datetime import date, datetime

from openpyxl import Workbook, load_workbook

from automations.forecast import service
from automations.forecast import tableau_service as ts
from automations.forecast import validation_service as vs

FORECAST_TYPE_TAB = """<?xml version='1.0' encoding='utf-8' ?>
<workbook>
  <datasources><datasource name='Parameters'>
    <column caption='Forecast Start' name='[Parameter 1]' param-domain-type='range' value='#2026-08-01#' />
  </datasource></datasources>
  <worksheets><worksheet name='Monthly Forecast'><table><view>
    <datasource-dependencies datasource='Parameters'>
      <column caption='Forecast Start' name='[Parameter 1]' />
    </datasource-dependencies>
    <filter class='categorical' column='[ds].[none:Company:nk]'>
      <groupfilter function='member' level='[none:Company:nk]' member='"ALIV"' />
    </filter>
    <filter class='categorical' column='[ds].[none:Forecast Type:nk]'>
      <groupfilter function='union'>
        <groupfilter function='member' level='[none:Forecast Type:nk]' member='"FY27 Forecast_092026"' />
      </groupfilter>
    </filter>
  </view></table></worksheet></worksheets>
</workbook>
"""


def test_forecast_type_tab_is_prepared_for_the_period(tmp_path):
    """This used to fail on every Forecast Type tab ("measure_name" not defined)."""
    source = tmp_path / "Legal.twb"
    source.write_text(FORECAST_TYPE_TAB, encoding="utf-8")
    out = tmp_path / "copy.twb"
    prepared = vs._prepare_individual_forecast_copy(
        source=source, destination=out, tab="Monthly Forecast", company="CBL",
        forecast_start=date(2026, 9, 1), forecast_type="FY27 Forecast_102026",
    )
    root = ET.parse(out).getroot()
    members = [g.get("member") for g in root.iter("groupfilter") if g.get("function") == "member"]
    assert prepared["period_mode"] == "Forecast Type"
    assert '"FY27 Forecast_102026"' in members
    # Company is left as saved: each individual tab is already built for its company.
    assert '"ALIV"' in members and '"CBL"' not in members
    assert next(root.iter("column")).get("value") == "#2026-09-01#"
    assert source.read_text(encoding="utf-8") == FORECAST_TYPE_TAB  # source untouched


def _forecast_sheet(path, sheet, total):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    months = [datetime(2026 if m >= 7 else 2027, m, 1) for m in (7, 8, 9, 10, 11, 12, 1, 2, 3, 4, 5, 6)]
    ws.append(["Forecast Type", "GL Account", *months, "Grand Total"])
    ws.append(["FY27 Forecast_102026", "6000", total, *[0] * 11, "=SUM(C2:N2)"])
    ws.append([None, None, *[f"=SUM({c}2:{c}2)" for c in "CDEFGHIJKLMNO"]])   # total row
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def test_one_category_is_validated_and_the_report_keeps_the_others(tmp_path, monkeypatch):
    twb_folder = tmp_path / "Workbooks"
    twb_folder.mkdir()
    (twb_folder / "Legal.twb").write_text(FORECAST_TYPE_TAB, encoding="utf-8")
    routes = [
        {"category": "OPEX", "workbook": "FY27 OPEX Legal_Aug Forecast.xlsx", "excel_sheet": "ALIV",
         "individual_tableau_folder": str(twb_folder), "individual_tableau_workbook": "Legal.twb",
         "individual_tableau_worksheet": "Monthly Forecast", "company_filter": "ALIV", "validation_status": "confirmed"},
        {"category": "Revenue", "workbook": "FY27 OPEX Legal_Aug Forecast.xlsx", "excel_sheet": "CBL",
         "individual_tableau_folder": str(twb_folder), "individual_tableau_workbook": "Legal.twb",
         "individual_tableau_worksheet": "Monthly Forecast", "company_filter": "CBL", "validation_status": "confirmed"},
    ]
    out = tmp_path / "out"
    wb_path = out / "FY27 OPEX Legal_Oct Forecast.xlsx"
    _forecast_sheet(wb_path, "ALIV", 100)
    wb = load_workbook(wb_path)
    cbl = wb.create_sheet("CBL")
    for row in wb["ALIV"].iter_rows(values_only=True):
        cbl.append(list(row))
    wb.save(wb_path)

    report_folder = tmp_path / "report"
    temp = tmp_path / "temp"
    temp.mkdir()
    monkeypatch.setattr(service, "routes_for", lambda cats: [r for r in routes if r["category"] in cats])
    monkeypatch.setattr(service, "get_forecast_validation_output_folder", lambda **k: report_folder)
    monkeypatch.setattr("shared.tableau_temp_copies.temp_twb_path", lambda src, tag: temp / f"{tag}.twb")
    opened = []
    monkeypatch.setattr(ts, "open_tableau_workbook", lambda tableau_workbook_path: opened.append(tableau_workbook_path))
    monkeypatch.setattr(ts, "get_tableau_crosstab_rows",
                        lambda worksheet_name, tableau_workbook_path: [{"Company": "Grand Total", "Grand Total": "100"}])
    monkeypatch.setattr(vs, "close_individual_tableau_workbook", lambda path: {"closed": True, "warning": None})
    monkeypatch.setattr("time.sleep", lambda s: None)

    first = service.validate_monthly_forecast(out, "FY27", 10, 2026, categories=["OPEX"])
    assert [(r["category"], r["status"]) for r in first["results"]] == [("OPEX", "PASS")]
    assert len(opened) == 1  # only OPEX's workbook was opened

    second = service.validate_monthly_forecast(out, "FY27", 10, 2026, categories=["Revenue"])
    report = load_workbook(second["validation_report"])
    detail = report["Validation Detail"]
    header = [c.value for c in detail[1]]
    categories = [row[header.index("Category")] for row in detail.iter_rows(min_row=2, values_only=True)]
    assert categories == ["OPEX", "Revenue"]  # OPEX kept after running Revenue
    summary = {row[0]: row[1] for row in report["Validation Summary"].iter_rows(values_only=True)}
    assert summary["Forecast Start"] == datetime(2026, 10, 1)
    assert summary["Passed"] == 2


def test_results_saved_by_an_earlier_version_are_read(tmp_path):
    """Live run: a .results.json from the earlier category branch crashed the report."""
    import json

    from shared.forecast_validation_report import merge_saved_results

    old = {"forecast_type": "FY27 Forecast_102026", "categories": {
        "OPEX": {"validated_at": "2026-10-15 18:00", "results": [{"status": "PASS", "category": "OPEX"}]},
    }}
    (tmp_path / "Forecast_Validation_Oct_2026.results.json").write_text(json.dumps(old), encoding="utf-8")
    new = [{"status": "ACTION REQUIRED", "category": "Direct Cost"}]
    merged = merge_saved_results(tmp_path, 10, 2026, "FY27 Forecast_102026", ["Direct Cost"], new)
    assert [(r["category"], r["status"]) for r in merged] == [("OPEX", "PASS"), ("Direct Cost", "ACTION REQUIRED")]


def test_company_filter_message_is_plain():
    from shared.friendly_errors import friendly

    text = friendly(ValueError('Expected one Company filter on individual forecast tab "Monthly Forecast".'))
    assert '"Monthly Forecast"' in text and "not validated" in text


def test_grand_total_found_on_a_tab_without_a_company_column():
    # Live run: ALIV Head Office's tab starts with "GL Account Mapping".
    rows = [{"GL Account Mapping": "Salaries", "Grand Total": "10"},
            {"GL Account Mapping": "Grand Total", "Grand Total": "1,234"}]
    assert vs.get_individual_tableau_grand_total(rows) == 1234


def _sheet_with_totals(path, values_by_row, extra_total=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "ALIV"
    months = [datetime(2026 if m >= 7 else 2027, m, 1) for m in (7, 8, 9, 10, 11, 12, 1, 2, 3, 4, 5, 6)]
    ws.append(["Forecast Type", "GL Account", *months, "Grand Total"])
    for i, values in enumerate(values_by_row, start=2):
        ws.append(["FY27 Forecast_102026", f"60{i}", *values, f"=SUM(C{i}:N{i})"])
    end = 1 + len(values_by_row)
    total = [None, None, *[f"=SUM({c}2:{c}{end})" for c in "CDEFGHIJKLMN"],
             extra_total or f"=SUM(O2:O{end})"]
    ws.append(total)
    wb.save(path)
    return path


def test_excel_grand_total_is_the_sheets_own_total_cell(tmp_path):
    path = _sheet_with_totals(tmp_path / "w.xlsx", [[100] * 12, [-50, *[None] * 11]])
    assert vs.get_excel_forecast_grand_total(path, "ALIV") == 1150
    # A total cell with its own formula is followed exactly (here: minus a fixed 150).
    path = _sheet_with_totals(tmp_path / "x.xlsx", [[100] * 12], extra_total="=SUM(O2:O2)-150")
    assert vs.get_excel_forecast_grand_total(path, "ALIV") == 1050


def test_unsupported_total_formula_is_explained(tmp_path):
    import pytest

    from shared.friendly_errors import PlainLanguageError

    path = _sheet_with_totals(tmp_path / "w.xlsx", [[1] * 12], extra_total="=SUMIF(A2:A2,\"x\",O2:O2)")
    with pytest.raises(PlainLanguageError, match="Open the workbook in Excel, save it"):
        vs.get_excel_forecast_grand_total(path, "ALIV")


def test_tolerance_is_1000_dollars():
    assert vs.compare_forecast_totals(10_000, 9_000.5)["status"] == "PASS"
    assert vs.compare_forecast_totals(10_000, 8_999)["status"] == "FAIL"


def test_needs_filter_lines_are_reported_not_skipped(tmp_path, monkeypatch):
    route = {"category": "OPEX", "workbook": "FY27 OPEX Engineering & Technology_Aug Forecast.xlsx",
             "excel_sheet": "Tech Ops - ALIV", "individual_tableau_folder": str(tmp_path),
             "individual_tableau_workbook": "Engineering.twb",
             "individual_tableau_worksheet": "GROUP Monthly Forecast by Sup Department",
             "validation_status": "needs_filter", "validation_filter_note": "Technical Operations subdepartment"}
    opened = []
    monkeypatch.setattr(ts, "open_tableau_workbook", lambda tableau_workbook_path: opened.append(1))
    result = vs.run_grouped_individual_validation([route], tmp_path, "Oct", 10, 2026)
    (item,) = result["results"]
    assert item["status"] == "NOT VALIDATED" and item["category"] == "OPEX"
    assert "Technical Operations subdepartment filter" in item["note"]
    assert opened == []


def test_tableau_forecast_type_is_checked_and_tidied():
    import pytest

    from shared.friendly_errors import PlainLanguageError

    assert service.check_tableau_forecast_type(" fy27  forecast_092026 ", "FY27") == "FY27 Forecast_092026"
    with pytest.raises(PlainLanguageError, match="for example FY27 Forecast_102026"):
        service.check_tableau_forecast_type("Forecast Oct", "FY27")


def test_step_3_uses_the_chosen_tableau_forecast_type(tmp_path, monkeypatch):
    seen = {}

    def fake_run(**kwargs):
        seen["forecast_type"] = kwargs["forecast_type"]
        return {"results": [], "workbooks": [], "passed": 0, "failed": 0, "action_required": 0,
                "status": "PASS"}

    monkeypatch.setattr(service, "run_grouped_individual_validation", fake_run)
    monkeypatch.setattr(service, "get_forecast_validation_output_folder", lambda **k: tmp_path)
    result = service.validate_monthly_forecast(tmp_path, "FY27", 10, 2026, categories=["OPEX"],
                                               tableau_forecast_type="FY27 Forecast_092026")
    assert seen["forecast_type"] == "FY27 Forecast_092026"
    summary = {row[0]: row[1] for row in load_workbook(result["validation_report"])["Validation Summary"]
               .iter_rows(values_only=True)}
    assert summary["Forecast Type"] == "FY27 Forecast_092026" and summary["Tolerance"] == 1000
