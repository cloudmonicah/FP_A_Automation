import xml.etree.ElementTree as ET

import pytest
from openpyxl import load_workbook

from automations.forecast import service
from tests.forecast.conftest import write_forecast_workbook

CONSOLIDATED = """<?xml version='1.0' encoding='utf-8' ?>
<workbook>
  <datasources><datasource name='Parameters'>
    <column caption='Forecast Start' name='[Parameter 1]' param-domain-type='range' value='#2026-09-01#' />
  </datasource></datasources>
  <worksheets><worksheet name='ALIV Legal'><table><view>
    <filter class='categorical' column='[ds].[none:Forecast Type:nk]'>
      <groupfilter function='member' level='[none:Forecast Type:nk]' member='"FY27 Forecast_092026"' />
    </filter>
  </view></table></worksheet></worksheets>
</workbook>
"""

MONTHS = ["Jul", "Aug", "Sep", "Oct", "Nov", "Dec", "Jan", "Feb", "Mar", "Apr", "May", "Jun"]
ROUTES = [
    {"tableau_worksheet": "ALIV Legal", "workbook": "FY27 OPEX Legal_Aug Forecast.xlsx", "excel_sheet": "ALIV"},
    {"tableau_worksheet": "CBL Legal", "workbook": "FY27 OPEX Legal_Aug Forecast.xlsx", "excel_sheet": "CBL"},
    {"tableau_worksheet": "ALIV HR", "workbook": "FY27 OPEX HR_Aug Forecast.xlsx", "excel_sheet": "ALIV"},
]


@pytest.fixture
def setup(tmp_path, monkeypatch):
    source = tmp_path / "source" / "FY27 FORECAST WORKBOOKS.twb"
    source.parent.mkdir()
    source.write_text(CONSOLIDATED, encoding="utf-8")
    temp_folder = tmp_path / "temp"
    temp_folder.mkdir()
    monkeypatch.setattr(
        service, "_temp_copy", lambda src, tag: temp_folder / f"{tag}.twb"
    )
    out = tmp_path / "out"
    # Legal workbook has ALIV (October) and CBL is missing from it on purpose
    write_forecast_workbook(out / "FY27 OPEX Legal_Oct Forecast.xlsx", sheet="ALIV")
    # HR workbook was left over from September (skipped in step 1)
    write_forecast_workbook(
        out / "FY27 OPEX HR_Oct Forecast.xlsx", sheet="ALIV", forecast_type="FY27 Forecast_092026"
    )
    return source, temp_folder, out


def test_population_uses_a_temp_copy_and_continues_past_bad_routes(setup):
    source, temp_folder, out = setup
    before = source.read_bytes()
    seen = {}

    def extractor(tableau_workbook_path, worksheet_names, progress=None):
        root = ET.parse(tableau_workbook_path).getroot()
        seen["start"] = [c.get("value") for c in root.iter("column")][0]
        seen["types"] = [g.get("member") for g in root.iter("groupfilter")]
        rows = [{"GL Account": "6000", **{m: "10" for m in MONTHS}, "Grand Total": "120"}]
        return {name: rows for name in worksheet_names if name != "ALIV HR"} | {"ALIV HR": rows}, {}

    events = []
    result = service.populate_monthly_forecast_from_tableau(
        forecast_folder=out,
        tableau_workbook_path=source,
        fiscal_year="FY27",
        forecast_month=10,
        forecast_year=2026,
        routes=ROUTES,
        opener=lambda tableau_workbook_path: events.append(("open", tableau_workbook_path)),
        waiter=lambda tableau_workbook_path: events.append(("wait", tableau_workbook_path)),
        extractor=extractor,
        closer=lambda path: events.append(("close", path)) or {"closed": True},
    )

    assert source.read_bytes() == before  # source never modified
    assert not list(source.parent.glob("*_BEFORE_*"))  # and no backup beside it
    assert not list(temp_folder.iterdir())  # temp copy removed
    assert [e[0] for e in events] == ["open", "wait", "close"]
    assert seen["start"] == "#2026-10-01#"
    assert seen["types"] == ['"FY27 Forecast_102026"']

    assert [p["worksheet"] for p in result["passed"]] == ["ALIV"]
    errors = {f["workbook"] + "/" + f["worksheet"]: f["error"] for f in result["failed"]}
    assert errors["FY27 OPEX Legal_Oct Forecast.xlsx/CBL"]
    hr = errors["FY27 OPEX HR_Oct Forecast.xlsx/ALIV"]
    assert "September 2026" in hr and "October 2026" in hr

    wb = load_workbook(out / "FY27 OPEX Legal_Oct Forecast.xlsx")
    assert wb["ALIV"].cell(4, 3).value == 10  # first month of the populated row
    wb.close()


def test_tableau_is_closed_and_copy_removed_when_extraction_fails(setup):
    source, temp_folder, out = setup
    closed = []

    def extractor(**kwargs):
        raise TimeoutError("Tableau did not finish opening")

    with pytest.raises(TimeoutError):
        service.populate_monthly_forecast_from_tableau(
            forecast_folder=out, tableau_workbook_path=source, fiscal_year="FY27",
            forecast_month=10, forecast_year=2026, routes=ROUTES,
            opener=lambda **k: None, waiter=lambda **k: None, extractor=extractor,
            closer=lambda path: closed.append(path) or {"closed": True},
        )
    assert closed and not list(temp_folder.iterdir())


def test_fiscal_year_must_match_the_month():
    assert service.fiscal_year_for(10, 2026) == "FY27"
    assert service.fiscal_year_for(6, 2027) == "FY27"
    assert service.fiscal_year_for(7, 2027) == "FY28"
    with pytest.raises(ValueError, match="October 2026 is in FY27"):
        service.build_forecast_period("FY28", 10, 2026)


def test_check_consolidated_workbook_leaves_no_files(setup):
    source, temp_folder, out = setup
    period = service.check_consolidated_workbook(source, "FY27", 10, 2026)
    assert period["forecast_type"] == "FY27 Forecast_102026"
    assert not list(temp_folder.iterdir())
    assert sorted(p.name for p in source.parent.iterdir()) == [source.name]
