import pytest

from automations.forecast.excel_service import (
    MissingMonthsError,
    _s4_tableau_months,
    populate_forecast_sheet_from_tableau_rows,
)
from automations.forecast.tableau_service import _parse_crosstab_text
from tests.forecast.conftest import write_forecast_workbook

FY27 = ["July 2026", "August 2026", "September 2026", "October 2026", "November 2026",
        "December 2026", "January 2027", "February 2027", "March 2027", "April 2027",
        "May 2027", "June 2027"]


def _row(headers, value="1"):
    return {"GL Account": "6000", **{h: value for h in headers}}


def test_all_twelve_months_are_mapped_by_month_and_year():
    lookup = _s4_tableau_months([_row(FY27)], "FY27")
    assert len(lookup) == 12
    assert lookup[9] == "September 2026" and lookup[6] == "June 2027"


def test_months_from_another_fiscal_year_are_never_used():
    # A rolling view from Oct 2026: Jul-Sep 2027 must not fill Jul-Sep 2026.
    rolling = FY27[3:] + ["July 2027", "August 2027", "September 2027"]
    with pytest.raises(MissingMonthsError) as error:
        _s4_tableau_months([_row(rolling)], "FY27", "Legal / ALIV")
    message = str(error.value)
    assert "9 of the 12 months" in message and "Missing: Jul, Aug, Sep" in message
    assert '"Legal / ALIV"' in message and "Jul 2026 to Jun 2027" in message


def test_partial_months_leave_the_sheet_unchanged(tmp_path):
    path = write_forecast_workbook(tmp_path / "x.xlsx", rows=((100, 200, 300),))
    before = path.read_bytes()
    rows = [_row(["September 2026", "October 2026"], "5")]
    with pytest.raises(MissingMonthsError):
        populate_forecast_sheet_from_tableau_rows(path, "ALIV", rows, "FY27")
    assert path.read_bytes() == before


def test_all_twelve_months_are_written(tmp_path):
    from openpyxl import load_workbook

    path = write_forecast_workbook(tmp_path / "x.xlsx", rows=((1,),))
    rows = [{"GL Account": "6000", **{h: str(i + 1) for i, h in enumerate(FY27)}}]
    populate_forecast_sheet_from_tableau_rows(path, "ALIV", rows, "FY27")
    ws = load_workbook(path)["ALIV"]
    assert [ws.cell(4, c).value for c in range(3, 15)] == list(range(1, 13))


def test_two_header_rows_are_combined():
    text = "\t".join(["", "", "July 2026", "", "August 2026", ""]) + "\n"
    text += "\t".join(["GL Account", "Company", "Actual", "Forecast", "Actual", "Forecast"]) + "\n"
    text += "\t".join(["6000", "ALIV", "1", "2", "3", "4"]) + "\n"
    rows = _parse_crosstab_text(text)
    assert list(rows[0]) == ["GL Account", "Company", "July 2026 Actual", "July 2026 Forecast",
                             "August 2026 Actual", "August 2026 Forecast"]
    assert rows[0]["August 2026 Forecast"] == "4"


def test_forecast_column_wins_when_a_month_has_two_measures():
    headers = [f"{m} {kind}" for m in FY27 for kind in ("Actual", "Forecast")]
    lookup = _s4_tableau_months([_row(headers)], "FY27")
    assert lookup[10] == "October 2026 Forecast"


def test_single_header_crosstab_is_unchanged():
    text = "GL Account\tOctober 2026\tGrand Total\n6000\t(1,250)\t5\n"
    assert _parse_crosstab_text(text) == [{"GL Account": "6000", "October 2026": "(1,250)", "Grand Total": "5"}]
