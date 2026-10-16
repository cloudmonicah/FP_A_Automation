"""Step 2: closed months come from Tableau, forecast months are kept."""

import xml.etree.ElementTree as ET
from datetime import datetime

import pytest
from openpyxl import Workbook, load_workbook

from automations.forecast import service
from automations.forecast.excel_service import ActualsUpdateError, closed_months, update_actual_months
from shared.friendly_errors import friendly

OCT = "FY27 Forecast_102026"
MONTHS = [datetime(2026 if m >= 7 else 2027, m, 1) for m in (7, 8, 9, 10, 11, 12, 1, 2, 3, 4, 5, 6)]


def _sheet(path, rows, forecast_type=OCT, sheet="ALIV"):
    """Forecast sheet as step 1 leaves it: last month's numbers for every month."""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(["Legal forecast"])
    ws.append(["Forecast Type", "GL Sub Segment code", "GL Sub Account Description", *MONTHS, "Grand Total"])
    for code, desc, values in rows:
        ws.append([forecast_type, code, desc, *values, "=SUM()"])
    ws.append([None, None, None, *["=SUM()"] * 12, "=SUM()"])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def _tableau(code, desc, jul, aug, sep, **extra):
    """A Tableau line as in the live run: only Jul-26, Aug-26, Sep-26."""
    return {"GL Sub Segment code": code, "GL Sub Account Description": desc, **extra,
            "Jul-26": jul, "Aug-26": aug, "Sep-26": sep, "Grand Total": "x"}


def _values(path, sheet="ALIV"):
    ws = load_workbook(path)[sheet]
    out = {}
    for r in range(3, ws.max_row + 1):
        if ws.cell(r, 1).value:
            out[str(ws.cell(r, 2).value)] = [ws.cell(r, c).value for c in range(4, 16)]
    return out


def test_closed_months_for_october():
    assert closed_months("FY27", 10) == [(2026, 7), (2026, 8), (2026, 9)]
    assert closed_months("FY27", 2) == [(2026, m) for m in range(7, 13)] + [(2027, 1)]


def test_actuals_replace_closed_months_and_forecast_months_are_kept(tmp_path):
    forecast = [10] * 12  # last month's workbook: 10 in every month
    path = _sheet(tmp_path / "FY27 OPEX Legal_Oct Forecast.xlsx", [
        (6000, "Salaries", forecast),
        (6100, "Travel", forecast),     # not in Tableau any more
    ])
    rows = [
        _tableau("6000", "Salaries", "1,000", "2,000", "(300)"),
        _tableau("6200", "Training", "50", "", "70"),  # new line
    ]
    result = update_actual_months(path, "ALIV", rows, "FY27", 10, OCT)
    values = _values(path)
    assert values["6000"] == [1000, 2000, -300] + [10] * 9   # Sep actual replaces Sep forecast
    assert values["6100"] == [None, None, None] + [10] * 9   # no actuals; forecast kept
    assert values["6200"] == [50, 0, 70] + [None] * 9        # new line, forecast left for the department
    assert result["lines_added"] == 1 and result["lines_without_actuals"] == 1
    ws = load_workbook(path)["ALIV"]
    assert ws.cell(5, 16).value == "=SUM(D5,E5,F5,G5,H5,I5,J5,K5,L5,M5,N5,O5)"  # row total for the new line
    assert ws.cell(6, 4).value == "=SUM(D3:D5)"  # bottom total covers the new line


def test_lines_split_in_tableau_are_added_together(tmp_path):
    path = _sheet(tmp_path / "FY27 CBL BANDWIDTH_Oct Forecast.xlsx", [(6000, "Bandwidth", [5] * 12)])
    rows = [
        _tableau("6000", "Bandwidth", "100", "100", "100", **{"Vendor Name": "A"}),
        _tableau("6000", "Bandwidth", "20", "30", "40", **{"Vendor Name": "B"}),
    ]
    update_actual_months(path, "ALIV", rows, "FY27", 10, OCT)
    assert _values(path)["6000"] == [120, 130, 140] + [5] * 9


def test_revenue_actuals_keep_the_revenue_sign(tmp_path):
    path = _sheet(tmp_path / "FY27 CONSUMER REVENUE_Oct Forecast.xlsx", [(4000, "Sales", [1] * 12)])
    update_actual_months(path, "ALIV", [_tableau("4000", "Sales", "-500", "-600", "-700")], "FY27", 10, OCT)
    assert _values(path)["4000"][:3] == [500, 600, 700]


def test_sheet_for_another_month_is_left_unchanged(tmp_path):
    path = _sheet(tmp_path / "FY27 OPEX HR_Oct Forecast.xlsx", [(6000, "Salaries", [10] * 12)],
                  forecast_type="FY27 Forecast_092026")
    before = path.read_bytes()
    with pytest.raises(ActualsUpdateError) as error:
        update_actual_months(path, "ALIV", [_tableau("6000", "Salaries", "1", "2", "3")], "FY27", 10, OCT)
    assert path.read_bytes() == before
    assert "replace" in friendly(error.value)


def test_lines_that_cannot_be_matched_leave_the_sheet_unchanged(tmp_path):
    path = _sheet(tmp_path / "x_Oct Forecast.xlsx", [(6000, "Salaries", [1] * 12)] * 2)
    before = path.read_bytes()
    rows = [_tableau("6000", "Salaries", "1", "1", "1")] * 3
    with pytest.raises(ActualsUpdateError, match="2 times in Excel and 3 times in Tableau"):
        update_actual_months(path, "ALIV", rows, "FY27", 10, OCT)
    assert path.read_bytes() == before


CONSOLIDATED = """<?xml version='1.0' encoding='utf-8' ?>
<workbook>
  <datasources><datasource name='Parameters'>
    <column caption='Forecast Start' name='[Parameter 1]' param-domain-type='range' value='#2026-09-01#' />
  </datasource></datasources>
  <worksheets><worksheet name='ALIV Legal'><table><view>
    <filter class='categorical' column='[ds].[none:Forecast Type:nk]'>
      <groupfilter function='member' level='[none:Forecast Type:nk]' member='"FY27 Forecast_092026"' />
    </filter>
  </view></table></worksheet></worksheets>
</workbook>
"""

ROUTES = [
    {"category": "OPEX", "tableau_worksheet": "ALIV Legal", "workbook": "FY27 OPEX Legal_Aug Forecast.xlsx", "excel_sheet": "ALIV"},
    {"category": "Revenue", "tableau_worksheet": "ALIV Revenue", "workbook": "FY27 REVENUE_Aug Forecast.xlsx", "excel_sheet": "ALIV"},
]


def test_step2_runs_one_category_on_a_temp_copy(tmp_path, monkeypatch):
    source = tmp_path / "source" / "FY27 FORECAST WORKBOOKS.twb"
    source.parent.mkdir()
    source.write_text(CONSOLIDATED, encoding="utf-8")
    before = source.read_bytes()
    temp = tmp_path / "temp"
    temp.mkdir()
    monkeypatch.setattr(service, "_temp_copy", lambda src, tag: temp / f"{tag}.twb")
    monkeypatch.setattr(service, "routes_for", lambda cats: [r for r in ROUTES if not cats or r["category"] in cats])
    out = tmp_path / "out"
    _sheet(out / "FY27 OPEX Legal_Oct Forecast.xlsx", [(6000, "Salaries", [10] * 12)])

    seen, events = {}, []

    def extractor(tableau_workbook_path, worksheet_names):
        root = ET.parse(tableau_workbook_path).getroot()
        seen["start"] = next(root.iter("column")).get("value")
        seen["names"] = worksheet_names
        return {"ALIV Legal": [_tableau("6000", "Salaries", "1", "2", "3")]}, {}

    result = service.populate_monthly_forecast_from_tableau(
        forecast_folder=out, tableau_workbook_path=source, fiscal_year="FY27",
        forecast_month=10, forecast_year=2026, categories=["OPEX"],
        opener=lambda tableau_workbook_path: events.append("open"),
        waiter=lambda tableau_workbook_path: events.append("wait"),
        extractor=extractor,
        closer=lambda path: events.append("close") or {"closed": True},
    )
    assert source.read_bytes() == before                # source never edited
    assert not list(temp.iterdir())                     # temp copy removed
    assert events == ["open", "wait", "close"]
    assert seen == {"start": "#2026-10-01#", "names": ["ALIV Legal"]}  # only OPEX tabs
    assert result["routes_passed"] == 1 and result["months_updated"] == ["Jul 2026", "Aug 2026", "Sep 2026"]
    assert result["audit_workbook"].endswith("Tableau_Forecast_Extract_2026-10_OPEX.xlsx")
    assert _values(out / "FY27 OPEX Legal_Oct Forecast.xlsx")["6000"] == [1, 2, 3] + [10] * 9


@pytest.mark.parametrize(
    "error, expected",
    [
        (RuntimeError('Excel sheet "CBL" not found. Available: [\'ALIV\']'), 'no sheet named "CBL"'),
        (RuntimeError('Tableau parameter "Forecast Start" could not be located in FY27.twb.'), '"Forecast Start" parameter'),
        (RuntimeError("Tableau extraction failed: RuntimeError: Tableau executed Copy Crosstab, but no Unicode text was available on the Windows clipboard."),
         "pop-up message"),
        (service.ForecastStepError("FY27 OPEX IT_Oct Forecast.xlsx is not in the Forecast Output Folder. Run step 1 for this month first."),
         "FY27 OPEX IT_Oct Forecast.xlsx is not in the Forecast Output Folder"),
    ],
)
def test_step2_errors_are_explained(error, expected):
    from shared.friendly_errors import GENERIC

    message = friendly(error)
    assert message != GENERIC and expected in message
