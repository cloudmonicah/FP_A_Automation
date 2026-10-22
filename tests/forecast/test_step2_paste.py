"""Step 2 pastes each consolidated Tableau tab into its sheet and keeps every formula."""

from datetime import datetime

from openpyxl import Workbook, load_workbook

from automations.forecast.excel_service import populate_forecast_sheet_from_tableau_rows

OCT = "FY27 Forecast_102026"


def _sheet(path, lines):
    """A forecast sheet like the real ones: =EDATE month headers, formula forecast months."""
    wb = Workbook()
    ws = wb.active
    ws.title = "ALIV"
    ws.append(["Legal forecast"])
    ws.append([
        "Forecast Type", "GL Sub Segment code", "GL Sub Account Description",
        datetime(2026, 7, 1), *[f"=EDATE(D2,{i})" for i in range(1, 12)], "Grand Total",
    ])
    for i, (code, desc) in enumerate(lines, start=3):
        # Jul-Sep values; Oct-Jun formulas (the department's forecast); row total.
        ws.append([OCT, code, desc, 1, 1, 1, *[f"=AVERAGE(D{i}:F{i})"] * 9, f"=SUM(D{i}:O{i})"])
    end = 2 + len(lines)
    ws.append([None, None, None, *[f"=SUM({c}3:{c}{end})" for c in "DEFGHIJKLMNO"], f"=SUM(P3:P{end})"])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


LATER = ["Oct-26", "Nov-26", "Dec-26", "Jan-27", "Feb-27", "Mar-27", "Apr-27", "May-27", "Jun-27"]


def _row(code, desc, jul, aug, sep, later="9"):
    """A line as Tableau copies it, with all 12 fiscal months."""
    return {"GL Business Code": "LEGAL", "GL Sub Segment code": code, "GL Sub Account Description": desc,
            "Jul-26": jul, "Aug-26": aug, "Sep-26": sep, **{m: later for m in LATER}, "Grand Total": "999"}


def test_tableau_values_are_pasted_and_formulas_stay(tmp_path):
    path = _sheet(tmp_path / "FY27 OPEX Legal_Oct Forecast.xlsx", [(6000, "Salaries"), (6100, "Travel")])
    rows = [_row("6000", "Salaries", "1,000", "2,000", "(300)"), _row("6100", "Travel", "5", "6", "7")]
    populate_forecast_sheet_from_tableau_rows(path, "ALIV", rows, "FY27")

    ws = load_workbook(path)["ALIV"]
    assert [ws.cell(3, c).value for c in (4, 5, 6)] == [1000, 2000, -300]   # pasted as Tableau shows
    assert [ws.cell(3, c).value for c in range(7, 16)] == [9] * 9            # all 12 periods present
    assert ws.cell(3, 16).value == "=SUM(D3:O3)"                           # row total formula kept
    assert ws.cell(2, 5).value == "=EDATE(D2,1)"                           # header formulas kept
    assert ws.cell(5, 4).value == "=SUM(D3:D4)"                            # bottom total


def test_extra_tableau_lines_get_the_formulas_filled_down(tmp_path):
    path = _sheet(tmp_path / "FY27 OPEX Legal_Oct Forecast.xlsx", [(6000, "Salaries")])
    rows = [_row("6000", "Salaries", "1", "2", "3"), _row("6200", "Training", "4", "5", "6")]
    populate_forecast_sheet_from_tableau_rows(path, "ALIV", rows, "FY27")

    ws = load_workbook(path)["ALIV"]
    assert [ws.cell(4, c).value for c in (1, 2, 3, 4, 5, 6)] == [OCT, 6200, "Training", 4, 5, 6]
    assert [ws.cell(4, c).value for c in range(7, 16)] == [9] * 9
    assert ws.cell(4, 16).value == "=SUM(D4:O4)"
    assert ws.cell(5, 4).value == "=SUM(D3:D4)"       # total now covers the added line
    assert ws.cell(5, 16).value == "=SUM(P3:P4)"


def test_fewer_tableau_lines_remove_rows(tmp_path):
    path = _sheet(tmp_path / "FY27 OPEX Legal_Oct Forecast.xlsx", [(6000, "Salaries"), (6100, "Travel"), (6200, "x")])
    populate_forecast_sheet_from_tableau_rows(path, "ALIV", [_row("6000", "Salaries", "1", "2", "3")], "FY27")

    ws = load_workbook(path)["ALIV"]
    assert ws.cell(3, 2).value == 6000 and ws.cell(4, 1).value is None
    assert ws.cell(4, 4).value == "=SUM(D3:D3)"


def test_revenue_keeps_the_existing_sign_rule(tmp_path):
    path = _sheet(tmp_path / "FY27 CONSUMER REVENUE_Oct Forecast.xlsx", [(4000, "Sales")])
    populate_forecast_sheet_from_tableau_rows(path, "ALIV", [_row("4000", "Sales", "-10", "-20", "-30")], "FY27")
    ws = load_workbook(path)["ALIV"]
    assert [ws.cell(3, c).value for c in (4, 5, 6)] == [10, 20, 30]


def test_month_from_another_fiscal_year_is_not_pasted(tmp_path):
    path = _sheet(tmp_path / "FY27 OPEX Legal_Oct Forecast.xlsx", [(6000, "Salaries")])
    row = _row("6000", "Salaries", "1", "2", "3")
    row["Jul-27"] = "777"  # FY28 - must not land in Jul 2026
    populate_forecast_sheet_from_tableau_rows(path, "ALIV", [row], "FY27")
    assert load_workbook(path)["ALIV"].cell(3, 4).value == 1


def test_months_the_tableau_tab_leaves_out_are_pasted_blank(tmp_path):
    # Live run: "INTERNET BANDWIDTH (ALIV BUSINESS)" shows only Oct-26..Jun-27.
    path = _sheet(tmp_path / "FY27 CBL BANDWIDTH_Oct Forecast.xlsx", [(6000, "Bandwidth")])
    row = {k: v for k, v in _row("6000", "Bandwidth", "1", "2", "3").items() if k not in ("Jul-26", "Aug-26", "Sep-26")}
    result = populate_forecast_sheet_from_tableau_rows(path, "ALIV", [row], "FY27")
    ws = load_workbook(path)["ALIV"]
    assert [ws.cell(3, c).value for c in range(4, 16)] == [None, None, None] + [9] * 9
    assert result["months_blank_in_tableau"] == ["Jul", "Aug", "Sep"]


def test_missing_sheet_message_is_plain():
    from shared.friendly_errors import friendly

    text = friendly(RuntimeError('Excel sheet "Media" not found. Available: [\'Media (MARKETING)\']'))
    assert '"Media"' in text and "Media (MARKETING)" in text and "unexpected" not in text


def test_media_marketing_route_uses_the_real_tab_name():
    from automations.forecast.mapping import FORECAST_ROUTES

    route = next(r for r in FORECAST_ROUTES if r["tableau_worksheet"] == "MEDIA (MARKETING)")
    assert route["excel_sheet"] == "Media (MARKETING)"


# --- Real layout (Engineering & Technology): Comments, gaps, date formats ----

import pytest

from shared.friendly_errors import PlainLanguageError

MONTHS = ["Jul-26", "Aug-26", "Sep-26", "Oct-26", "Nov-26", "Dec-26",
          "Jan-27", "Feb-27", "Mar-27", "Apr-27", "May-27", "Jun-27"]


def _real_sheet(path, lines, comments=None, ft_blank=(), date_cells=(), merge=None):
    """Header in row 1, real month dates, Grand Total + Comments, SUM total row."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Tech Ops - CBL"
    ws.append(["Forecast Type", "Bus Code Name", "Company", "GL Sub Segment code",
               "GL Sub Account Description",
               *[datetime(2026 + (m < 7), m, 1) for m in (7, 8, 9, 10, 11, 12, 1, 2, 3, 4, 5, 6)],
               "Grand Total", "Comments"])
    for i, (bus, code, desc) in enumerate(lines, start=2):
        ws.append(["FY27 Forecast_102026", bus, "CBL", code, desc, *[1] * 12,
                   f"=SUM(F{i}:Q{i})", (comments or {}).get(code)])
        for col in range(6, 18):
            ws.cell(i, col).number_format = "#,##0"
    end = 1 + len(lines)
    ws.append([None] * 5 + [f"=SUM({c}2:{c}{end})" for c in "FGHIJKLMNOPQR"])
    for row in ft_blank:
        ws.cell(row, 1).value = None
    for ref in date_cells:
        ws[ref].number_format = "mmm-yy"
    if merge:
        ws.merge_cells(merge)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def _trow(bus, code, desc, values=None, months=MONTHS):
    values = values or {}
    return {"Bus Code Name": bus, "Company": "CBL", "GL Sub Segment code": code,
            "GL Sub Account Description": desc, **{m: values.get(m, "10") for m in months}}


def _lines(ws):
    return [(ws.cell(r, 2).value, ws.cell(r, 4).value, ws.cell(r, 19).value)
            for r in range(2, ws.max_row + 1) if ws.cell(r, 4).value is not None]


def test_comments_follow_their_account_when_tableau_reorders(tmp_path):
    path = _real_sheet(tmp_path / "FY27 OPEX Eng_Oct Forecast.xlsx",
                       [("AF", 600, "Pay"), ("AF", 610, "Travel"), ("FP", 620, "Fuel")],
                       comments={610: "Travel freeze", 620: "Moved from P&F"})
    rows = [_trow("FP", "620", "Fuel"), _trow("AF", "630", "New"), _trow("AF", "600", "Pay")]
    result = populate_forecast_sheet_from_tableau_rows(path, "Tech Ops - CBL", rows, "FY27")
    ws = load_workbook(path)["Tech Ops - CBL"]
    assert _lines(ws) == [("FP", 620, "Moved from P&F"), ("AF", 630, None), ("AF", 600, None)]
    # The account that left Tableau took its comment with it - reported, not lost silently.
    assert result["comments_dropped"] == [{"account": "AF / CBL / 610 / Travel", "text": "Travel freeze"}]


def test_missing_months_are_blank_and_present_months_land_in_their_column(tmp_path):
    path = _real_sheet(tmp_path / "FY27 OPEX Eng_Oct Forecast.xlsx", [("AF", 600, "Pay")])
    present = [m for m in MONTHS if m not in ("Jul-26", "Jun-27")]
    rows = [_trow("AF", "600", "Pay", {"Oct-26": "1,234", "May-27": "(50)"}, months=present)]
    result = populate_forecast_sheet_from_tableau_rows(path, "Tech Ops - CBL", rows, "FY27")
    ws = load_workbook(path)["Tech Ops - CBL"]
    assert ws["F2"].value is None and ws["Q2"].value is None      # Jul, Jun blank
    assert ws["I2"].value == 1234 and ws["P2"].value == -50       # Oct, May in place
    assert result["months_blank_in_tableau"] == ["Jul", "Jun"]


def test_revenue_flips_line_values_only_and_total_row_stays_a_sum(tmp_path):
    path = _real_sheet(tmp_path / "FY27 REVENUE Mobile_Oct Forecast.xlsx", [("AF", 400, "Sales")])
    rows = [_trow("AF", "400", "Sales", {"Oct-26": "500"}), _trow("AF", "401", "Fees", {"Oct-26": "(20)"})]
    populate_forecast_sheet_from_tableau_rows(path, "Tech Ops - CBL", rows, "FY27")
    ws = load_workbook(path)["Tech Ops - CBL"]
    assert ws["I2"].value == -500 and ws["I3"].value == 20
    assert ws["I4"].value == "=SUM(I2:I3)" and ws["R4"].value == "=SUM(R2:R3)"   # not negated


def test_amounts_formatted_as_dates_get_the_number_format(tmp_path):
    path = _real_sheet(tmp_path / "FY27 OPEX Eng_Oct Forecast.xlsx", [("AF", 600, "Pay")],
                       date_cells=("I2", "Q2"))
    result = populate_forecast_sheet_from_tableau_rows(
        path, "Tech Ops - CBL", [_trow("AF", "600", "Pay", {"Oct-26": "67,500"})], "FY27")
    ws = load_workbook(path)["Tech Ops - CBL"]
    assert ws["I2"].value == 67500 and ws["I2"].number_format == "#,##0"
    assert ws["Q2"].number_format == "#,##0" and result["date_formats_fixed"] == 2


def test_last_line_with_blank_forecast_type_is_still_a_line(tmp_path):
    path = _real_sheet(tmp_path / "FY27 OPEX Eng_Oct Forecast.xlsx",
                       [("AF", 600, "Pay"), ("IC", 610, "Recharge")], ft_blank=(3,))
    rows = [_trow("AF", "600", "Pay"), _trow("IC", "610", "Recharge"), _trow("IC", "620", "More")]
    populate_forecast_sheet_from_tableau_rows(path, "Tech Ops - CBL", rows, "FY27")
    ws = load_workbook(path)["Tech Ops - CBL"]
    assert [ws.cell(r, 4).value for r in (2, 3, 4)] == [600, 610, 620]
    assert ws["F5"].value == "=SUM(F2:F4)"                           # one total row, below
    assert all(ws.cell(r, c).value is None for r in range(6, ws.max_row + 1) for c in range(1, 20))
    assert ws["A3"].value == "FY27 Forecast_102026"                   # filled in


def test_total_row_a_row_below_the_lines_stays_a_row_below(tmp_path):
    path = _real_sheet(tmp_path / "FY27 OPEX Eng_Oct Forecast.xlsx", [("AF", 600, "Pay")])
    wb = load_workbook(path)
    wb.active.insert_rows(3)                                          # blank spacer row
    wb.save(path)
    rows = [_trow("AF", "600", "Pay"), _trow("AF", "610", "Travel")]
    result = populate_forecast_sheet_from_tableau_rows(path, "Tech Ops - CBL", rows, "FY27")
    ws = load_workbook(path)["Tech Ops - CBL"]
    assert result["total_row"] == 5 and ws["F5"].value == "=SUM(F2:F3)"
    assert ws["F4"].value is None


def test_merged_comment_cells_are_unmerged(tmp_path):
    path = _real_sheet(tmp_path / "FY27 OPEX Eng_Oct Forecast.xlsx",
                       [("AF", 600, "Pay"), ("AF", 610, "Travel")], comments={600: "Both lines"},
                       merge="S2:S3")
    result = populate_forecast_sheet_from_tableau_rows(
        path, "Tech Ops - CBL", [_trow("AF", "610", "Travel"), _trow("AF", "600", "Pay")], "FY27")
    ws = load_workbook(path)["Tech Ops - CBL"]
    assert result["unmerged"] == ["S2:S3"] and not ws.merged_cells.ranges
    assert _lines(ws) == [("AF", 610, None), ("AF", 600, "Both lines")]


def test_layout_problems_are_explained(tmp_path):
    path = _real_sheet(tmp_path / "FY27 OPEX Eng_Oct Forecast.xlsx", [("AF", 600, "Pay")])
    with pytest.raises(PlainLanguageError, match='no tab called "Tech Ops - ALIV".*Tech Ops - CBL'):
        populate_forecast_sheet_from_tableau_rows(path, "Tech Ops - ALIV", [_trow("AF", "600", "Pay")], "FY27")
    wb = load_workbook(path)
    wb.active["R1"] = "Total FY"
    wb.save(path)
    with pytest.raises(PlainLanguageError, match='no "Grand Total" column'):
        populate_forecast_sheet_from_tableau_rows(path, "Tech Ops - CBL", [_trow("AF", "600", "Pay")], "FY27")
