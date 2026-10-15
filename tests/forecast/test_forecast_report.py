import os

from openpyxl import load_workbook

from automations.forecast.validation_service import summarize_results
from shared import forecast_validation_report as report

RESULTS = [
    {"forecast_workbook": "A_Oct Forecast.xlsx", "worksheet": "ALIV", "status": "PASS",
     "excel_grand_total": 10.0, "tableau_grand_total": 10.0, "variance": 0.0},
    {"forecast_workbook": "B_Oct Forecast.xlsx", "worksheet": "CBL", "status": "NOT VALIDATED",
     "note": "B_Oct Forecast.xlsx / CBL is for September 2026"},
    {"forecast_workbook": "C_Oct Forecast.xlsx", "worksheet": "ALIV", "status": "ACTION REQUIRED",
     "note": "RuntimeError: Tableau workbook not ready yet"},
]


def _write(folder):
    return report.write_forecast_validation_report(
        validation_result=summarize_results(RESULTS),
        output_folder=folder,
        fiscal_year="FY27",
        forecast_month=10,
        forecast_year=2026,
        forecast_type="FY27 Forecast_102026",
    )


def _columns(ws):
    return {cell.value: index for index, cell in enumerate(ws[1]) if cell.value}


def _summary(path):
    """Label -> value from the summary (details and KPI tiles), and detail statuses."""
    wb = load_workbook(path)
    try:
        ws = wb["Validation Summary"]
        values = {ws.cell(r, 1).value: ws.cell(r, 2).value for r in range(1, ws.max_row + 1)}
        for col in range(1, ws.max_column + 1):  # tiles: number above label
            if ws.cell(7, col).value:
                values[ws.cell(7, col).value] = ws.cell(6, col).value
        detail = wb["Validation Detail"]
        status = _columns(detail)["Status"]
        details = [row[status] for row in detail.iter_rows(min_row=2, values_only=True)]
        return values, details
    finally:
        wb.close()


def test_every_route_and_status_is_reported(tmp_path):
    path = _write(tmp_path)
    assert path.name == "Forecast_Validation_Oct_2026.xlsx"
    values, details = _summary(path)
    assert sorted(details) == ["ACTION REQUIRED", "NOT VALIDATED", "PASS"]
    assert details[0] == "ACTION REQUIRED"  # problems are listed first
    assert values["Not Validated"] == 1
    assert values["Action Required"] == 1
    assert values["Overall Status"] == "REVIEW REQUIRED"
    assert not list(tmp_path.glob("~staging_*"))


def test_report_open_in_excel_is_saved_under_another_name(tmp_path, monkeypatch):
    first = _write(tmp_path)
    real_replace = os.replace

    def locked(src, dst):
        if os.fspath(dst) == os.fspath(first):
            raise PermissionError("in use")
        return real_replace(src, dst)

    monkeypatch.setattr(report.os, "replace", locked)
    second = _write(tmp_path)
    assert second != first and second.exists()
    assert second.name.startswith("Forecast_Validation_Oct_2026 (")
    assert not list(tmp_path.glob("~staging_*"))


def _run(folder, category, statuses, forecast_type="FY27 Forecast_102026", **flags):
    results = [
        {"category": category, "forecast_workbook": f"{category}_{i}_Oct Forecast.xlsx",
         "worksheet": "ALIV", "status": status, **flags}
        for i, status in enumerate(statuses)
    ]
    return report.write_forecast_validation_report(
        validation_result={"results": results},
        output_folder=folder,
        fiscal_year="FY27",
        forecast_month=10,
        forecast_year=2026,
        forecast_type=forecast_type,
        categories_run=[category],
    )


def _detail(path, sheet="Validation Detail"):
    wb = load_workbook(path)
    try:
        ws = wb[sheet]
        cols = _columns(ws)
        return [(row[cols["Category"]], row[cols["Status"]])
                for row in ws.iter_rows(min_row=2, values_only=True) if row[cols["Status"]]]
    finally:
        wb.close()


def test_running_one_category_keeps_the_others(tmp_path):
    _run(tmp_path, "OPEX", ["PASS", "FAIL"])
    _run(tmp_path, "Revenue", ["PASS"])
    path = _run(tmp_path, "OPEX", ["PASS", "PASS"])  # rerun OPEX after a fix
    assert _detail(path) == [("OPEX", "PASS"), ("OPEX", "PASS"), ("Revenue", "PASS")]
    values, _ = _summary(path)
    assert values["Categories Not Run Yet"] == "Direct Cost, SAC, Below the Line"
    assert values["Overall Status"] == "INCOMPLETE"


def test_results_from_another_month_are_not_merged(tmp_path):
    _run(tmp_path, "OPEX", ["FAIL"], forecast_type="FY27 Forecast_092026")
    path = _run(tmp_path, "Revenue", ["PASS"])
    assert _detail(path) == [("Revenue", "PASS")]


def test_routes_not_set_up_do_not_block_reconciled(tmp_path):
    for category in ("OPEX", "Revenue", "Direct Cost", "SAC"):
        _run(tmp_path, category, ["PASS"])
    path = _run(tmp_path, "Below the Line", ["NOT VALIDATED"], not_configured=True)
    values, _ = _summary(path)
    assert values["Not Set Up"] == 1
    assert values["Not Validated"] == 0
    assert values["Overall Status"] == "RECONCILED"


def test_needs_review_sheet_lists_only_problems(tmp_path):
    _run(tmp_path, "OPEX", ["PASS", "FAIL"])
    path = _run(tmp_path, "Revenue", ["NOT VALIDATED", "NOT VALIDATED"], not_configured=True)
    assert _detail(path, "Needs Review") == [("OPEX", "FAIL")]
    assert _detail(path) == [("OPEX", "FAIL"), ("OPEX", "PASS"), ("Revenue", "NOT SET UP"), ("Revenue", "NOT SET UP")]


def test_report_uses_one_font_throughout(tmp_path):
    from shared.excel_style import FONT_NAME

    path = _run(tmp_path, "OPEX", ["PASS", "FAIL"])
    wb = load_workbook(path)
    try:
        fonts = {cell.font.name for ws in wb.worksheets for row in ws.iter_rows() for cell in row
                 if cell.value is not None}
    finally:
        wb.close()
    assert fonts == {FONT_NAME}
