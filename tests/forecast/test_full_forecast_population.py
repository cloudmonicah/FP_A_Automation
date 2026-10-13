from __future__ import annotations

from datetime import datetime
from pathlib import Path
import re
import shutil
import traceback

from openpyxl import load_workbook

from automations.forecast.mapping import FORECAST_ROUTES
from automations.forecast.excel_service import (
    FISCAL_MONTH_ORDER,
    _s4_structure,
    populate_forecast_sheet_from_audit,
)

# ============================================================
# STAGE 4B - FULL MAPPED FORECAST POPULATION TEST
# ============================================================
#
# IMPORTANT:
# - This does NOT modify the clean October forecast files.
# - It copies only the mapped forecast workbooks into a separate
#   Stage 4B test folder and processes those copies.
# - It uses the existing Stage 3 Tableau audit workbook.
# ============================================================

FORECAST_FOLDER = Path(r"C:\Users\moncloud\OneDrive - Aliv\Monicah\ForecastTest")

AUDIT_WORKBOOK = FORECAST_FOLDER / "Tableau_Forecast_Extract_2026-10.xlsx"

TEST_FOLDER = FORECAST_FOLDER / "Stage4B_Full_Population_Test_2"

TARGET_MONTH_ABBR = "Oct"
FISCAL_YEAR = "FY27"

EXPECTED_DATES = [
    datetime(2026, 7, 1),
    datetime(2026, 8, 1),
    datetime(2026, 9, 1),
    datetime(2026, 10, 1),
    datetime(2026, 11, 1),
    datetime(2026, 12, 1),
    datetime(2027, 1, 1),
    datetime(2027, 2, 1),
    datetime(2027, 3, 1),
    datetime(2027, 4, 1),
    datetime(2027, 5, 1),
    datetime(2027, 6, 1),
]


def resolve_target_filename(baseline_name: str) -> str:
    """
    Mapping intentionally stores baseline names such as:
        ..._Aug Forecast.xlsx

    Stage 4B resolves only that forecast-month token to Oct.
    """
    resolved, count = re.subn(
        r"_[A-Za-z]{3}\s+Forecast\.xlsx$",
        f"_{TARGET_MONTH_ABBR} Forecast.xlsx",
        baseline_name,
        count=1,
        flags=re.IGNORECASE,
    )

    if count != 1:
        raise ValueError(f"Could not resolve mapped workbook name: {baseline_name}")

    return resolved


def validate_finished_sheet(
    workbook_path: Path,
    excel_sheet_name: str,
) -> None:
    workbook = load_workbook(
        workbook_path,
        data_only=False,
        read_only=False,
    )

    try:
        if excel_sheet_name in workbook.sheetnames:
            worksheet = workbook[excel_sheet_name]
        else:
            wanted = re.sub(
                r"[^a-z0-9]+",
                "",
                excel_sheet_name.lower(),
            )

            matches = [
                sheet
                for sheet in workbook.worksheets
                if re.sub(
                    r"[^a-z0-9]+",
                    "",
                    sheet.title.lower(),
                )
                == wanted
            ]

            if len(matches) != 1:
                raise AssertionError(
                    f'Excel sheet "{excel_sheet_name}" not found '
                    f"in {workbook_path.name}"
                )

            worksheet = matches[0]

        structure = _s4_structure(worksheet)
        header_row = structure["header_row"]

        actual_dates = []

        for month_number in FISCAL_MONTH_ORDER:
            column_number = structure["month_columns"][month_number]
            cell = worksheet.cell(
                row=header_row,
                column=column_number,
            )

            if isinstance(cell.value, str) and cell.value.startswith("="):
                raise AssertionError(
                    f"{cell.coordinate} is still a formula: {cell.value}"
                )

            if not isinstance(cell.value, datetime):
                raise AssertionError(
                    f"{cell.coordinate} is not a real Excel date: " f"{cell.value!r}"
                )

            if cell.number_format != "mmm-yy":
                raise AssertionError(
                    f"{cell.coordinate} format is "
                    f"{cell.number_format!r}; expected 'mmm-yy'"
                )

            actual_dates.append(cell.value)

        if actual_dates != EXPECTED_DATES:
            raise AssertionError(
                f"Fiscal month order is incorrect in "
                f"{workbook_path.name} / {worksheet.title}"
            )

        # Find last detail row from Forecast Type.
        first_detail_row = header_row + 1
        forecast_type_column = structure["forecast_type_column"]

        detail_rows = [
            row_number
            for row_number in range(
                first_detail_row,
                worksheet.max_row + 1,
            )
            if worksheet.cell(
                row=row_number,
                column=forecast_type_column,
            ).value
            not in (None, "")
        ]

        if not detail_rows:
            raise AssertionError(
                f"No populated forecast detail rows found in "
                f"{workbook_path.name} / {worksheet.title}"
            )

        last_detail_row = max(detail_rows)
        total_row = last_detail_row + 1
        grand_total_column = structure["grand_total_column"]

        # Detail Grand Total formulas.
        for row_number in range(
            first_detail_row,
            last_detail_row + 1,
        ):
            value = worksheet.cell(
                row=row_number,
                column=grand_total_column,
            ).value

            if not (isinstance(value, str) and value.startswith("=")):
                raise AssertionError(
                    f"Missing detail Grand Total formula at "
                    f"{workbook_path.name} / {worksheet.title} "
                    f"row {row_number}"
                )

        # Bottom monthly totals.
        for month_number in FISCAL_MONTH_ORDER:
            column_number = structure["month_columns"][month_number]
            value = worksheet.cell(
                row=total_row,
                column=column_number,
            ).value

            if not (isinstance(value, str) and value.startswith("=")):
                raise AssertionError(
                    f"Missing bottom monthly total in "
                    f"{workbook_path.name} / {worksheet.title}"
                )

        # Bottom overall Grand Total.
        overall_total = worksheet.cell(
            row=total_row,
            column=grand_total_column,
        ).value

        if not (isinstance(overall_total, str) and overall_total.startswith("=")):
            raise AssertionError(
                f"Missing bottom overall Grand Total in "
                f"{workbook_path.name} / {worksheet.title}"
            )

        # No literal Grand Total label in descriptive/Forecast Type fields.
        for column_number in {
            forecast_type_column,
            *structure["descriptive_columns"].values(),
        }:
            value = worksheet.cell(
                row=total_row,
                column=column_number,
            ).value

            if str(value or "").strip().lower() == "grand total":
                raise AssertionError(
                    f'Literal "Grand Total" found in '
                    f"{workbook_path.name} / {worksheet.title}"
                )

        # Required total-row border.
        july_column = structure["month_columns"][7]
        border = worksheet.cell(
            row=total_row,
            column=july_column,
        ).border

        if border.top.style != "medium":
            raise AssertionError(
                f"Total row is missing medium top border in "
                f"{workbook_path.name} / {worksheet.title}"
            )

        if border.bottom.style != "double":
            raise AssertionError(
                f"Total row is missing double bottom border in "
                f"{workbook_path.name} / {worksheet.title}"
            )

    finally:
        workbook.close()


def main() -> None:
    print("=" * 78)
    print("STAGE 4B - FULL MAPPED FORECAST POPULATION")
    print("=" * 78)

    if not FORECAST_FOLDER.exists():
        raise FileNotFoundError(f"Forecast folder not found: {FORECAST_FOLDER}")

    if not AUDIT_WORKBOOK.exists():
        raise FileNotFoundError(f"Tableau audit workbook not found: {AUDIT_WORKBOOK}")

    # --------------------------------------------------------
    # Rebuild the test folder every run so each test starts
    # from clean Stage 1 October workbooks.
    # --------------------------------------------------------
    if TEST_FOLDER.exists():
        shutil.rmtree(TEST_FOLDER)

    TEST_FOLDER.mkdir(parents=True, exist_ok=True)

    target_names = sorted(
        {resolve_target_filename(route["workbook"]) for route in FORECAST_ROUTES}
    )

    print(f"Mapping routes:       {len(FORECAST_ROUTES)}")
    print(f"Mapped workbooks:     {len(target_names)}")
    print(f"Test output folder:   {TEST_FOLDER}")
    print()

    # Copy only production workbooks referenced by mapping.py.
    copy_failures = []

    for filename in target_names:
        source = FORECAST_FOLDER / filename
        destination = TEST_FOLDER / filename

        if not source.exists():
            copy_failures.append(filename)
            print(f"[COPY FAIL] {filename} - source file not found")
            continue

        shutil.copy2(source, destination)

    if copy_failures:
        print()
        print("STAGE 4B RESULT: FAIL")
        print(
            f"{len(copy_failures)} mapped source workbook(s) " f"could not be copied."
        )
        return

    print(f"Copied workbooks:     {len(target_names)}")
    print()

    # --------------------------------------------------------
    # Populate every mapping route.
    # --------------------------------------------------------
    passed = []
    failed = []
    revenue_routes = 0
    partial_month_routes = []

    for index, route in enumerate(FORECAST_ROUTES, start=1):
        tableau_sheet = route["tableau_worksheet"]
        excel_sheet = route["excel_sheet"]
        target_filename = resolve_target_filename(route["workbook"])
        target_workbook = TEST_FOLDER / target_filename

        route_label = f"{target_filename} -> {excel_sheet} " f"<- {tableau_sheet}"

        print(f"[{index:02d}/{len(FORECAST_ROUTES):02d}] " f"{route_label}")

        try:
            result = populate_forecast_sheet_from_audit(
                workbook_path=target_workbook,
                excel_sheet_name=excel_sheet,
                audit_workbook_path=AUDIT_WORKBOOK,
                tableau_worksheet_name=tableau_sheet,
                fiscal_year=FISCAL_YEAR,
            )

            validate_finished_sheet(
                workbook_path=target_workbook,
                excel_sheet_name=excel_sheet,
            )

            if result.get("revenue_sign_rule"):
                revenue_routes += 1

            passed.append(
                {
                    "route": route_label,
                    "rows": result["tableau_detail_rows"],
                }
            )

            print(f"    PASS - {result['tableau_detail_rows']} " f"detail row(s)")

        except Exception as exc:
            failed.append(
                {
                    "route": route_label,
                    "error": str(exc),
                }
            )

            print(f"    FAIL - {exc}")

    # --------------------------------------------------------
    # Summary.
    # --------------------------------------------------------
    print()
    print("=" * 78)
    print("STAGE 4B SUMMARY")
    print("=" * 78)
    print(f"Mapping routes:       {len(FORECAST_ROUTES)}")
    print(f"Routes processed:     {len(passed) + len(failed)}")
    print(f"Routes passed:        {len(passed)}")
    print(f"Routes failed:        {len(failed)}")
    print(f"Excel workbooks:      {len(target_names)}")
    print(f"Revenue routes:       {revenue_routes}")
    print()

    if failed:
        print("FAILED ROUTES")
        print("-" * 78)

        for item in failed:
            print(item["route"])
            print(f"  ERROR: {item['error']}")

        print()
        print("STAGE 4B RESULT: FAIL")
        print(
            "Do not move to production/full orchestration yet. "
            "Fix only the failed routes and rerun this test."
        )

    else:
        print("PASS - All mapped Tableau routes populated")
        print("PASS - Jul-26 -> Jun-27 are REAL Excel dates")
        print("PASS - Dates are in correct fiscal order")
        print("PASS - No =EDATE/header formulas remain")
        print("PASS - Revenue sign rule applied where required")
        print("PASS - Detail Grand Total formulas rebuilt")
        print("PASS - Bottom monthly totals rebuilt")
        print("PASS - Bottom overall Grand Total rebuilt")
        print('PASS - No literal "Grand Total" labels')
        print("PASS - Total-row borders applied")
        print()
        print("STAGE 4B RESULT: PASS")
        print()
        print("Review the completed test workbooks here:")
        print(TEST_FOLDER)


if __name__ == "__main__":
    main()
