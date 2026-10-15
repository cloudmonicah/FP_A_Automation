from openpyxl import load_workbook
from openpyxl.styles import Font

from automations.forecast.excel_service import (
    populate_forecast_sheet_from_tableau_rows,
    update_forecast_workbook,
)
from shared.excel_style import BODY_SIZE, FONT_NAME
from tests.forecast.conftest import write_forecast_workbook
from tests.forecast.test_forecast_months import FY27


def _mixed_fonts(path):
    wb = load_workbook(path)
    ws = wb["ALIV"]
    ws["A1"].font = Font(name="Calibri", size=14, bold=True)  # a title
    for row in ws.iter_rows(min_row=3):
        for cell in row:
            cell.font = Font(name="Arial", size=9, bold=cell.row == 3)
    wb.save(path)


def _fonts(path):
    wb = load_workbook(path)
    ws = wb["ALIV"]
    fonts = {(c.font.name, c.font.sz, bool(c.font.b)) for row in ws.iter_rows(min_row=2) for c in row if c.value is not None}
    title = (ws["A1"].font.name, ws["A1"].font.sz, ws["A1"].font.b)
    wb.close()
    return fonts, title


def test_step_2_writes_one_font_and_size(tmp_path):
    path = write_forecast_workbook(tmp_path / "x.xlsx")
    _mixed_fonts(path)
    rows = [{"GL Account": "6000", **{h: "1" for h in FY27}}]
    populate_forecast_sheet_from_tableau_rows(path, "ALIV", rows, "FY27")
    fonts, title = _fonts(path)
    assert {(name, size) for name, size, _ in fonts} == {(FONT_NAME, BODY_SIZE)}
    assert (FONT_NAME, BODY_SIZE, True) in fonts  # bold header kept
    assert title == (FONT_NAME, 14, True)  # title keeps its size


def test_step_1_gives_the_new_workbook_one_font(tmp_path):
    path = write_forecast_workbook(tmp_path / "x.xlsx", forecast_type="FY27 Forecast_092026")
    _mixed_fonts(path)
    update_forecast_workbook(path, "FY27", 10, 2026)
    fonts, _ = _fonts(path)
    assert {(name, size) for name, size, _ in fonts} == {(FONT_NAME, BODY_SIZE)}
