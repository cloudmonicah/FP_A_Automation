from automations.forecast import preflight
from automations.forecast.service import build_forecast_period
from tests.forecast.conftest import FORECAST_TYPE_TAB, write_forecast_workbook
from tests.forecast.test_forecast_population import CONSOLIDATED

PERIOD = build_forecast_period("FY27", 10, 2026)


def _route(**overrides):
    route = {
        "category": "OPEX",
        "tableau_worksheet": "ALIV Legal",
        "workbook": "FY27 OPEX Legal_Aug Forecast.xlsx",
        "excel_sheet": "ALIV",
        "individual_tableau_workbook": "Legal.twb",
        "individual_tableau_worksheet": "Monthly Forecast ",
        "validation_status": "confirmed",
    }
    route.update(overrides)
    return route


def test_update_check_finds_each_problem_without_tableau(tmp_path):
    twb = tmp_path / "FY27 FORECAST.twb"
    twb.write_text(CONSOLIDATED, encoding="utf-8")
    out = tmp_path / "out"
    write_forecast_workbook(out / "FY27 OPEX Legal_Oct Forecast.xlsx", sheet="ALIV")
    write_forecast_workbook(out / "FY27 OPEX HR_Oct Forecast.xlsx", sheet="ALIV", forecast_type="FY27 Forecast_092026")
    routes = [
        _route(),  # fine
        _route(excel_sheet="CBL"),  # sheet missing
        _route(tableau_worksheet="CBL Legal"),  # tab missing in consolidated
        _route(workbook="FY27 OPEX HR_Aug Forecast.xlsx"),  # September file
        _route(workbook="FY27 OPEX IT_Aug Forecast.xlsx"),  # not created
    ]
    problems = [p["problem"] for p in preflight.check_before_update(routes, out, twb, PERIOD)]
    assert len(problems) == 4
    assert 'no sheet named "CBL"' in problems[0]
    assert 'no tab called "CBL Legal"' in problems[1]
    assert "September 2026" in problems[2] and "replace" in problems[2]
    assert "Run step 1" in problems[3]


def test_validation_check_needs_step_2_and_the_individual_tab(tmp_path, twb):
    source = twb(FORECAST_TYPE_TAB)
    out = tmp_path / "out"
    write_forecast_workbook(out / "FY27 OPEX Legal_Oct Forecast.xlsx", sheet="ALIV")
    folder = str(source.parent)
    routes = [
        _route(individual_tableau_folder=folder),  # fine (tab name has a trailing space in the mapping)
        _route(individual_tableau_folder=folder, individual_tableau_worksheet="CBL Monthly Forecast"),
        _route(individual_tableau_folder=folder, individual_tableau_workbook="Missing.twb"),
        _route(validation_status="needs_filter"),  # skipped: reported by the run
    ]
    problems = [p["problem"] for p in preflight.check_before_validation(routes, out, PERIOD)]
    assert len(problems) == 2
    assert 'no tab called "CBL Monthly Forecast"' in problems[0]
    assert "Missing.twb" in problems[1]


def test_validation_check_flags_sheets_step_2_has_not_finished(tmp_path, twb):
    from openpyxl import load_workbook

    path = write_forecast_workbook(tmp_path / "out" / "FY27 OPEX Legal_Oct Forecast.xlsx", sheet="ALIV")
    wb = load_workbook(path)
    wb["ALIV"].cell(3, 3).value = "Jul"  # month header still text, as before step 2
    wb.save(path)
    source = twb(FORECAST_TYPE_TAB)
    problems = preflight.check_before_validation(
        [_route(individual_tableau_folder=str(source.parent))], tmp_path / "out", PERIOD
    )
    assert "Run step 2" in problems[0]["problem"]
