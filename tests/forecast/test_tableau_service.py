from pathlib import Path
import re

from openpyxl import Workbook

from automations.forecast.mapping import FORECAST_ROUTES
from automations.forecast.tableau_service import (
    get_tableau_crosstab_rows,
)

# ============================================================
# CONFIGURATION
# ============================================================

TABLEAU_WORKBOOK = Path(
    r"E:\Workbooks\FY27 Forecast Workbooks\FY27 FORECAST WORKBOOKS.twb"
)

OUTPUT_FOLDER = Path(r"C:\Users\moncloud\OneDrive - Aliv\Monicah\ForecastTest")

FORECAST_MONTH = 10
FORECAST_YEAR = 2026


# ============================================================
# SAFE EXCEL SHEET NAME
# ============================================================


def safe_sheet_name(name, used_names):

    cleaned_name = re.sub(
        r"[:\\/?*\[\]]",
        "_",
        name,
    ).strip()

    if not cleaned_name:
        cleaned_name = "Tableau"

    base_name = cleaned_name[:31]

    candidate = base_name
    counter = 2

    while candidate.lower() in used_names:

        suffix = f"_{counter}"

        candidate = base_name[: 31 - len(suffix)] + suffix

        counter += 1

    used_names.add(candidate.lower())

    return candidate


# ============================================================
# WRITE TABLEAU CROSSTAB
# ============================================================


def write_rows(ws, rows):
    """
    Write exactly what Tableau returned.

    No extra columns.
    No extra rows.
    No totals.
    No notes.
    """

    if not rows:
        return

    headers = list(rows[0].keys())

    # Headers
    for column_number, header in enumerate(
        headers,
        start=1,
    ):
        ws.cell(
            row=1,
            column=column_number,
            value=header,
        )

    # Data
    for row_number, row in enumerate(
        rows,
        start=2,
    ):

        for column_number, header in enumerate(
            headers,
            start=1,
        ):

            ws.cell(
                row=row_number,
                column=column_number,
                value=row.get(header),
            )


# ============================================================
# UNIQUE TABLEAU WORKSHEETS
# ============================================================


def get_unique_tableau_worksheets():

    worksheet_names = []
    seen = set()

    for route in FORECAST_ROUTES:

        worksheet_name = str(route.get("tableau_worksheet") or "").strip()

        if not worksheet_name:
            continue

        normalized = " ".join(worksheet_name.lower().split())

        if normalized in seen:
            continue

        seen.add(normalized)

        worksheet_names.append(worksheet_name)

    return worksheet_names


# ============================================================
# MAIN
# ============================================================


def main():

    print()
    print("=" * 72)
    print("STAGE 3C - CONSOLIDATED TABLEAU EXTRACTION")
    print("=" * 72)

    if not TABLEAU_WORKBOOK.exists():

        raise FileNotFoundError(f"Tableau workbook not found:\n" f"{TABLEAU_WORKBOOK}")

    OUTPUT_FOLDER.mkdir(
        parents=True,
        exist_ok=True,
    )

    worksheet_names = get_unique_tableau_worksheets()

    print(f"Mapping routes             : " f"{len(FORECAST_ROUTES)}")

    print(f"Unique Tableau worksheets  : " f"{len(worksheet_names)}")

    print(f"Tableau workbook           : " f"{TABLEAU_WORKBOOK}")

    print()
    print("IMPORTANT:")
    print("Open FY27 FORECAST WORKBOOKS.twb in Tableau Desktop")
    print("and maximize the Tableau window before continuing.")
    print()

    input("Press ENTER when Tableau is fully open and maximized...")

    # ========================================================
    # OUTPUT WORKBOOK
    # ========================================================

    output_file = OUTPUT_FOLDER / (
        f"Tableau_Forecast_Extract_" f"{FORECAST_YEAR}-" f"{FORECAST_MONTH:02d}.xlsx"
    )

    wb = Workbook()

    default_ws = wb.active

    used_names = set()

    passed = 0
    failed = 0

    failures = []

    # ========================================================
    # EXTRACT
    # ========================================================

    for sequence, worksheet_name in enumerate(
        worksheet_names,
        start=1,
    ):

        print()
        print("-" * 72)

        print(f"[{sequence}/{len(worksheet_names)}] " f"{worksheet_name}")

        try:

            rows = get_tableau_crosstab_rows(
                worksheet_name=worksheet_name,
                tableau_workbook_path=TABLEAU_WORKBOOK,
            )

            if not rows:

                raise RuntimeError("Tableau returned no rows.")

            sheet_name = safe_sheet_name(
                worksheet_name,
                used_names,
            )

            ws = wb.create_sheet(title=sheet_name)

            write_rows(
                ws,
                rows,
            )

            passed += 1

            print(f"PASS - " f"{len(rows)} rows x " f"{len(rows[0])} columns")

        except Exception as exc:

            failed += 1

            failures.append(
                (
                    worksheet_name,
                    str(exc),
                )
            )

            print(f"FAILED - {exc}")

    # ========================================================
    # REMOVE DEFAULT SHEET
    # ========================================================

    if passed > 0:

        try:
            wb.remove(default_ws)
        except Exception:
            pass

    # ========================================================
    # SAVE
    # ========================================================

    if passed > 0:

        wb.save(output_file)

    # ========================================================
    # RESULTS
    # ========================================================

    print()
    print("=" * 72)
    print("STAGE 3C RESULTS")
    print("=" * 72)

    print(f"Unique worksheets : {len(worksheet_names)}")

    print(f"Passed            : {passed}")

    print(f"Failed            : {failed}")

    if passed > 0:

        print(f"Audit workbook    : {output_file}")

    if failures:

        print()
        print("FAILED WORKSHEETS")
        print("-" * 72)

        for worksheet_name, error in failures:

            print()
            print(worksheet_name)
            print(f"  {error}")

    print()
    print("=" * 72)

    if failed == 0:
        print("STAGE 3C RESULT: PASS")
    else:
        print("STAGE 3C RESULT: REVIEW REQUIRED")

    print("=" * 72)


# ============================================================
# RUN
# ============================================================

if __name__ == "__main__":
    main()
