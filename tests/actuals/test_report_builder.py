from openpyxl import Workbook, load_workbook

from automations.actuals.report_builder import (
    STATUS_PASS,
    STATUS_PENDING,
    STATUS_REVIEW,
    PeriodValues,
    finalize_report,
    read_report,
    report_status,
    update_individual_values,
)

LINES = [
    # name, monthly (fs, cons, ind), ytd (fs, cons, ind), note
    ("Legal", (100.0, 100.4, 100.4), (800.0, 800.0, 800.0), "Reconciled."),
    ("Facilities", (50.0, 50.0, 47.0), (400.0, 400.0, 400.0), "Mapped to two lines"),
    ("Operating Expenses", (150.0, 150.4, None), (1200.0, 1200.0, None), None),
]


def _capture_workbook(path):
    """The A:S layout written by the OPEX/Direct Cost/Revenue capture step."""
    wb = Workbook()
    wb.active.title = "Executive Summary"
    for company in ("ALIV", "CBL"):
        ws = wb.create_sheet(f"{company} OPEX")
        ws["C2"] = "August 2026"
        ws["C3"] = "09/01/2026"
        ws["A6"] = "Financial Statement Line"
        for row, (name, monthly, ytd, note) in enumerate(LINES, 7):
            ws.cell(row, 1, name)
            ws.cell(row, 2, monthly[0])
            ws.cell(row, 3, monthly[1])
            ws.cell(row, 12, monthly[2])
            ws.cell(row, 8, ytd[0])
            ws.cell(row, 9, ytd[1])
            ws.cell(row, 16, ytd[2])
            ws.cell(row, 6, note)
    wb.save(path)


def test_period_status():
    assert PeriodValues(100, 100.9, 100).status == STATUS_PASS
    assert PeriodValues(100, 102, 102).status == STATUS_REVIEW
    assert PeriodValues(100, 100, None).status == STATUS_PENDING
    assert PeriodValues(100, 102, 102).v1 == 2
    assert PeriodValues(100, 102, 99).v2 == -3


def test_finalize_builds_final_layout(tmp_path):
    path = tmp_path / "report.xlsx"
    _capture_workbook(path)

    finalize_report(path, forecast_type="5+7", forecast_start="09/01/2026")

    wb = load_workbook(path)
    assert wb.sheetnames == ["Executive Summary", "ALIV OPEX", "CBL OPEX", "_meta"]
    ws = wb["ALIV OPEX"]
    assert ws["A6"].value == "Financial Statement Line"
    assert ws["G7"].value == STATUS_PASS  # Legal monthly
    assert ws["G8"].value == STATUS_REVIEW  # Facilities monthly, V2 = -3
    assert ws["F8"].value == -3
    assert ws["G9"].value == STATUS_PENDING  # no individual value
    assert ws["H8"].value.startswith("Mapped to two lines")
    assert ws["H7"].value is None  # "Reconciled." is dropped
    assert report_status(path) == STATUS_REVIEW

    meta, sheets = read_report(path)
    assert meta.closed_period == "August 2026"
    assert meta.forecast_type == "5+7"
    assert [line.name for line in sheets[0].lines] == [l[0] for l in LINES]


def test_finalize_is_repeatable(tmp_path):
    path = tmp_path / "report.xlsx"
    _capture_workbook(path)
    finalize_report(path, forecast_type="5+7")
    first = read_report(path)
    finalize_report(path)
    second = read_report(path)
    assert first[1] == second[1]
    assert second[0].forecast_type == "5+7"


def test_update_individual_values_recomputes_status(tmp_path):
    path = tmp_path / "report.xlsx"
    _capture_workbook(path)
    finalize_report(path)

    outcome = update_individual_values(
        path,
        "ALIV",
        {"Facilities": (50.0, 400.0, None), "Operating Expenses": (150.0, 1200.0, None), "Unknown": (1, 1, None)},
    )

    assert outcome == {"Facilities": "updated", "Operating Expenses": "updated", "Unknown": "missing"}
    _, sheets = read_report(path)
    aliv = sheets[0]
    assert aliv.status == STATUS_PASS
    assert sheets[1].status == STATUS_REVIEW  # CBL untouched
    assert report_status(path) == STATUS_REVIEW
