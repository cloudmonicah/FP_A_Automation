from datetime import date

import pytest
from openpyxl import Workbook

from automations.actuals.adaptive_aliv_statement_reader import read_aliv_period
from automations.actuals.direct_cost_financial_statement_reader import read_aliv_direct_cost


def _save(tmp_path, build):
    wb = Workbook()
    ws = wb.active
    ws.title = "P & L_"
    build(ws)
    path = tmp_path / "aliv.xlsx"
    wb.save(path)
    return path


LABELS = (
    "Interconnect",
    "Reseller Commissions",
    "Roaming Service Costs",
    "Service Fees",
    "Direct Cost",
)


def _lines(ws, row, month_col, ytd_col, month_total, ytd_total):
    """Five Direct Cost lines; the section total is the last row."""
    for offset, label in enumerate(LABELS):
        last = offset == len(LABELS) - 1
        ws.cell(row + offset, 2, label)
        ws.cell(row + offset, month_col, month_total if last else 1.0)
        ws.cell(row + offset, ytd_col, ytd_total if last else 1.0)


def august_layout(ws):
    ws["D1"] = date(2026, 8, 1)
    ws["C5"], ws["G5"], ws["H3"] = "Actual ", "Actual ", "Actual vs. Budget YTD"
    _lines(ws, 14, 3, 7, 863.0, 1947.0)


def july_layout(ws):
    # 12-month trend block first; month Actual is column M, YTD Actual column P.
    ws["J5"] = date(2026, 7, 31)
    ws["C9"] = date(2025, 7, 1)
    ws["M11"], ws["P11"], ws["Q9"] = "Actual ", "Actual ", "Actual vs. Budget YTD"
    _lines(ws, 20, 3, 3, 938.0, 938.0)  # July 2025 trend column
    _lines(ws, 20, 13, 16, 1084.0, 9999.0)  # real July / odd YTD


def test_august_layout_uses_actual_column(tmp_path):
    path = _save(tmp_path, august_layout)
    assert read_aliv_direct_cost(path, date(2026, 8, 1))["Direct Cost"] == 863.0


def test_july_layout_uses_header_columns_not_column_c(tmp_path):
    path = _save(tmp_path, july_layout)
    assert read_aliv_direct_cost(path, date(2026, 7, 1))["Direct Cost"] == 1084.0
    # In July the YTD equals the monthly actual (start of the fiscal year).
    assert read_aliv_period(path, date(2026, 7, 1))["DIRECT COST"] == (1084.0, 1084.0)


def test_wrong_month_is_rejected(tmp_path):
    path = _save(tmp_path, july_layout)
    with pytest.raises(ValueError, match="selected close"):
        read_aliv_period(path, date(2026, 8, 1))
