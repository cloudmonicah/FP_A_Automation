from datetime import date, datetime

from openpyxl import Workbook

from automations.actuals.cbl_statement_periods import read_cbl_area


def _statement(path):
    wb = Workbook()
    fin = wb.active
    fin.title = "Financials"
    fin["C2"] = datetime(2026, 9, 30)
    fin.append([])
    rows = [
        (None, None, "Actual YTD"),
        (None, "Income Statement ", None),
        (None, "Gross Revenues", None),
        (None, "Residential", 18000.0),
        (None, "Smart Home ", 11.0),
        (None, "MSA-ALIV", 3000.0),
        (None, "Total Service Revenues", 21011.0),
        (None, "Corporate - One off", 1000.0),
        (None, "Gross Revenues", 22011.0),
        (None, "Cost of Goods Sold", None),
        (None, "Programming fees", 4000.0),
        (None, "Cost of Goods Sold", 4000.0),
    ]
    for row in rows:
        fin.append(list(row))
    soe = wb.create_sheet("Statement of Earnings")
    soe["D3"] = datetime(2026, 9, 30)
    soe["D6"] = "Current Month"
    soe["D8"] = "Actual"
    data = wb.create_sheet("Data")
    data.append([])
    header = [None, "x", None, None, None, "CY Actual"]
    data.append(header)
    data.append([None, None, None, None, None, datetime(2026, 9, 1)])
    for label, value in (
        ("Revenue Breakdown for P&L Summary", None),
        ("Residential", 6000.0),
        ("Other", 5.0),
        ("Intercompany", 1000.0),
        ("Corporate- One Off", 250.0),
        ("Gross Revenue", 7255.0),
    ):
        data.append([None, label if value is None else None, None if value is None else label, None, None, value])
    data.append([None, "Opex Breakdown for Income stmt mthly"])
    wb.save(path)


def test_september_names_are_matched(tmp_path):
    path = tmp_path / "cbl.xlsx"
    _statement(path)
    result = read_cbl_area(path, "Revenue", date(2026, 9, 1))
    assert result["YTD"]["Gross Revenue"] == 22011.0
    assert result["YTD"]["Other"] == 11.0  # "Smart Home" in the statement
    assert "Total Service Revenues" not in result["YTD"]
    assert result["MONTH"]["MSA-ALIV"] == 1000.0  # "Intercompany" on the Data tab
    assert result["MONTH"]["Other"] == 5.0
    assert result["MONTH"]["Gross Revenue"] == 7255.0
