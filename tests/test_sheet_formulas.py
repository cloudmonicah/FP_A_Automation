"""Totals calculated without Excel match what Excel shows."""

import pytest
from openpyxl import Workbook

from shared.sheet_formulas import SheetCalculator, UnsupportedFormula


def test_sums_ranges_cells_and_arithmetic():
    ws = Workbook().active
    ws["A1"], ws["B1"], ws["C1"] = 1, 2.5, "=SUM(A1:B1)"
    ws["A2"], ws["B2"], ws["C2"] = -3, None, "=SUM(A2,B2)*2-(1)"
    ws["C3"] = "=SUM(C1:C2)+$A$1"
    assert SheetCalculator(ws).value(3, 3) == -2.5


def test_amount_formatted_as_a_date_counts_as_its_number():
    ws = Workbook().active
    ws["A1"] = 67500
    ws["A1"].number_format = "mmm-yy"        # shows "Oct-84" in Excel
    ws["A2"] = "=SUM(A1:A1)"
    path_ws = ws
    # openpyxl reads such cells back as dates; simulate by saving and reloading.
    from io import BytesIO

    from openpyxl import load_workbook

    buffer = BytesIO()
    path_ws.parent.save(buffer)
    reloaded = load_workbook(BytesIO(buffer.getvalue())).active
    assert hasattr(reloaded["A1"].value, "year")
    assert SheetCalculator(reloaded).value(2, 1) == 67500


def test_text_and_blanks_are_ignored_like_excel():
    ws = Workbook().active
    ws["A1"], ws["A2"], ws["A3"] = "n/a", None, 5
    ws["A4"] = "=SUM(A1:A3)"
    assert SheetCalculator(ws).value(4, 1) == 5


def test_other_functions_are_reported():
    ws = Workbook().active
    ws["A1"] = "=VLOOKUP(1,B1:C2,2)"
    with pytest.raises(UnsupportedFormula):
        SheetCalculator(ws).value(1, 1)
